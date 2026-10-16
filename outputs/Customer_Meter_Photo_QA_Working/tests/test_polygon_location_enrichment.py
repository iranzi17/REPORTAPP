"""Regression tests for read-only polygon location spreadsheet enrichment."""
from __future__ import annotations
import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from meter_photo_qa.polygon_location_enrichment import enrich_report,_prepare,_make_live_resolver


def sample(path,rows,group=7):
    book=Workbook()
    summary=book.active
    summary.title="Summary"
    summary.append(["Feature Layer", 12 if group==7 else 7])
    summary.append(["Asset Group",group])
    summary.append(["Location source","Recorded fields only"])
    data=book.create_sheet("Asset Details")
    data.append(["Status","OBJECTID","GlobalID","Asset Type",
                 "Surveyor Username","Recorded District","Recorded Sector",
                 "Recorded Cell","Outstanding Issues","Unverified Checks"])
    for row in rows:
        data.append(row)
    data.add_table(Table(displayName="LVPoleDetails",ref="A1:J"+str(data.max_row)))
    book.create_sheet("By Recorded Location").append(["Recorded District","Assets Checked"])
    book.save(path)
    book.close()


class PolygonLocationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.folder=Path(self.tmp.name)
        self.input=self.folder/"pole_quality.xlsx"
        sample(self.input,[
            ["GAP",3910247,"{A}","Wooden","surveyor",None,None,None,"Missing",""],
            ["GAP",3911043,"{B}","Wooden","surveyor","Gakenke",None,None,"Missing",""],
            ["GAP",3911446,"{C}","Wooden","surveyor","Rulindo",None,None,"Missing",""],
            ["GAP",3915443,"{D}","Wooden","surveyor",None,None,None,"Missing",""],
            ["NO_GAPS_FOUND",33,"{E}","Wooden","surveyor","Burera","Cyanika","Gasiza","",""],
        ])
        self.before=self.input.read_bytes()

    def tearDown(self):
        self.tmp.cleanup()

    def locator(self,oid,gid):
        return {
            3910247:("POLYGON_MATCH",("Gakenke","Nemba","Mugunga")),
            3911043:("POLYGON_MATCH",("Gakenke","Nemba","Mugunga")),
            3911446:("POLYGON_MATCH",("Gakenke","Nemba","Mugunga")),
            3915443:("AMBIGUOUS_POLYGON_MATCH",None),
        }[oid]

    def test_dry_run_never_writes_files(self):
        results=enrich_report(self.input,resolver=self.locator)
        self.assertEqual(results["status"],"DRY_RUN")
        self.assertEqual(results["records_filled"],2)
        self.assertEqual(results["cells_filled"],5)
        self.assertFalse(list(self.folder.glob("*enriched*")))
        self.assertEqual(self.input.read_bytes(),self.before)

    def test_fill_blank_only_and_audit_conflicts(self):
        result=enrich_report(self.input,resolver=self.locator,apply=True)
        self.assertEqual(result["status"],"OUTPUT_CREATED")
        self.assertEqual(result["records_filled"],2)
        self.assertEqual(self.input.read_bytes(),self.before)
        wb=load_workbook(result["output_xlsx"])
        try:
            sh=wb["Asset Details"]
            self.assertEqual([sh.cell(2,i).value for i in (6,7,8)],
                             ["Gakenke","Nemba","Mugunga"])
            self.assertEqual([sh.cell(3,i).value for i in (6,7,8)],
                             ["Gakenke","Nemba","Mugunga"])
            self.assertEqual([sh.cell(4,i).value for i in (6,7,8)],
                             ["Rulindo",None,None])
            self.assertEqual([sh.cell(5,i).value for i in (6,7,8)],
                             [None,None,None])
            self.assertEqual(sh.cell(6,8).value,"Gasiza")
            self.assertEqual(sh.cell(2,11).value,"POLYGON_MATCH")
            self.assertEqual(sh.cell(4,11).value,"LOCATION_CONFLICT_REVIEW")
            self.assertEqual(sh.cell(5,11).value,"AMBIGUOUS_POLYGON_MATCH")
            audit=wb["Polygon Location Audit"]
            self.assertEqual(audit.cell(2,2).value,"a")
            self.assertEqual(audit.cell(4,12).value,"LOCATION_CONFLICT_REVIEW")
            self.assertIn("LVPoleDetails",sh.tables)
            self.assertEqual(sh.tables["LVPoleDetails"].ref,"A1:J6")
        finally: wb.close()

    def test_no_overwriting_or_reenrichment(self):
        output=self.folder/"filled.xlsx"
        enrich_report(self.input,resolver=self.locator,output_path=output,apply=True)
        with self.assertRaises(FileExistsError):
            enrich_report(self.input,resolver=self.locator,output_path=output,apply=True)
        with self.assertRaises(ValueError):
            enrich_report(output,apply=False,resolver=self.locator)

    def test_unsupported_asset_group_fails_closed(self):
        bad=self.folder/"line.xlsx"
        sample(bad,[["GAP",12,"{L}","Line","user",None,None,None,"",""]],group=10)
        with self.assertRaisesRegex(ValueError,"point assets only"):
            enrich_report(bad,resolver=lambda *_: self.fail("unexpected"))

    def test_malformed_report_blocks_without_gis_calls(self):
        bad=self.folder/"unknown.xlsx"
        book=Workbook()
        book.save(bad)
        book.close()
        with self.assertRaisesRegex(ValueError,"REPORTAPP asset QA"):
            enrich_report(bad,resolver=lambda *_: self.fail("unexpected"))

    def test_no_guess_with_invalid_polygons(self):
        with self.assertRaisesRegex(ValueError,"incomplete"):
            enrich_report(self.input,resolver=lambda *_:("POLYGON_MATCH",("Nemba","","")),
                          apply=True)
        self.assertEqual(self.input.read_bytes(),self.before)

    def test_source_location_lookup_checks_objectids_and_source_group(self):
        module_arcgis=types.ModuleType("arcgis")
        feature_module=types.ModuleType("arcgis.features")
        geo_module=types.ModuleType("arcgis.geometry")
        filters_module=types.ModuleType("arcgis.geometry.filters")
        request_log=[]
        point_rows=[
            SimpleNamespace(attributes={"OBJECTID":3910247,"GLOBALID":"{A}","ASSETGROUP":7},
                            geometry={"x":30.2,"y":-1.6,"spatialReference":{"wkid":4326}}),
            SimpleNamespace(attributes={"OBJECTID":3911043,"GLOBALID":"{B}","ASSETGROUP":13},
                            geometry={"x":30.3,"y":-1.7,"spatialReference":{"wkid":4326}}),
        ]
        def layer(url,gis):
            if url.endswith("/12"):
                return SimpleNamespace(query=lambda **kwargs:(
                    request_log.append(("source",kwargs)) or SimpleNamespace(features=point_rows)))
            if url.endswith("/1"):
                return SimpleNamespace(query=lambda **kwargs:(
                    request_log.append(("cells",kwargs)) or SimpleNamespace(features=[
                        SimpleNamespace(attributes={"district":"Gakenke","sector":"Nemba","cell":"Mugunga"})
                    ])))
            raise AssertionError("Wrong source layer: "+url)
        feature_module.FeatureLayer=layer
        filters_module.intersects=lambda point,sr:{"point":point,"sr":sr}
        with patch.dict(sys.modules,{
            "arcgis":module_arcgis,"arcgis.features":feature_module,
            "arcgis.geometry":geo_module,"arcgis.geometry.filters":filters_module,
        }),patch("meter_photo_qa.progress_report._get_gis",return_value=object()):
            locator=_make_live_resolver([
                {"objectid":3910247,"globalid":"a"},
                {"objectid":3911043,"globalid":"b"},
            ],7)
            self.assertEqual(locator(3910247,"a"),
                             ("POLYGON_MATCH",("Gakenke","Nemba","Mugunga")))
            self.assertEqual(locator(3911043,"b"),("WRONG_ASSET_GROUP",None))
            self.assertEqual(len([x for x in request_log if x[0]=="cells"]),1)
            self.assertIn("gdb_version",request_log[0][1])
            self.assertEqual(request_log[0][1]["out_sr"],4326)
            self.assertEqual(request_log[0][1]["object_ids"],"3910247,3911043")

if __name__=="__main__":
    unittest.main()
