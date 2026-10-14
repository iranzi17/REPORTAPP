from __future__ import annotations
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook

from meter_photo_qa.geometry_feeder_runner import (
    _get_nationwide_line_ids,run_provisional_geometry_feeder_report,
)
from meter_photo_qa.lv_feeder_report import run_lv_feeder_report


def feature(attrs, geom):
    return SimpleNamespace(attributes=attrs,geometry=geom)


class FakeLayer:
    def __init__(self,kind):
        self.kind=kind
        self.properties=SimpleNamespace(fields=[
            {"name":name,"alias":name} for name in (
                ["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE",
                 "created_user","Transformer_Name","Phases"]
                if kind=="line" else
                ["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE","Transfomer_name_"]
            )
        ])
        self.calls=[]

    def query(self,**opts):
        self.calls.append(opts)
        if self.kind=="line":
            rows=[
                feature({"OBJECTID":1,"GLOBALID":"{L1}","ASSETGROUP":9,"ASSETTYPE":1,
                         "Transformer_Name":"Central","created_user":"ninamenye9683"},
                        {"paths":[[[0,0],[10,0]]],"spatialReference":{"wkid":3857}}),
                feature({"OBJECTID":2,"GLOBALID":"{L2}","ASSETGROUP":9,"ASSETTYPE":1,
                         "Transformer_Name":"Central","created_user":"ninamenye9683"},
                        {"paths":[[[10,0],[20,0]]],"spatialReference":{"wkid":3857}}),
                feature({"OBJECTID":3,"GLOBALID":"{L3}","ASSETGROUP":9,"ASSETTYPE":1,
                         "Transformer_Name":"Central","created_user":"other"},
                        {"paths":[[[20,0],[30,0]]],"spatialReference":{"wkid":3857}}),
                feature({"OBJECTID":4,"GLOBALID":"{L4}","ASSETGROUP":9,"ASSETTYPE":1,
                         "Transformer_Name":"Central","created_user":"ninamenye9683"},
                        {"paths":[[[90,0],[100,0]]],"spatialReference":{"wkid":3857}}),
            ]
            filtered=[x for x in rows if (
                "created_user IN" not in opts.get("where","")
                or x.attributes["created_user"]=="ninamenye9683"
            )]
            if opts.get("return_count_only"):
                return len(filtered)
            if opts.get("return_ids_only"):
                return {"objectIds":[v.attributes["OBJECTID"] for v in filtered]}
            ids={int(q) for q in opts.get("object_ids","").split(",") if q}
            return SimpleNamespace(features=[x for x in rows
                     if x.attributes["OBJECTID"] in ids],
                     spatial_reference={"wkid":3857})
        if self.kind=="transformer":
            return SimpleNamespace(features=[
                feature({
                    "OBJECTID":100,"GLOBALID":"{T100}","ASSETGROUP":12,
                    "ASSETTYPE":4,"Transfomer_name_":"Central"
                },{"x":0,"y":0,"spatialReference":{"wkid":3857}})
            ],spatial_reference={"wkid":3857})
        raise AssertionError(self.kind)


class GeometryFeederRunnerTests(unittest.TestCase):
    def test_default_feeder_report_uses_no_trace_and_crosses_boundary(self):
        line_layer=FakeLayer("line")
        transformer_layer=FakeLayer("transformer")
        layers={"line":line_layer,"transformer":transformer_layer}
        feature_mod=types.ModuleType("arcgis.features")
        feature_mod.FeatureLayer=lambda url,gis: (
            layers["line"] if url.endswith("/10") else
            layers["transformer"] if url.endswith("/7") else
            FakeLayer("boundary")
        )
        start=[{"OBJECTID":100,"GLOBALID":"{T100}","ASSETGROUP":12,
                 "ASSETTYPE":4,"Transfomer_name_":"Central"}]
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch.dict(sys.modules,{"arcgis":types.ModuleType("arcgis"),
                    "arcgis.features":feature_mod}), \
             patch("meter_photo_qa.geometry_feeder_runner._get_sector_features",
                   return_value=[object()]), \
             patch("meter_photo_qa.geometry_feeder_runner._filters_for_area",
                   return_value=[("Gicumbi",object())]), \
             patch("meter_photo_qa.geometry_feeder_runner.query_attributes_in_batches",
                   return_value=start), \
             patch("meter_photo_qa.connectivity_trace.trace_connected_lv_lines",
                   side_effect=AssertionError("NO TOPOLOGY TRACE ALLOWED")), \
             patch("meter_photo_qa.lv_feeder_report.trace_connected_lv_lines",
                   side_effect=AssertionError("NO TOPOLOGY TRACE ALLOWED")):
            result=run_lv_feeder_report(folder,"Gicumbi",
                usernames=["ninamenye9683"],gis=object())
            self.assertEqual(result["mode"],"PROVISIONAL_GEOMETRY_NO_TRACE")
            self.assertEqual(result["starting_transformers"],1)
            self.assertEqual(result["nationwide_candidate_lv_lines"],3)
            self.assertEqual(result["connected_lv_line_ids"],2)
            self.assertEqual(result["unassociated_lines"],1)
            self.assertTrue(all(
                "geometry_filter" not in kw for kw in line_layer.calls
            ))
            self.assertTrue(all(
                "out_sr" not in kw or kw["out_sr"]==3857
                for kw in line_layer.calls
            ))
            wb=load_workbook(result["output_xlsx"],read_only=True)
            connected=list(wb["Provisional Connected Lines"].values)
            self.assertEqual(len(connected),3)
            self.assertEqual({connected[1][1],connected[2][1]},{1,2})
            self.assertEqual(
                wb["Other Surveyor Lines"].max_row,2,
            )
            self.assertIn("NO UTILITY NETWORK TRACE",wb["Summary"]["B2"].value)
            self.assertIn("Suggested Endpoint Gaps",wb.sheetnames)
            wb.close()

    def test_unbounded_all_surveyors_is_limited_instead_of_partial(self):
        class HugeLayer:
            def query(self,**opts):
                return 40000 if opts.get("return_count_only") else {"objectIds":[]}
        with self.assertRaisesRegex(RuntimeError,"Select a surveyor"):
            _get_nationwide_line_ids(HugeLayer(),None,"sde.DEFAULT",25000)

    def test_fail_when_service_geometry_not_properly_projected(self):
        from meter_photo_qa.geometry_feeder_runner import _ensure_projected_geometry
        with self.assertRaisesRegex(RuntimeError,"metre tolerances"):
            _ensure_projected_geometry(
                {"paths":[[[30,1],[31,1]]],"spatialReference":{"wkid":4326}},
                SimpleNamespace(spatial_reference={"wkid":4326}),
                "LV line",1,
            )

if __name__=="__main__":
    unittest.main()
