from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch


SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(SRC))

from meter_photo_qa.progress_report import (
    _build_connected_transformer_groups,
    _build_line_where,
    _create_workbook,
    _normalize_transformer_name,
    _query_lines,
    _query_local_transformers,
    _query_line_geometries,
    build_report_title,
    generate_progress_report,
    list_cells,
    list_districts,
    list_sectors,
    load_boundary_index,
    load_cell_index,
)


class ProgressReportTests(unittest.TestCase):
    def test_boundary_index_covers_the_whole_country(self):
        index = load_boundary_index()
        self.assertEqual(len(index), 30)
        self.assertEqual(sum(len(sectors) for sectors in index.values()), 416)
        self.assertEqual(len(index["Gasabo"]), 15)
        self.assertIn("Remera", index["Gasabo"])
        self.assertIn("Bweyeye", index["Rusizi"])

        cell_index = load_cell_index()
        self.assertEqual(len(cell_index), 30)
        self.assertEqual(sum(len(sectors) for sectors in cell_index.values()), 416)
        self.assertEqual(
            sum(
                len(cells)
                for sectors in cell_index.values()
                for cells in sectors.values()
            ),
            2148,
        )
        self.assertEqual(
            cell_index["Gasabo"]["Remera"],
            ["Nyabisindu", "Nyarutarama", "Rukiri I", "Rukiri Ii"],
        )

    def test_boundary_selectors_have_local_fallback(self):
        from unittest.mock import patch

        with patch(
            "meter_photo_qa.progress_report._get_gis",
            side_effect=RuntimeError("portal unavailable"),
        ):
            self.assertEqual(len(list_districts()), 30)
            self.assertEqual(len(list_sectors("Gasabo")), 15)
            self.assertIn("Rutunga", list_sectors("gasabo"))
            self.assertEqual(
                list_cells("Gasabo", ["Remera"]),
                ["Nyabisindu", "Nyarutarama", "Rukiri I", "Rukiri Ii"],
            )

    def test_line_where_uses_inclusive_end_date(self):
        where = _build_line_where(
            date(2026, 6, 26),
            date(2026, 7, 10),
            ["alpha", "beta"],
        )
        self.assertIn("ASSETGROUP IN (9,10)", where)
        self.assertIn("created_date >= TIMESTAMP '2026-06-26 00:00:00'", where)
        self.assertIn("created_date < TIMESTAMP '2026-07-11 00:00:00'", where)
        self.assertIn("created_user IN ('alpha','beta')", where)

    def test_title_matches_sample_structure(self):
        title = build_report_title(
            "2026-06-26",
            "2026-07-10",
            "Gasabo",
            ["Remera", "Rusororo", "Rutunga"],
        )
        self.assertEqual(
            title,
            "Progress Report GASABO DISTRICT: Remera, Rusororo and Rutunga "
            "Sectors (From 26th June to 10th July 2026)",
        )

    def test_title_handles_one_sector(self):
        title = build_report_title(
            "2026-07-01", "2026-07-15", "Gasabo", ["Remera"]
        )
        self.assertIn("Remera Sector", title)

    def test_title_includes_selected_cells(self):
        title = build_report_title(
            "2026-07-01",
            "2026-07-15",
            "Gasabo",
            ["Remera"],
            ["Nyabisindu", "Nyarutarama"],
        )
        self.assertIn(
            "Remera Sector; Nyabisindu and Nyarutarama Cells",
            title,
        )

    def test_title_identifies_custom_boundary(self):
        title = build_report_title(
            "2026-07-01",
            "2026-07-15",
            "Gasabo",
            ["Remera"],
            ["Nyabisindu"],
            custom_boundary=True,
        )
        self.assertIn("Progress Report CUSTOM BOUNDARY", title)
        self.assertNotIn("GASABO", title)

    def test_custom_boundary_report_excludes_users_without_matching_lines(self):
        records = {
            1: {
                "OBJECTID": 1,
                "created_user": "inside.user",
                "Shape__Length": 1250.0,
                "Transformer_Name": "Transformer A",
            }
        }
        custom_geometry = {
            "rings": [[[500000, 4800000], [500100, 4800000], [500000, 4800100]]],
            "spatialReference": {"wkid": 4326},
        }
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, patch(
            "meter_photo_qa.progress_report._get_gis", return_value=object()
        ), patch(
            "meter_photo_qa.progress_report._feature_layers",
            return_value=(object(), object(), object()),
        ), patch(
            "meter_photo_qa.progress_report._transformer_layer",
            return_value=object(),
        ), patch(
            "meter_photo_qa.progress_report._query_local_transformers",
            return_value=[{"OBJECTID": 100, "Transfomer_name_": "Transformer A"}],
        ), patch(
            "meter_photo_qa.progress_report.trace_connected_lv_lines",
            return_value={100: {
                "line_objectids": {1},
                "connected_transformer_objectids": {100},
                "network_element_count": 4,
            }},
        ), patch(
            "meter_photo_qa.progress_report._query_lines",
            side_effect=[records, records, {1: {"OBJECTID": 1}}],
        ), patch(
            "meter_photo_qa.progress_report._resolve_full_names",
            return_value={"inside.user": "Inside User"},
        ):
            summary = generate_progress_report(
                Path(folder) / "custom_boundary.xlsx",
                "2026-07-01",
                "2026-07-15",
                ["inside.user", "outside.user"],
                district="Gasabo",
                sectors=["Remera"],
                custom_boundary_geometry=custom_geometry,
            )
        self.assertEqual(summary["boundary_mode"], "custom")
        self.assertEqual(summary["selected_users"], ["inside.user", "outside.user"])
        self.assertEqual(summary["reported_users"], ["inside.user"])
        self.assertEqual([item["username"] for item in summary["users"]], ["inside.user"])
        self.assertEqual(summary["localized_transformers"], 1)
        self.assertEqual(summary["boundary_crossing_records"], 0)

    def test_progress_line_query_rejects_incomplete_and_duplicate_source_records(self):
        class FakeLine:
            def __init__(self, response):
                self.response=response
                self.calls=[]
            def query(self, **kwargs):
                self.calls.append(kwargs)
                return self.response

        def row(oid):
            return SimpleNamespace(attributes={
                "OBJECTID":oid, "GLOBALID":"{SOURCE}",
                "created_user":"operator", "Shape__Length":100,
            },geometry=None)

        for response, expected in (
            (SimpleNamespace(features=[row(1)], exceeded_transfer_limit=True),
             "transfer limit"),
            (SimpleNamespace(features=[row(1),row(1)],
                             exceeded_transfer_limit=False),
             "returned twice"),
            (SimpleNamespace(features=[SimpleNamespace(
                attributes={"GLOBALID":"{SOURCE}"},geometry=None
             )],exceeded_transfer_limit=False), "missing source OBJECTID"),
            (SimpleNamespace(), "no feature list"),
        ):
            with self.subTest(error=expected):
                layer=FakeLine(response)
                with self.assertRaisesRegex(RuntimeError,expected):
                    _query_lines(layer, "1=1", "sde.DEFAULT", None)
                self.assertEqual(len(layer.calls),1)  # no extra REG requests
                self.assertFalse(layer.calls[0]["return_geometry"])
        layer=FakeLine(SimpleNamespace(
            features=[row(1),row(2)],exceeded_transfer_limit=False
        ))
        data=_query_lines(layer, "1=1", "sde.DEFAULT", None)
        self.assertEqual(set(data),{1,2})

    def test_progress_transformer_location_query_fails_closed(self):
        class FakeTransformer:
            def __init__(self, rows, limited=False):
                self.rows=rows
                self.limited=limited
            def query(self,**kwargs):
                return SimpleNamespace(
                    features=[SimpleNamespace(attributes=attrs) for attrs in self.rows],
                    exceeded_transfer_limit=self.limited,
                )
        geom={"rings":[[[30,-2],[30.1,-2],[30,-1.9]]],
              "spatialReference":{"wkid":4326}}
        with patch(
            "meter_photo_qa.progress_report._geometry_filter",
            return_value=object(),
        ):
            with self.assertRaisesRegex(RuntimeError,"transfer limit"):
                _query_local_transformers(
                    FakeTransformer([{"OBJECTID":100}],True),geom,"sde.DEFAULT")
            with self.assertRaisesRegex(RuntimeError,"repeated OBJECTID"):
                _query_local_transformers(
                    FakeTransformer([{"OBJECTID":100},{"OBJECTID":100}]),
                    geom,"sde.DEFAULT")
            found=_query_local_transformers(
                FakeTransformer([{"OBJECTID":100},{"OBJECTID":101}]),
                geom,"sde.DEFAULT")
            self.assertEqual([v["OBJECTID"] for v in found],[100,101])

    def test_progress_geometry_batches_must_return_all_requested_objectids(self):
        class FakeLines:
            def __init__(self, ids, missing_crs=False):
                self.ids=ids
                self.missing_crs=missing_crs
                self.calls=[]
            def query(self,**kwargs):
                self.calls.append(kwargs)
                return SimpleNamespace(
                    features=[
                        SimpleNamespace(
                            attributes={"OBJECTID":oid},
                            geometry={"paths":[[[30,-2],[30.1,-2]]],
                                      **({} if self.missing_crs else
                                         {"spatialReference":{"wkid":4326}})},
                        )
                        for oid in self.ids
                    ],
                    spatial_reference={} if self.missing_crs else {"wkid":4326},
                    exceeded_transfer_limit=False,
                )
        scenarios=[
            ([12], [12,13], "omitted"),
            ([12,12], [12], "repeated"),
            ([12,99], [12], "unexpected"),
            ([12], [12], None),
        ]
        for response_ids, request_ids, error in scenarios:
            with self.subTest(response=response_ids,required=request_ids):
                layer=FakeLines(response_ids)
                if error:
                    with self.assertRaisesRegex(RuntimeError,error):
                        _query_line_geometries(layer,request_ids,"sde.DEFAULT")
                else:
                    actual=_query_line_geometries(layer,request_ids,"sde.DEFAULT")
                    self.assertEqual(list(actual),[12])
                self.assertEqual(len(layer.calls),1)
        with self.assertRaisesRegex(RuntimeError,"no verified CRS"):
            _query_line_geometries(FakeLines([12],missing_crs=True),
                                   [12],"sde.DEFAULT")

    def test_transformer_name_normalization_is_conservative(self):
        self.assertEqual(
            _normalize_transformer_name("  Remera   2A "),
            "remera 2a",
        )

    def test_connectivity_keeps_transformers_separate_and_numbered(self):
        groups, by_key = _build_connected_transformer_groups(
            [
                {"OBJECTID": 100, "Transfomer_name_": "Same", "ASSETGROUP": 12, "ASSETTYPE": 1},
                {"OBJECTID": 200, "Transfomer_name_": "Same", "ASSETGROUP": 12, "ASSETTYPE": 1},
            ],
            {
                100: {"line_objectids": {1, 2}, "connected_transformer_objectids": {100}, "network_element_count": 8},
                200: {"line_objectids": {3}, "connected_transformer_objectids": {200}, "network_element_count": 4},
            },
        )
        self.assertEqual(len(groups), 2)
        self.assertEqual(groups[0]["transformer_no"], "T-001")
        self.assertEqual(groups[1]["transformer_no"], "T-002")
        self.assertEqual(by_key["100"]["trace_line_objectids"], {1, 2})
        self.assertTrue(groups[0]["trackable"])

    def test_multi_transformer_trace_candidates_are_reviewed_not_counted(self):
        from openpyxl import load_workbook

        inside = {"OBJECTID": 1, "created_user": "pilot.user", "Shape__Length": 100.0,
                  "Transformer_Name": "Mburamazi trf"}
        outside = {"OBJECTID": 2, "created_user": "pilot.user", "Shape__Length": 200.0,
                   "Transformer_Name": "Mburamazi trf"}
        custom_geometry = {
            "rings": [[[500000, 4800000], [500100, 4800000], [500000, 4800100]]],
            "spatialReference": {"wkid": 4326},
        }
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, patch(
            "meter_photo_qa.progress_report._get_gis", return_value=object()
        ), patch(
            "meter_photo_qa.progress_report._feature_layers",
            return_value=(object(), object(), object()),
        ), patch(
            "meter_photo_qa.progress_report._transformer_layer", return_value=object(),
        ), patch(
            "meter_photo_qa.progress_report._query_local_transformers",
            return_value=[{"OBJECTID": 100, "Transfomer_name_": "Mburamazi trf"}],
        ), patch(
            "meter_photo_qa.progress_report.trace_connected_lv_lines",
            return_value={100: {
                "line_objectids": {1, 2},
                "connected_transformer_objectids": {100, 200},
                "network_element_count": 7,
            }},
        ), patch(
            "meter_photo_qa.progress_report._query_lines",
            side_effect=[{1: inside}, {1: inside, 2: outside}, {1: {"OBJECTID": 1}}],
        ), patch(
            "meter_photo_qa.progress_report._resolve_full_names",
            return_value={"pilot.user": "Pilot User"},
        ):
            output = Path(folder) / "ambiguous.xlsx"
            summary = generate_progress_report(
                output, "2026-09-14", "2026-09-14", ["pilot.user"],
                custom_boundary_geometry=custom_geometry,
            )
            self.assertEqual(summary["matched_line_records"], 0)
            self.assertEqual(summary["ambiguous_connected_transformers"], 1)
            self.assertEqual(summary["localization_warnings"], 2)
            workbook = load_workbook(output)
            log = workbook["Transformer Tracking Log"]
            self.assertIn("reaches 2 transformer points", log["J8"].value)
            self.assertTrue(any(
                isinstance(cell.value, str) and "Unassigned LV Candidates" in cell.value
                for row in log for cell in row
            ))
            workbook.close()

    def test_workbook_adds_transformer_tracking_log_without_changing_main_sheet(self):
        from openpyxl import load_workbook

        rows = [{
            "full_name": "Example Linesman",
            "username": "example.user",
            "length_km": 2.5,
        }]
        register = [{
            "transformer_no": "T-001",
            "transformer_name": "Urwego",
            "objectids": "100",
            "surveyors": "Example Linesman (example.user)",
            "line_records": 2,
            "total_km": 2.5,
            "crossing_records": 1,
            "outside_records": 0,
            "exception_km": 0.4,
        }]
        boundary_log = [{
            "transformer_no": "T-001",
            "transformer_name": "Urwego",
            "objectids": "100",
            "full_name": "Example Linesman",
            "username": "example.user",
            "line_objectid": 55,
            "boundary_status": "CROSSES SELECTED BOUNDARY",
            "length_km": 0.4,
            "districts": "Gasabo",
            "sectors": "Remera; Kimihurura",
            "cells": "Nyabisindu; Kimihurura",
            "created_date": None,
        }]
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            output = Path(folder) / "tracking_test.xlsx"
            _create_workbook(
                output,
                "Progress Report GASABO DISTRICT",
                rows,
                transformer_register=register,
                boundary_log=boundary_log,
                localization_warnings=[],
                tracking_context={
                    "boundary_label": "Gasabo / Remera / Nyabisindu",
                    "surveyors_followed": 1,
                },
            )
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Sheet2", "Transformer Tracking Log"])
            self.assertEqual(workbook["Sheet2"]["D6"].value, "=SUM(D5:D5)")
            log = workbook["Transformer Tracking Log"]
            self.assertEqual(log["A6"].value, "Localized Transformer Register")
            self.assertEqual(log["A8"].value, "T-001")
            self.assertEqual(log["B4"].value, "=COUNTA(A8:A8)")
            self.assertIn("Boundary Continuation Detail", [cell.value for cell in log[10]])
            workbook.close()

    def test_workbook_matches_approved_template(self):
        from openpyxl import load_workbook

        rows = [
            {
                "full_name": "Example Linesman",
                "username": "example.user",
                "length_km": 15.2,
            },
            {
                "full_name": "Second Linesman",
                "username": "second.user",
                "length_km": 8.22,
            },
        ]
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            output = Path(folder) / "template_test.xlsx"
            _create_workbook(output, "Progress Report GASABO DISTRICT", rows)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Sheet2"])
            sheet = workbook["Sheet2"]
            self.assertEqual(
                {str(item) for item in sheet.merged_cells.ranges},
                {"A2:D2", "A7:C7"},
            )
            self.assertEqual(sheet["A2"].fill.fgColor.rgb, "00C6E0B4")
            self.assertEqual(sheet["A4"].fill.fgColor.rgb, "00B4C6E7")
            self.assertEqual(sheet["A5"].fill.fgColor.rgb, "00C6E0B4")
            self.assertEqual(sheet["A7"].fill.fgColor.rgb, "00FFD966")
            self.assertEqual(sheet["D7"].value, "=SUM(D5:D6)")
            self.assertEqual(sheet["D7"].number_format, "0.00")
            self.assertEqual(sheet["D5"].number_format, "General")
            self.assertEqual(sheet["D4"].value, "Surveyed LV Line Records(km)")
            self.assertEqual(sheet.column_dimensions["B"].width, 63.140625)
            self.assertIsNone(sheet.freeze_panes)
            self.assertIsNone(sheet.auto_filter.ref)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
