from __future__ import annotations

import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from meter_photo_qa.template_audit_runner import run_template_surveyor_audit


class GPResult:
    def __init__(self,v):self.v=v
    def getOutput(self,i):return self.v


class FakeArcPy:
    def __init__(self,count=3,geographic_count=1,raise_spatial=False):
        self.calls=[]
        self.count=count
        self.geographic_count=geographic_count
        self.raise_spatial=raise_spatial
        self.management=SimpleNamespace(
            MakeFeatureLayer=self.make,
            GetCount=self.get_count,
            SelectLayerByLocation=self.select,
            Delete=self.delete,
        )
        self.da=SimpleNamespace(SearchCursor=self.search)
        self.rows=[
            (101,"{ID-1}","ninamenye9683",13),
            (102,"{ID-1}","ninamenye9683",14),
            (103,"{ID-3}","unmapped_user",16),
        ]

    def make(self,url,name,where):
        self.calls.append(("make",url,name,where))
        return GPResult(name)

    def get_count(self,layer):
        self.calls.append(("count",layer))
        return GPResult(str(self.geographic_count if "boundary" in layer else self.count))

    def select(self,layer,relation,polygon,selection_type):
        self.calls.append(("select",layer,relation,polygon,selection_type))
        if self.raise_spatial:
            raise RuntimeError("Boundary FeatureServer 503")

    def delete(self,layer):
        self.calls.append(("delete",layer))

    def search(self,layer,fields):
        self.calls.append(("cursor",layer,tuple(fields)))
        position={"OBJECTID":0,"GLOBALID":1,"created_user":2,"Pole_Height":3}
        rows=[[row[position[f]] for f in fields] for row in self.rows]
        class Cursor:
            def __enter__(self):return iter(rows)
            def __exit__(self,*args):return False
        return Cursor()


class TemplateAuditWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.arcpy=FakeArcPy()
        mod=types.ModuleType("arcgis.features")
        mod.FeatureLayer=lambda url,gis:SimpleNamespace(
            properties=SimpleNamespace(
                fields=[
                    {"name":"OBJECTID","type":"esriFieldTypeOID"},
                    {"name":"GLOBALID","type":"esriFieldTypeGlobalID"},
                    {"name":"created_user","type":"esriFieldTypeString"},
                    {"name":"ASSETGROUP","type":"esriFieldTypeInteger"},
                    {"name":"ASSETTYPE","type":"esriFieldTypeInteger"},
                    {"name":"Pole_Height","alias":"Pole Height","type":"esriFieldTypeDouble"},
                ],
                types=[],
            )
        )
        self.feature_mod=mod

    def invoke(self,arcpy=None):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder,patch.dict(
            sys.modules,{
                "arcgis":types.ModuleType("arcgis"),
                "arcgis.features":self.feature_mod,
            }
        ):
            output=Path(folder)/"audit.json"
            result=run_template_surveyor_audit(
                "Give me surveyors who recorded pole height greater than 12 in Gicumbi",
                output_folder=folder,output_json=output,arcpy_module=arcpy or self.arcpy,
                gis=object(),
            )
            self.assertEqual(result["matching_features"],3)
            self.assertEqual(result["matching_surveyors"],2)
            self.assertEqual(result["duplicate_globalid_groups"],1)
            self.assertEqual(result["notifications_sent"],0)
            self.assertTrue(Path(result["output_xlsx"]).exists())
            self.assertTrue(output.exists())
            parsed=json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(parsed["interpreted"]["district"],"Gicumbi")
            from openpyxl import load_workbook
            wb=load_workbook(result["output_xlsx"],read_only=True)
            self.assertEqual(wb["Surveyors"].max_row,3)
            self.assertEqual(wb["Matching Records"].max_row,4)
            self.assertEqual(wb["Notification Readiness"]["E2"].value,"NOT_SENT")
            self.assertEqual(wb["Surveyors"]["C2"].value,2)
            self.assertEqual(wb["Matching Records"]["H2"].value,"YES")
            self.assertIn("Resolved Template Criteria",wb.sheetnames)
            wb.close()
            return result

    def test_read_only_grouped_and_spatial_filtered(self):
        self.invoke()
        self.assertEqual(sum(x[0]=="select" for x in self.arcpy.calls),1)
        self.assertEqual(sum(x[0]=="cursor" for x in self.arcpy.calls),1)
        self.assertEqual(sum(x[0]=="delete" for x in self.arcpy.calls),2)
        self.assertTrue(all(x[0] in ("make","count","select","cursor","delete")
                            for x in self.arcpy.calls))
        self.assertTrue(any("ASSETTYPE IN (1,2,3)" in x[3]
                            for x in self.arcpy.calls if x[0]=="make"))

    def test_empty_boundary_aborts_before_cursor(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder,patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),
                         "arcgis.features":self.feature_mod}):
            fake=FakeArcPy(geographic_count=0)
            with self.assertRaisesRegex(RuntimeError,"zero matching polygons"):
                run_template_surveyor_audit(
                    "Show surveyors with pole height >12 in Gicumbi",
                    folder,arcpy_module=fake,gis=object()
                )
            self.assertFalse(any(x[0]=="cursor" for x in fake.calls))
            self.assertFalse(list(Path(folder).glob("*.xlsx")))

    def test_service_count_mismatch_does_not_publish_partial_report(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder,patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),
                         "arcgis.features":self.feature_mod}):
            fake=FakeArcPy(count=4)
            with self.assertRaisesRegex(RuntimeError,"Feature count changed"):
                run_template_surveyor_audit(
                    "Show surveyors with pole height >12 in Gicumbi",
                    folder,arcpy_module=fake,gis=object()
                )
            self.assertFalse(list(Path(folder).glob("*.xlsx")))

    def test_maximum_features_fails_closed(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder,patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),
                         "arcgis.features":self.feature_mod}):
            fake=FakeArcPy(count=30000)
            with self.assertRaisesRegex(ValueError,"audit maximum"):
                run_template_surveyor_audit(
                    "Show surveyors with pole height >12 in Gicumbi",
                    folder,arcpy_module=fake,gis=object()
                )
            self.assertFalse(any(x[0]=="cursor" for x in fake.calls))
            self.assertFalse(list(Path(folder).glob("*.xlsx")))

    def test_boundary_failure_aborts_before_reporting(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder,patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),
                         "arcgis.features":self.feature_mod}):
            fake=FakeArcPy(raise_spatial=True)
            with self.assertRaisesRegex(RuntimeError,"503"):
                run_template_surveyor_audit(
                    "Show surveyors with pole height >12 in Gicumbi",
                    folder,arcpy_module=fake,gis=object()
                )
            self.assertFalse(list(Path(folder).glob("*.xlsx")))

if __name__=="__main__":
    unittest.main()
