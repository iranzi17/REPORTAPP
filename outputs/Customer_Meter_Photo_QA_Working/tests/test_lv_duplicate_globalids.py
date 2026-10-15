"""Duplicate source GlobalIDs must not abort LV QA or merge review rows."""
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from meter_photo_qa.quality_watcher import run_quality_watcher
from test_quality_watcher import typical_attributes
from test_quality_watcher_workflow import fake_line_layer

class DuplicateLinesTests(unittest.TestCase):
    def test_duplicate_globalids_are_reported_and_reviewed_by_objectid(self):
        rows=[]
        for oid in (916046,922050):
            attrs=typical_attributes()
            attrs.update({
                "OBJECTID":oid,
                "GLOBALID":"{720A4222-DBF5-4562-93CA-6E5D7747044A}" if oid==916046 else "720a4222-dbf5-4562-93ca-6e5d7747044a",
                "created_user":"julesmugae72",
                "created_date":1789420000000,
                "last_edited_user":"julesmugae72",
                "last_edited_date":1789420000000,
                "ASSETTYPE":1,
                "Notes":"",
                "Transformer_Name":None,
            })
            rows.append(attrs)
        layer=fake_line_layer(rows[0])
        def query(**kwargs):
            if kwargs.get("return_ids_only"):
                return {"objectIds":[916046,922050]}
            ids={int(x) for x in str(kwargs.get("object_ids","")).split(",") if x}
            return SimpleNamespace(features=[
                SimpleNamespace(attributes=dict(row)) for row in rows
                if row["OBJECTID"] in ids
            ],exceeded_transfer_limit=False)
        layer.query=query
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch("meter_photo_qa.quality_watcher._get_gis",return_value=object()), \
             patch("meter_photo_qa.quality_watcher._feature_layers",return_value=(layer,None,None)), \
             patch("meter_photo_qa.quality_watcher._resolve_full_names",return_value={"julesmugae72":"Jules"}):
            output=Path(folder)/"LV_QA.xlsx"
            first=run_quality_watcher(output,"2026-09-11","2026-10-15")
            self.assertEqual(first["total_lines"],2)
            self.assertEqual(first["duplicate_globalid_groups"],1)
            self.assertEqual(first["duplicate_feature_records"],2)
            self.assertEqual(first["outstanding_lv_issues"],2)
            wb=load_workbook(first["output_xlsx"])
            self.assertEqual(wb["Duplicate GlobalIDs"]["B2"].value,"916046; 922050")
            ids=[r[1] for r in list(wb["Duplicate Line Records"].values)[1:]]
            self.assertEqual(ids,[916046,922050])
            queue=wb["Review Queue"]
            keys=[r[0] for r in list(queue.values)[4:]]
            self.assertEqual(len(set(keys)),2)
            self.assertTrue(all("|oid:" in key for key in keys))
            self.assertTrue(all("Duplicate GlobalID" in (r[7] or "")
                                for r in list(queue.values)[4:]))
            for row_no in range(5, queue.max_row + 1):
                oid = queue.cell(row_no, 3).value
                queue.cell(row_no, 13).value = (
                    "RESOLVED" if oid == 916046 else "IN_REVIEW"
                )
            wb.save(first["output_xlsx"])
            wb.close()
            second=run_quality_watcher(output,"2026-09-11","2026-10-15")
            wb=load_workbook(second["output_xlsx"],read_only=True)
            statuses={r[2]:r[12] for r in list(wb["Review Queue"].values)[4:]}
            self.assertEqual(statuses,{916046:"REOPENED",922050:"IN_REVIEW"})
            wb.close()
            # A duplicate is still reviewable when the mandatory fields are filled.
            for row in rows:
                row["Transformer_Name"]="Transformer A"
            third=run_quality_watcher(Path(folder)/"fresh.xlsx","2026-09-11","2026-10-15")
            wb=load_workbook(third["output_xlsx"],read_only=True)
            self.assertEqual([r[4] for r in list(wb["Review Queue"].values)[4:]],
                             ["IDENTIFIER_CONFLICT","IDENTIFIER_CONFLICT"])
            wb.close()

if __name__=="__main__":
    unittest.main()
