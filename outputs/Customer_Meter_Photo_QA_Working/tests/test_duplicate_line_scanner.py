from __future__ import annotations

import re
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from meter_photo_qa.duplicate_line_scanner import (
    collect_duplicate_ids, fetch_duplicate_features, assign_admin_areas,
    scan_electric_line_duplicate_globalids,
)

G1 = "{9FD1C5F4-AF23-43A3-90E3-CB1E7D8CA08C}"
G2 = "{11111111-2222-3333-4444-555555555555}"
G3 = "{AAAAAAAA-BBBB-CCCC-DDDD-EEEEEEEEEEEE}"


def feature(oid, gid, xs, group=9):
    return SimpleNamespace(
        attributes={
            "OBJECTID": oid, "GLOBALID": gid, "ASSETGROUP": group,
            "ASSETTYPE": 1, "SHAPE_Length": 19.028842,
            "created_user": "field.user", "Transformer_Name": "=HYPERLINK(0)",
        },
        geometry={
            "paths": [[[xs, 10], [xs + 1, 10]]],
            "spatialReference": {"wkid": 4326},
        },
    )


class FakeElectricLayer:
    def __init__(self, records):
        self.rows = sorted(records, key=lambda v: v.attributes["OBJECTID"])
        self.calls = []
        self.count_change = 0

    def query(self, **opts):
        self.calls.append(opts)
        eligible = [
            row for row in self.rows
            if "ASSETGROUP IN (9,10)" not in opts.get("where", "")
            or row.attributes["ASSETGROUP"] in (9,10)
        ]
        if opts.get("return_count_only"):
            return len(eligible) + self.count_change
        if opts.get("object_ids"):
            ids = {int(i) for i in opts["object_ids"].split(",")}
            return SimpleNamespace(
                features=[row for row in eligible if row.attributes["OBJECTID"] in ids],
                spatial_reference={"wkid": 4326},
            )
        greater = re.search(r"OBJECTID > (-?\d+)", opts["where"])
        assert greater, opts
        last = int(greater.group(1))
        limit = opts["result_record_count"]
        return SimpleNamespace(
            features=[
                SimpleNamespace(attributes=dict(row.attributes))
                for row in eligible if row.attributes["OBJECTID"] > last
            ][:limit],
            exceeded_transfer_limit=False,
        )


class FakeCells:
    def __init__(self):
        self.calls = []

    def query(self, **opts):
        self.calls.append(opts)
        geometry = opts["geometry_filter"]
        x = geometry["paths"][0][0][0]
        if x == 5:
            areas = [("Kicukiro","Masaka","Ayabaraya"),
                     ("Kicukiro","Masaka","Gako")]
        elif x == 8:
            areas = [("Gasabo","Bumbogo","Munyiginya")]
        else:
            areas = []
        return SimpleNamespace(features=[
            SimpleNamespace(attributes={"district":d,"sector":s,"cell":c})
            for d,s,c in areas
        ])


class ScannerTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            feature(122510,G1,5),
            feature(122944,G1,5),
            feature(200001,G2,8),
            feature(200002,G2,8),
            feature(200003,G3,30),
            feature(300000,"",12),
        ]

    def test_keyset_paging_finds_every_duplicate_even_across_pages(self):
        layer=FakeElectricLayer(self.rows)
        messages=[]
        groups,stats=collect_duplicate_ids(
            layer, "sde.DEFAULT", page_size=2, progress=messages.append,
        )
        self.assertEqual(groups,{G1:[122510,122944],G2:[200001,200002]})
        self.assertEqual(stats["scanned"],6)
        self.assertEqual(stats["invalid_or_missing_globalids"],1)
        pages=[c for c in layer.calls if c.get("result_record_count")]
        self.assertEqual(len(pages),4)
        self.assertTrue(all(c["order_by_fields"]=="OBJECTID ASC" for c in pages))
        self.assertTrue(any("duplicate groups" in x for x in messages))

    def test_duplicate_only_geometry_fetch_preserves_all_objectids(self):
        layer=FakeElectricLayer(self.rows)
        records=fetch_duplicate_features(layer,{G1:[122510,122944]}, "sde.DEFAULT")
        self.assertEqual(sorted(records),[122510,122944])
        geo=[c for c in layer.calls if c.get("object_ids")]
        self.assertEqual(len(geo),1)
        self.assertTrue(geo[0]["return_geometry"])
        self.assertNotIn("200001",geo[0]["object_ids"])

    def test_admin_areas_all_intersections_and_geometry_cache(self):
        cells=FakeCells()
        geometry=fetch_duplicate_features(
            FakeElectricLayer(self.rows),{G1:[122510,122944],G2:[200001,200002]},
            "sde.DEFAULT",
        )
        filters=types.ModuleType("arcgis.geometry.filters")
        filters.intersects=lambda geom,sr=None: geom
        geometry_module=types.ModuleType("arcgis.geometry")
        with patch.dict(sys.modules,{
            "arcgis":types.ModuleType("arcgis"),
            "arcgis.geometry":geometry_module,
            "arcgis.geometry.filters":filters,
        }):
            areas=assign_admin_areas(cells,geometry)
        self.assertEqual(len(cells.calls),2)
        self.assertEqual(areas[122510]["district"],["Kicukiro"])
        self.assertEqual(areas[122944]["cell"],["Ayabaraya","Gako"])
        self.assertEqual(areas[122510]["status"],"MULTIPLE_CELLS")
        self.assertEqual(areas[200001]["sector"],["Bumbogo"])

    def test_incomplete_server_scan_refuses_result(self):
        layer=FakeElectricLayer(self.rows)
        layer.count_change=1
        with self.assertRaisesRegex(RuntimeError,"initial"):
            collect_duplicate_ids(layer,"sde.DEFAULT",page_size=2)

    def test_different_features_same_length_do_not_prove_identical_geometry(self):
        rows=[feature(10,G1,1), feature(11,G1,4)]
        from meter_photo_qa.duplicate_line_inspector import _canonical_paths
        self.assertNotEqual(_canonical_paths(rows[0].geometry), _canonical_paths(rows[1].geometry))

    def test_end_to_end_national_all_line_scan_excel_has_geography(self):
        electric=FakeElectricLayer(self.rows)
        cells=FakeCells()
        feature_mod=types.ModuleType("arcgis.features")
        geometry_mod=types.ModuleType("arcgis.geometry")
        filters=types.ModuleType("arcgis.geometry.filters")
        filters.intersects=lambda geom,sr=None: geom
        feature_mod.FeatureLayer=lambda url,gis: (
            electric if url.endswith("/10") else cells if url.endswith("/1")
            else (_ for _ in ()).throw(AssertionError(url))
        )
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, patch.dict(
            sys.modules,{
                "arcgis":types.ModuleType("arcgis"),
                "arcgis.features":feature_mod,
                "arcgis.geometry":geometry_mod,
                "arcgis.geometry.filters":filters,
            }
        ):
            result=scan_electric_line_duplicate_globalids(
                Path(folder)/"duplicates.xlsx",gis=object(),page_size=2,
            )
            self.assertEqual(result["scanned"],6)
            self.assertEqual(result["duplicate_globalid_groups"],2)
            self.assertEqual(result["duplicate_feature_records"],4)
            self.assertEqual(result["additional_copies"],2)
            book=load_workbook(result["output_xlsx"],read_only=True)
            self.assertIn("All Duplicate Records",book.sheetnames)
            self.assertIn("Duplicate GlobalIDs",book.sheetnames)
            self.assertIn("By District Sector Cell",book.sheetnames)
            all_records=list(book["All Duplicate Records"].values)
            self.assertEqual(len(all_records),5)
            self.assertEqual({row[1] for row in all_records[1:]},
                             {122510,122944,200001,200002})
            masaka=[row for row in all_records[1:] if row[1]==122510][0]
            self.assertEqual(masaka[4],"Kicukiro")
            self.assertEqual(masaka[5],"Masaka")
            self.assertIn("Gako",masaka[6])
            self.assertIn("Ayabaraya",masaka[6])
            groups=list(book["Duplicate GlobalIDs"].values)
            self.assertEqual(len(groups),3)
            self.assertTrue(all(row[3]=="YES" for row in groups[1:]))
            survey=list(book["By District Sector Cell"].values)
            self.assertTrue(any(row[0]=="Kicukiro" and row[1]=="Masaka" for row in survey))
            book.close()

    def test_lv_only_scan_excludes_mv_duplicates(self):
        rows=self.rows+[feature(500001,G3,1,group=2), feature(500002,G3,1,group=2)]
        layer=FakeElectricLayer(rows)
        groups,stat=collect_duplicate_ids(
            layer,"sde.DEFAULT",where="ASSETGROUP IN (9,10)",page_size=2
        )
        self.assertEqual(stat["scanned"],6)
        self.assertNotIn(G3,groups)

    def test_zero_duplicates_can_still_produce_empty_register(self):
        groups,stat=collect_duplicate_ids(
            FakeElectricLayer([feature(1,G1,1)]),"sde.DEFAULT",page_size=2
        )
        self.assertEqual(groups,{})
        self.assertEqual(stat["scanned"],1)

    def test_admin_query_failure_is_not_silently_assigned(self):
        geometry=fetch_duplicate_features(
            FakeElectricLayer(self.rows),{G1:[122510,122944]},"sde.DEFAULT"
        )
        geometry_mod=types.ModuleType("arcgis.geometry")
        filters=types.ModuleType("arcgis.geometry.filters")
        filters.intersects=lambda geom,sr=None: geom
        class FailingCell:
            def query(self,**opts):
                raise RuntimeError("503 unavailable")
        with patch.dict(sys.modules, {
            "arcgis":types.ModuleType("arcgis"),
            "arcgis.geometry":geometry_mod,
            "arcgis.geometry.filters":filters,
        }):
            with self.assertRaisesRegex(RuntimeError,"503"):
                assign_admin_areas(FailingCell(),geometry)

if __name__=="__main__":
    unittest.main()
