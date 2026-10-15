from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook

from meter_photo_qa.combined_qa_report import run_combined_qa_report


def _fixture_workbook(path, kind):
    wb=Workbook()
    first=wb.active
    first.title="Linesman Summary" if kind=="lv_line" else "Summary"
    first["A1"]="Original "+kind+" report"
    first["A2"]="Source calculation"
    if kind=="lv_line":
        detail=wb.create_sheet("Line Details")
        detail.append(["Status","OBJECTID","GlobalID","Surveyor"])
        detail.append(["FLAG",11,"{LV1}","ninamenye9683"])
        first["B2"]='=COUNTA(\'Line Details\'!C2:C2)'
        wb.create_sheet("Outstanding LV Gaps").append(["GlobalID","Gap"])
        wb.create_sheet("Review Queue").append(["GlobalID","Review"])
        wb.create_sheet("Duplicate GlobalIDs").append(["GlobalID","OBJECTIDs"])
        wb.create_sheet("Duplicate Line Records").append(["GlobalID","OBJECTID"])
    else:
        detail=wb.create_sheet("Asset Details")
        detail.append(["Status","OBJECTID","GlobalID","Surveyor"])
        detail.append(["GAP",21 if kind=="transformer" else 31,
                       "{ASSET1}","ninamenye9683"])
        wb.create_sheet("Duplicate GlobalIDs").append(["GlobalID","OBJECTIDs"])
    wb.save(path)
    wb.close()


class CombinedQAWorkflowTests(unittest.TestCase):
    def fake_line(self,**kwargs):
        self.calls.append(("lv_line",kwargs))
        base=Path(kwargs["output_xlsx"])
        base.parent.mkdir(parents=True,exist_ok=True)
        file=base.with_name("original_lv.xlsx")
        _fixture_workbook(file,"lv_line")
        return {
            "output_xlsx":str(file),"total_lines":1,"flagged_lines":1,
            "outstanding_lv_lines":1,"outstanding_lv_issues":2,
            "duplicate_globalid_groups":1,
            "scored_field_gap_records":1,
            "scored_missing_entries":2,
            "scored_invalid_entries":1,
            "user_summary":[{"username":"ninamenye9683","total_lines":1,
                             "outstanding_lines":1}],
        }

    def fake_asset(self,**kwargs):
        key=kwargs["asset"]
        self.calls.append((key,kwargs))
        base=Path(kwargs["output_xlsx"])
        base.parent.mkdir(parents=True,exist_ok=True)
        file=base.with_name("original_"+key+".xlsx")
        _fixture_workbook(file,key)
        return {
            "output_xlsx":str(file),
            "total":1,"with_gaps":1,"issue_count":2,
            "original_attribute_gap_records":1,
            "original_attribute_issue_count":1,
            "duplicate_globalid_groups":1,
            "not_verified":0,
            "per_surveyor":[{"username":"ninamenye9683",
                             "total":1,"with_gaps":1}],
        }

    def setUp(self):
        self.calls=[]

    def test_all_three_original_qa_workbooks_combined_without_geometry(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch("meter_photo_qa.combined_qa_report.run_quality_watcher",
                   side_effect=self.fake_line), \
             patch("meter_photo_qa.combined_qa_report.run_asset_quality",
                   side_effect=self.fake_asset), \
             patch("meter_photo_qa.combined_qa_report.run_lv_feeder_report",
                   side_effect=AssertionError("geometry must not run")):
            r=run_combined_qa_report(
                folder,"Gicumbi",usernames=["ninamenye9683"],
                start_date="2026-09-14",end_date="2026-10-14",
                include_continuity=False,
            )
            self.assertEqual([k for k,_ in self.calls],
                             ["lv_line","transformer","pole"])
            for _,kw in self.calls:
                self.assertEqual(kw["district"],"Gicumbi")
                self.assertEqual(kw["usernames"],["ninamenye9683"])
                self.assertEqual(kw["start_date"].isoformat(),"2026-09-14")
                self.assertEqual(kw["end_date"].isoformat(),"2026-10-14")
            self.assertEqual(r["status"],"COMPLETE")
            wb=load_workbook(r["output_xlsx"])
            self.assertIn("LV Line Details",wb.sheetnames)
            self.assertIn("DT Asset Details",wb.sheetnames)
            self.assertIn("Pole Asset Details",wb.sheetnames)
            self.assertIn("LV Outstanding LV Gaps",wb.sheetnames)
            self.assertIn("LV Review Queue",wb.sheetnames)
            self.assertIn("LV Duplicate GlobalIDs",wb.sheetnames)
            self.assertIn("LV Duplicate Line Records",wb.sheetnames)
            self.assertTrue(any(
                row[0]=="LV line records with missing/invalid scored fields"
                and row[1]==1 for row in wb["Combined Summary"].values
            ))
            self.assertEqual(wb["Combined Summary"]["E13"].value,1)
            self.assertIn("DT Duplicate GlobalIDs",wb.sheetnames)
            self.assertIn("Pole Duplicate GlobalIDs",wb.sheetnames)
            self.assertEqual(wb["Combined Summary"]["B13"].value,1)
            self.assertEqual(wb["Combined Summary"]["B14"].value,1)
            self.assertEqual(wb["Combined Summary"]["B15"].value,1)
            self.assertEqual(wb["Surveyor Comparison"]["B2"].value,1)
            self.assertEqual(wb["Surveyor Comparison"]["D2"].value,1)
            self.assertEqual(wb["Surveyor Comparison"]["F2"].value,1)
            self.assertIn("'LV Line Details'!",
                          wb["LV Linesman Summary"]["B2"].value)
            wb.close()
            self.assertEqual(len(list(Path(r["components_folder"]).glob("*.xlsx"))),3)
            self.assertIn("COMPLETE",Path(r["manifest"]).read_text())

    def test_optional_continuity_failure_does_not_replace_qa_results(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch("meter_photo_qa.combined_qa_report.run_quality_watcher",
                   side_effect=self.fake_line), \
             patch("meter_photo_qa.combined_qa_report.run_asset_quality",
                   side_effect=self.fake_asset), \
             patch("meter_photo_qa.combined_qa_report.run_lv_feeder_report",
                   side_effect=RuntimeError("Geometry service timeout")):
            r=run_combined_qa_report(
                folder,"Gicumbi",usernames=["ninamenye9683"],
                include_continuity=True
            )
            self.assertEqual(r["geometry"]["status"],"UNAVAILABLE")
            wb=load_workbook(r["output_xlsx"],read_only=True)
            self.assertEqual(wb["Combined Summary"]["B13"].value,1)
            self.assertTrue(any(row[1]=="UNAVAILABLE"
                                for row in wb["Combined Summary"].values))
            self.assertTrue(any("Geometry service timeout" in str(row[1])
                                for row in wb["Combined Summary"].values))
            wb.close()

    def test_optional_continuity_is_appendix_not_a_substitute_for_qa(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch("meter_photo_qa.combined_qa_report.run_quality_watcher",
                   side_effect=self.fake_line), \
             patch("meter_photo_qa.combined_qa_report.run_asset_quality",
                   side_effect=self.fake_asset):
            def geometry(**kwargs):
                self.assertEqual(kwargs["district"],"Gicumbi")
                self.assertEqual(kwargs["usernames"],["ninamenye9683"])
                self.assertEqual(kwargs["method"],"geometry")
                path=Path(kwargs["output_folder"])/"geom.xlsx"
                _fixture_workbook(path,"geometry")
                return {
                    "output_xlsx":str(path),"starting_transformers":2,
                    "nationwide_candidate_lv_lines":50,
                    "connected_lv_line_ids":8,"unassociated_lines":42,
                    "suggested_gap_count":3,
                }
            with patch("meter_photo_qa.combined_qa_report.run_lv_feeder_report",
                       side_effect=geometry):
                r=run_combined_qa_report(
                    folder,"Gicumbi",usernames=["ninamenye9683"]
                )
            self.assertEqual(r["geometry"]["status"],"COMPLETE")
            wb=load_workbook(r["output_xlsx"],read_only=True)
            self.assertIn("Geo Summary",wb.sheetnames)
            self.assertEqual(wb["Combined Summary"]["B13"].value,1)
            self.assertTrue(any("PROVISIONAL" in str(row[1])
                                for row in wb["Combined Summary"].values))
            wb.close()

    def test_core_asset_failure_fails_closed_and_preserves_components(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, \
             patch("meter_photo_qa.combined_qa_report.run_quality_watcher",
                   side_effect=self.fake_line), \
             patch("meter_photo_qa.combined_qa_report.run_asset_quality",
                   side_effect=RuntimeError("Server 503")):
            with self.assertRaisesRegex(RuntimeError,"Consolidated QA failed"):
                run_combined_qa_report(
                    folder,"Gicumbi",usernames=["ninamenye9683"],
                    include_continuity=False,
                )
            dirs=list(Path(folder).glob("REG_Combined_QA_*"))
            self.assertEqual(len(dirs),1)
            self.assertIn("INCOMPLETE",(dirs[0]/"Run_Manifest.txt").read_text())
            self.assertEqual(len(list((dirs[0]/"Components").glob("*.xlsx"))),1)
            self.assertEqual(len(list(dirs[0].glob("*Combined*.xlsx"))),0)


if __name__=="__main__":
    unittest.main()
