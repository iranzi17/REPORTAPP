from __future__ import annotations

import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from meter_photo_qa.transformer_line_audit import run_transformer_line_name_audit


class FakeLayer:
    def __init__(self, rows):
        self.rows = list(rows)
        self.properties = SimpleNamespace(
            fields=[{"name":key,"alias":key}
                    for key in sorted({key for row in rows for key in row})],
            maxRecordCount=2000,
        )
        self.calls = []

    def query(self, **kwargs):
        self.calls.append(kwargs)
        if kwargs.get("return_ids_only"):
            return {"objectIds":[row["OBJECTID"] for row in self.rows]}
        if kwargs.get("object_ids"):
            ids = {int(x) for x in kwargs["object_ids"].split(",")}
            return SimpleNamespace(features=[
                SimpleNamespace(attributes=dict(row))
                for row in self.rows if row["OBJECTID"] in ids
            ])
        raise RuntimeError("Unexpected large query")


def li(oid, name):
    return {
        "OBJECTID": oid, "GLOBALID":f"{{L-{oid}}}",
        "Transformer_Name": name, "created_user":"isaac",
        "ASSETGROUP":9, "ASSETTYPE":1,
    }


def tr(oid, name):
    return {
        "OBJECTID": oid, "GLOBALID":f"{{T-{oid}}}",
        "ASSETGROUP":12, "ASSETTYPE":4, "Transfomer_name_":name,
    }


class TransformerLineNameWorkflowTests(unittest.TestCase):
    def test_cross_layer_report_from_single_traced_transformer(self):
        lines = FakeLayer([li(1,"Gatovu"),li(2,"Kisaro"),li(3,None)])
        points = FakeLayer([tr(101,"Kisaro")])
        def layer(url,gis):
            if url.endswith("/10"):
                return lines
            if url.endswith("/7"):
                return points
            return FakeLayer([])
        arcgis=types.ModuleType("arcgis")
        features=types.ModuleType("arcgis.features")
        features.FeatureLayer=layer
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, patch.dict(
            sys.modules, {"arcgis":arcgis,"arcgis.features":features},
        ), patch(
            "meter_photo_qa.transformer_line_audit.trace_connected_lv_lines",
            return_value={
                101:{
                    "line_objectids":{1,2,3},
                    "connected_transformer_objectids":{101},
                    "network_element_count":4,
                },
            },
        ) as trace:
            result=run_transformer_line_name_audit(
                Path(folder)/"name_consistency.xlsx",
                "2026-10-01", "2026-10-13", gis=object(),
            )
            self.assertEqual(result["name_mismatches_for_review"], 1)
            self.assertEqual(result["missing_line_transformer_names"], 1)
            self.assertEqual(result["matches"],1)
            self.assertEqual(result["lines_checked"],3)
            self.assertEqual(result["transformer_points_traced"],1)
            self.assertIn("created_date",lines.calls[0]["where"])
            self.assertIn("ASSETGROUP = 12",points.calls[0]["where"])
            self.assertNotIn("created_date",points.calls[0]["where"])
            self.assertEqual(trace.call_count,1)
            wb=load_workbook(result["output_xlsx"],read_only=True)
            self.assertIn("Name Mismatch Review",wb.sheetnames)
            self.assertIn("Surveyor Summary",wb.sheetnames)
            review=list(wb["Name Mismatch Review"].values)
            self.assertEqual(len(review),3) # header plus mismatch and blank name
            self.assertEqual({row[1] for row in review[1:]},{"{L-1}","{L-3}"})
            self.assertIn("{T-101}",[row[6] for row in review[1:]])
            wb.close()

    def test_trace_error_does_not_produce_partial_excel(self):
        lines=FakeLayer([li(1,"Wrong")])
        points=FakeLayer([tr(101,"Kisaro")])
        features=types.ModuleType("arcgis.features")
        features.FeatureLayer=lambda url,gis: lines if url.endswith("/10") else points
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder, patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),"arcgis.features":features},
        ), patch(
            "meter_photo_qa.transformer_line_audit.trace_connected_lv_lines",
            side_effect=RuntimeError("Dirty area / invalid connected trace"),
        ):
            with self.assertRaisesRegex(RuntimeError,"Dirty area"):
                run_transformer_line_name_audit(
                    Path(folder)/"audit.xlsx", "2026-10-01", "2026-10-13", gis=object(),
                )
            self.assertEqual(list(Path(folder).glob("*.xlsx")),[])


if __name__ == "__main__":
    unittest.main()
