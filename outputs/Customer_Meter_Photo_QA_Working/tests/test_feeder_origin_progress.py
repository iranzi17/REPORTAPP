"""Offline engineering acceptance tests for transformer-origin LV coverage.

No live GIS queries and no scheduled task changes.
"""
import ast
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from meter_photo_qa.feeder_inventory import (
    assess_source_feeder_inventory, add_feeder_inventory_sheets,
)
from meter_photo_qa.progress_report import _create_workbook


class TransformerOriginCoverageTests(unittest.TestCase):
    def setUp(self):
        self.groups = [
            {
                "key":"100","transformer_no":"T-001",
                "objectids":[100],"transformer_name":"Source A",
                "trace_line_objectids":{1,2,3},"trackable":True,
            },
            {
                "key":"200","transformer_no":"T-002",
                "objectids":[200],"transformer_name":"Source C",
                "trace_line_objectids":{2,4},"trackable":True,
            },
        ]
        self.lines = [
            {"OBJECTID":1,"GLOBALID":"{ONE}","ASSETGROUP":9,
             "Shape__Length":1000,"created_user":"alice","Transformer_Name":"Source A"},
            {"OBJECTID":2,"GLOBALID":"{TWO}","ASSETGROUP":10,
             "Shape__Length":500,"created_user":"alice","Transformer_Name":"Source A"},
            {"OBJECTID":3,"GLOBALID":"{THREE}","ASSETGROUP":9,
             "Shape__Length":2000,"created_user":"bob","Transformer_Name":"Source A"},
            {"OBJECTID":4,"GLOBALID":"{FOUR}","ASSETGROUP":9,
             "Shape__Length":400,"created_user":"bob","Transformer_Name":"Source C"},
        ]
        self.areas = {
            1: {"district":["Gasabo"],"sector":["Remera","Kacyiru"],"cell":["A","B"]},
            2: {"district":["Gasabo"],"sector":["Remera","Kacyiru"],"cell":["B"]},
            3: {"district":["Gasabo"],"sector":["Kacyiru"],"cell":["C"]},
            4: {"district":["Gasabo"],"sector":["Remera"],"cell":["A"]},
        }

    def assess(self, **kwargs):
        return assess_source_feeder_inventory(
            self.groups,self.lines,self.areas,"Gasabo",
            period_record_ids={1},selected_origin_sectors=["Remera"],**kwargs,
        )

    def test_origin_coverage_includes_outside_sector_but_no_double_count(self):
        rows, by_source, stats = self.assess()
        self.assertEqual([row["line_objectid"] for row in rows],[1,2,3,4])
        self.assertAlmostEqual(stats["all_trace_connected_length_km"],3.9)
        self.assertAlmostEqual(stats["unique_connected_source_length_km"],3.4)
        self.assertAlmostEqual(stats["ambiguous_or_unverified_source_length_km"],0.5)
        self.assertAlmostEqual(
            stats["unique_connected_source_length_km"]+
            stats["ambiguous_or_unverified_source_length_km"],
            stats["all_trace_connected_length_km"],
        )
        sources={row["source_transformer_objectid"]:row for row in by_source}
        self.assertAlmostEqual(sources[100]["unique_trace_length_km"],3.0)
        self.assertAlmostEqual(sources[200]["unique_trace_length_km"],0.4)
        self.assertEqual(rows[1]["source_status"],"AMBIGUOUS_MULTIPLE_TRANSFORMERS")
        self.assertIsNone(rows[1]["source_transformer_objectid"])
        self.assertEqual(rows[2]["sector_continuation_status"],"TOUCHES_OTHER_SECTOR")
        self.assertEqual(rows[2]["sectors"],"Kacyiru")
        self.assertEqual(stats["sector_continuation_segment_records"],3)
        self.assertAlmostEqual(
            stats["full_length_of_sector_continuation_segments_km"],3.5,
        )
        self.assertEqual(stats["surveyed_user_period_segment_records"],1)
        self.assertEqual(sum(int(row["in_requested_surveyor_period"]) for row in rows),1)

    def test_primary_sheet_is_origin_overview_preserves_approved_surveyor_sheet(self):
        rows, summaries, stats = self.assess()
        stats.update(district="Gasabo",version="sde.DEFAULT")
        with tempfile.TemporaryDirectory() as folder:
            output=Path(folder)/"origin.xlsx"
            _create_workbook(
                output,"Progress Report GASABO - Remera",
                [{"full_name":"Alice","username":"alice","length_km":1.0}],
                full_feeder_inventory={
                    "rows":rows,"summaries":summaries,"scope":stats,
                },
            )
            book=load_workbook(output)
            try:
                self.assertEqual(book.sheetnames[0],"Transformer-Origin Overview")
                self.assertEqual(book.sheetnames[1],"Sheet2")
                self.assertEqual(book["Sheet2"]["D6"].value,"=SUM(D5:D5)")
                info=dict(row[:2] for row in
                          book["Transformer-Origin Overview"].values
                          if len(row)>=2 and row[0] is not None)
                self.assertAlmostEqual(info["All traced LV length (km)"],3.9)
                self.assertAlmostEqual(info["Unassigned/multiple-source length (km)"],0.5)
                self.assertIn("NOT split",str(info["Coverage measurement"]))
                detailed=list(book["Full Feeder Segments"].values)
                header=detailed[0]
                self.assertIn("Sector Continuation Status",header)
                self.assertEqual(len(detailed),5)
                self.assertIn("Kacyiru",str(detailed[3]))
                caveats=list(book["Feeder Evidence Notes"].values)
                self.assertTrue(any("not a geographic allocation" in
                                    str(row[1]).lower() for row in caveats
                                    if row[1]))
            finally:
                book.close()

    def test_missing_nonfinite_length_fails_before_report(self):
        for value in (None,"",float("nan"),float("inf"),-10):
            with self.subTest(value=value):
                lines=[dict(row) for row in self.lines]
                lines[0]["Shape__Length"]=value
                with self.assertRaises(RuntimeError):
                    assess_source_feeder_inventory(
                        self.groups,lines,self.areas,"Gasabo",
                        selected_origin_sectors=["Remera"],
                    )

    def test_toolbox_exposes_full_origin_coverage_as_default_and_manual_only(self):
        code=(Path(__file__).resolve().parents[1]/"CustomerMeterPhotoQA.pyt").read_text(
            encoding="utf-8"
        )
        module=ast.parse(code)
        tool=next(
            node for node in module.body
            if isinstance(node,ast.ClassDef) and node.name=="GenerateSurveyProgressReport"
        )
        source=ast.get_source_segment(code,tool)
        self.assertIn('name="include_full_feeder_inventory"',source)
        self.assertIn('origin.value = True',source)
        self.assertIn("include_full_feeder_inventory=origin_enabled",source)
        self.assertIn("Transformer-Origin Coverage",source)
        self.assertNotIn("schtasks",source)


if __name__ == "__main__":
    unittest.main()
