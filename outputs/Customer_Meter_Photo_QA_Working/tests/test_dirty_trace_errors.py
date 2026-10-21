from __future__ import annotations

import json
import ast
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from types import SimpleNamespace
from unittest.mock import patch

from meter_photo_qa.connectivity_trace import (
    TopologyDirtyAreaError, _trace_dirty_error, trace_connected_lv_lines,
)


GUID = "{9FD1C5F4-AF23-43A3-90E3-CB1E7D8CA08C}"
POINT = {
    "OBJECTID": 74507, "GLOBALID": "{11111111-2222-3333-4444-555555555555}",
    "ASSETGROUP": 12, "ASSETTYPE": 4,
}
SCHEMA = {"domainNetworks": []}


class DirtyAreaTraceTests(unittest.TestCase):
    def test_arcgis_dirty_area_message_identifies_affected_features(self):
        error = Exception(f"One or more dirty areas were discovered. [ElectricLine: {GUID}] (Error Code: 0)")
        wrapped = _trace_dirty_error(error, POINT, "sde.DEFAULT", 2, 48)
        self.assertIsInstance(wrapped, TopologyDirtyAreaError)
        self.assertEqual(wrapped.transformer_objectid, 74507)
        self.assertEqual(wrapped.trace_position, 2)
        self.assertEqual(wrapped.trace_total, 48)
        self.assertEqual(wrapped.dirty_line_globalids, (GUID,))
        self.assertIn("NOT a transformer-name mismatch", str(wrapped))
        self.assertIn("No name-consistency", str(wrapped))
        self.assertIn(GUID, str(wrapped))

    def test_reg_multiple_dirty_electric_lines_all_preserved(self):
        other = "{7BFBCDFA-F565-4D6A-A75D-A0ACCE45CE5B}"
        first = "{14B864E9-87C5-4B20-9939-908F35C1B7F3}"
        server_message = (
            "One or more dirty areas were discovered. "
            f"[ElectricLine: {first}, {other}, {first}], Class not registered"
        )
        blocked = _trace_dirty_error(
            Exception(server_message), {
                **POINT,
                "OBJECTID":830384,
                "GLOBALID":"{35FA3DEC-9AC1-4605-9FC8-3A529D50BE20}",
            }, "sde.DEFAULT",21,22
        )
        self.assertIsInstance(blocked,TopologyDirtyAreaError)
        self.assertEqual(blocked.dirty_line_globalids,(first,other))
        self.assertIn(first,str(blocked))
        self.assertIn(other,str(blocked))
        self.assertEqual(blocked.trace_position,21)
        self.assertEqual(blocked.trace_total,22)

    def test_diagnostic_excel_is_not_a_progress_report_and_never_overwrites_it(self):
        from meter_photo_qa.progress_report import write_topology_blocker_diagnostic
        error = _trace_dirty_error(
            Exception(
                "One or more dirty areas were discovered. "
                "[ElectricLine: {14B864E9-87C5-4B20-9939-908F35C1B7F3}, "
                "{7BFBCDFA-F565-4D6A-A75D-A0ACCE45CE5B}], Class not registered"
            ),
            {**POINT,"OBJECTID":830384},
            "sde.DEFAULT",21,22,
        )
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder)/"survey_progress_report_20261021.xlsx"
            target.write_bytes(b"EXISTING REPORT MUST REMAIN")
            generated=write_topology_blocker_diagnostic(
                target,error,
                start_date="10/1/2026",
                end_date="10/20/2026",
                districts=["Musanze"],
                sectors=["Gataraga","Nyange","Shingiro"],
                cells=[],
                usernames=["janviersafari515","julesmugae72"],
            )
            self.assertEqual(target.read_bytes(),b"EXISTING REPORT MUST REMAIN")
            self.assertTrue(generated.exists())
            self.assertNotEqual(generated,target)
            self.assertIn("TRACE_BLOCKED",generated.name)
            book=load_workbook(generated,read_only=True)
            try:
                self.assertEqual(
                    book.sheetnames,
                    ["TRACE BLOCKED","Dirty Electric Lines","What this means"],
                )
                rows=list(book["TRACE BLOCKED"].values)
                data=dict(row[:2] for row in rows[1:])
                self.assertIn("NO — do not use",data["Report generated"])
                self.assertEqual(data["Trace position"],"21 of 22")
                self.assertEqual(data["Starting sector(s)"],
                                 "Gataraga; Nyange; Shingiro")
                self.assertEqual(data["Dirty ElectricLine GlobalIDs returned"],"2")
                offenders=list(book["Dirty Electric Lines"].values)
                self.assertEqual(len(offenders),3)
                self.assertEqual(offenders[1][1],error.dirty_line_globalids[0])
                self.assertEqual(offenders[2][1],error.dirty_line_globalids[1])
                self.assertEqual(offenders[1][2],830384)
                self.assertFalse(any("km total" in str(cell).lower()
                                     for row in rows for cell in row))
                self.assertIn("Class not registered",
                              data["Additional server message"])
            finally:
                book.close()

    def test_progress_tool_catches_dirty_error_only_and_still_fails(self):
        toolbox=Path(__file__).resolve().parents[1]/"CustomerMeterPhotoQA.pyt"
        source=toolbox.read_text(encoding="utf-8")
        tree=ast.parse(source)
        progress=next(
            item for item in tree.body
            if isinstance(item,ast.ClassDef)
            and item.name=="GenerateSurveyProgressReport"
        )
        execute=next(item for item in progress.body
                     if isinstance(item,ast.FunctionDef) and item.name=="execute")
        code=ast.get_source_segment(source,execute)
        self.assertIn("except _progress_report.TopologyDirtyAreaError as blocked:",code)
        self.assertIn("write_topology_blocker_diagnostic(",code)
        self.assertIn("TRACE BLOCKED DIAGNOSTIC (NOT A PROGRESS REPORT)",code)
        self.assertIn("            raise",code)
        self.assertNotIn("validateConsistency\": False",code)

    def test_non_dirty_errors_are_not_mislabeled(self):
        self.assertIsNone(_trace_dirty_error(
            Exception("HTTP 503 timeout"), POINT, "sde.DEFAULT", 1, 1
        ))

    def test_dirty_trace_preserves_validate_consistency_and_fails_closed(self):
        captured = {}
        def failing_post(url, params):
            captured["config"] = json.loads(params["traceConfiguration"])
            raise Exception(f"One or more dirty areas were discovered. [ElectricLine: {GUID}]")
        fake_gis = SimpleNamespace(_con=SimpleNamespace(post=failing_post))
        with patch("meter_photo_qa.connectivity_trace._network_schema", return_value=SCHEMA), \
             patch("meter_photo_qa.connectivity_trace._source", return_value={"sourceId": 3}), \
             patch("meter_photo_qa.connectivity_trace._terminal_id", return_value=2):
            with self.assertRaises(TopologyDirtyAreaError) as cm:
                trace_connected_lv_lines(
                    fake_gis, "https://example.org/FeatureServer",
                    [POINT], "sde.DEFAULT",
                )
        self.assertTrue(captured["config"]["validateConsistency"])
        self.assertEqual(cm.exception.dirty_line_globalids, (GUID,))

    def test_returned_service_error_dictionary_is_converted_to_actionable_error(self):
        fake_gis = SimpleNamespace(_con=SimpleNamespace(
            post=lambda url, params: {
                "error": {"message": f"One or more dirty areas were discovered. [ElectricLine: {GUID}]"}
            }
        ))
        with patch("meter_photo_qa.connectivity_trace._network_schema", return_value=SCHEMA), \
             patch("meter_photo_qa.connectivity_trace._source", return_value={"sourceId": 3}), \
             patch("meter_photo_qa.connectivity_trace._terminal_id", return_value=2):
            with self.assertRaises(TopologyDirtyAreaError):
                trace_connected_lv_lines(
                    fake_gis, "https://example.org/FeatureServer",
                    [POINT], "sde.DEFAULT",
                )


if __name__ == "__main__":
    unittest.main()
