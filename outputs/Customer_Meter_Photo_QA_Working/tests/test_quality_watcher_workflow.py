from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from meter_photo_qa.quality_watcher import run_quality_watcher
from test_quality_watcher import DOMAINS, typical_attributes


def fake_line_layer(attributes: dict):
    fields = {name for name in attributes}
    fields.update({"GLOBALID", "OBJECTID", "created_user", "created_date", "last_edited_user", "last_edited_date", "ASSETTYPE", "Notes"})
    service_types = []
    for group in (9, 10):
        service_types.append({
            "id": group,
            "domains": {
                name: {"codedValues": [{"code": code, "name": label} for code, label in values.items()]}
                for name, values in DOMAINS[9].items()
            },
        })
    layer = SimpleNamespace(properties={
        "fields": [{"name": name} for name in fields],
        "types": service_types,
    })
    layer.query = lambda **kwargs: SimpleNamespace(features=[SimpleNamespace(attributes=dict(attributes))])
    return layer


class QualityWatcherWorkflowTests(unittest.TestCase):
    def test_review_decisions_carry_forward_and_reopen_on_repeat_flag(self):
        attributes = typical_attributes()
        attributes.update({
            "OBJECTID": 101,
            "GLOBALID": "{EXAMPLE-GID}",
            "created_user": "alpha",
            "created_date": 1789420000000,
            "last_edited_user": "alpha",
            "last_edited_date": 1789420000000,
            "ASSETTYPE": 1,
            "Notes": "",
            "Transformer_Name": None,
        })
        line_layer = fake_line_layer(attributes)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            output = Path(folder) / "my_team_quality_report.xlsx"
            with patch("meter_photo_qa.quality_watcher._get_gis", return_value=object()), patch(
                "meter_photo_qa.quality_watcher._feature_layers", return_value=(line_layer, None, None)
            ), patch("meter_photo_qa.quality_watcher._resolve_full_names", return_value={"alpha": "Alpha"}):
                first = run_quality_watcher(output, "2026-09-14", "2026-09-15")
                self.assertEqual(first["flagged_lines"], 1)
                self.assertEqual(first["district"], "")
                self.assertEqual(first["sectors"], [])
                self.assertNotIn("created_user IS NOT NULL", first["where"])
                workbook = load_workbook(first["output_xlsx"])
                self.assertIn("Review Queue", workbook.sheetnames)
                self.assertIn("Outstanding Gap Summary", workbook.sheetnames)
                self.assertIn("Outstanding LV Gaps", workbook.sheetnames)
                self.assertEqual(first["outstanding_lv_lines"], 1)
                self.assertEqual(first["outstanding_lv_issues"], 1)
                gap_row = workbook["Outstanding LV Gaps"]
                self.assertEqual(gap_row["C2"].value, "{EXAMPLE-GID}")
                self.assertIn("Missing transformer name", gap_row["J2"].value)
                queue = workbook["Review Queue"]
                self.assertEqual(queue["E5"].value, "FLAG")
                queue["M5"] = "RESOLVED"
                queue["N5"] = "Team Lead"
                queue["Q5"] = "Corrected in source"
                workbook.save(first["output_xlsx"])
                workbook.close()

                attributes["Transformer_Name"] = "Transformer A"
                second = run_quality_watcher(output, "2026-09-14", "2026-09-15")
                self.assertNotEqual(first["output_xlsx"], second["output_xlsx"])
                self.assertEqual(second["met_threshold_lines"], 1)
                self.assertEqual(second["previous_review_report"], first["output_xlsx"])
                workbook = load_workbook(second["output_xlsx"], read_only=True)
                self.assertEqual(workbook["Review Queue"]["E5"].value, "MEETS_THRESHOLD")
                self.assertEqual(workbook["Review Queue"]["M5"].value, "RESOLVED")
                self.assertEqual(workbook["Review Queue"]["Q5"].value, "Corrected in source")
                workbook.close()

                attributes["Transformer_Name"] = None
                third = run_quality_watcher(output, "2026-09-14", "2026-09-15")
                workbook = load_workbook(third["output_xlsx"], read_only=True)
                self.assertEqual(workbook["Review Queue"]["M5"].value, "REOPENED")
                self.assertEqual(workbook["Review Queue"]["K5"].value, 2)
                self.assertEqual(workbook["Review Queue"]["N5"].value, "Team Lead")
                workbook.close()


if __name__ == "__main__":
    unittest.main()
