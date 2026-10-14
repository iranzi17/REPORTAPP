from __future__ import annotations
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from meter_photo_qa.lv_feeder_report import run_lv_feeder_report

class FakeLayer:
    def __init__(self,fields):
        self.properties=SimpleNamespace(fields=[{"name":x,"alias":x} for x in fields])

class FeederWorkflowTests(unittest.TestCase):
    def test_cross_district_trace_not_clipped_by_requested_surveyor(self):
        line_fields=[
            "OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE",
            "Transformer_Name","created_user","Phases",
        ]
        point_fields=["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE","Transfomer_name_"]
        points=[{"OBJECTID":101,"GLOBALID":"{T101}","ASSETGROUP":12,
                 "ASSETTYPE":4,"Transfomer_name_":"Kisaro"}]
        trace={101:{"line_objectids":{1,2,3},
                    "connected_transformer_objectids":{101},
                    "warnings":[]}}
        source={
            1:{"OBJECTID":1,"GLOBALID":"{L1}","ASSETGROUP":9,
               "Transformer_Name":"Kisaro","created_user":"ninamenye9683"},
            2:{"OBJECTID":2,"GLOBALID":"{L2}","ASSETGROUP":9,
               "Transformer_Name":"Kisaro","created_user":"someone_else"},
            3:{"OBJECTID":3,"GLOBALID":"{L3}","ASSETGROUP":9,
               "Transformer_Name":"Wrong","created_user":"ninamenye9683"},
        }
        passed={}
        def fetch(layer,oids,version,fields,usernames):
            passed["all_trace_ids"]=set(oids)
            passed["users"]=usernames
            return [source[oid] for oid in sorted(oids)
                    if not usernames or source[oid]["created_user"] in usernames]
        def query(layer,where,version,fields,filters,*args,**kwargs):
            passed["transformer_where"]=where
            passed["spatial_filters"]=filters
            self.assertNotIn("created_user",where)
            return points
        feature_module=types.ModuleType("arcgis.features")
        feature_module.FeatureLayer=lambda url,gis: (
            FakeLayer(line_fields) if url.endswith("/10")
            else FakeLayer(point_fields) if url.endswith("/7")
            else FakeLayer(["district","sector","cell"])
        )
        with tempfile.TemporaryDirectory() as folder,patch.dict(
            sys.modules,{"arcgis":types.ModuleType("arcgis"),
                         "arcgis.features":feature_module}
        ), patch(
            "meter_photo_qa.lv_feeder_report._get_sector_features",
            return_value=[object()],
        ), patch(
            "meter_photo_qa.lv_feeder_report._filters_for_area",
            return_value=[("Gicumbi",object())],
        ), patch(
            "meter_photo_qa.lv_feeder_report.query_attributes_in_batches",
            side_effect=query,
        ), patch(
            "meter_photo_qa.lv_feeder_report.trace_connected_lv_lines",
            return_value=trace,
        ), patch(
            "meter_photo_qa.lv_feeder_report._fetch_line_attributes_by_ids",
            side_effect=fetch,
        ):
            result=run_lv_feeder_report(folder,"Gicumbi",
                         usernames=["ninamenye9683"],gis=object(),method="topology")
            self.assertEqual(result["starting_transformers"],1)
            self.assertEqual(result["connected_lv_line_ids"],3)
            self.assertEqual(result["surveyed_connected_lines"],2)
            self.assertEqual(result["engineering_review_lines"],1)
            self.assertEqual(passed["all_trace_ids"],{1,2,3})
            self.assertEqual(passed["users"],["ninamenye9683"])
            wb=load_workbook(result["output_xlsx"],read_only=True)
            self.assertEqual(wb["Connected LV Lines"].max_row,3)
            self.assertIn("Engineering Review",wb.sheetnames)
            self.assertEqual(wb["Engineering Review"].max_row,2)
            self.assertEqual(wb["Connected LV Lines"]["B2"].value,"{L1}")
            self.assertEqual(wb["Connected LV Lines"]["B3"].value,"{L3}")
            self.assertIn("cross borders",result["boundary_rule"])
            wb.close()

    def test_empty_destination_fails_explicitly(self):
        with self.assertRaisesRegex(ValueError,"district is required"):
            run_lv_feeder_report("/tmp","")
if __name__=="__main__": unittest.main()
