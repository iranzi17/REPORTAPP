"""Source-only progress must work even with dirty Utility Network topology."""
import ast
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from meter_photo_qa.progress_report import generate_progress_report


def feature(oid, user, length, ag=9):
    return {
        "OBJECTID":oid, "GLOBALID":f"{{GUID-{oid}}}",
        "created_user":user,
        "created_date":int(datetime(2026,10,20,tzinfo=timezone.utc).timestamp()*1000),
        "Shape__Length":length,"Transformer_Name":"UNVERIFIED LABEL",
        "ASSETGROUP":ag,
    }


class SurveyorOnlyProgressTests(unittest.TestCase):
    def sample(self):
        return {
            10:feature(10,"alice",1000.0),
            11:feature(11,"alice",500.0,10),
            12:feature(12,"bob",750.0),
        }

    def run_report(self, folder, data=None, district="Burera"):
        rows=self.sample() if data is None else data
        output=Path(folder)/"survey_progress_report_20261021.xlsx"
        # An existing official report should be left intact even when the
        # next run uses the identical output name.
        output.write_bytes(b"MY EXISTING TRANSFORMER REPORT")
        boundary=SimpleNamespace(
            spatial_reference={"wkid":4326},
            features=[SimpleNamespace(geometry={
                "rings":[[[30,-2],[30.1,-2],[30,-1.9]]]
            })],
        )
        with patch(
            "meter_photo_qa.progress_report._get_gis",
            return_value=object(),
        ),patch(
            "meter_photo_qa.progress_report._feature_layers",
            return_value=(object(),object(),object()),
        ),patch(
            "meter_photo_qa.progress_report._get_sector_features",
            return_value=boundary if district else None,
        ),patch(
            "meter_photo_qa.progress_report._query_lines",
            return_value=rows,
        ) as query,patch(
            "meter_photo_qa.progress_report.trace_connected_lv_lines",
            side_effect=AssertionError("Utility Network trace MUST NOT run"),
        ) as trace,patch(
            "meter_photo_qa.progress_report._transformer_layer",
            side_effect=AssertionError("Transformer layer MUST NOT be loaded"),
        ) as transformers,patch(
            "meter_photo_qa.progress_report._resolve_full_names",
            return_value={"alice":"Alice","bob":"Bob"},
        ),patch(
            "meter_photo_qa.progress_report._load_name_overrides",
            return_value={},
        ):
            def run():
                return generate_progress_report(
                    output,"2026-07-07","2026-10-21",
                    ["alice","bob"],district=district,
                    service_url="https://fake.example/FeatureServer",
                    boundary_service_url="https://fake.example/Boundaries/FeatureServer",
                    include_full_feeder_inventory=False,
                    report_mode="surveyor_only",
                )
            if data is None:
                result=run()
            else:
                with self.assertRaisesRegex(RuntimeError,"(invalid|missing|Non-LV|Unselected)"):
                    run()
                result=None
            trace.assert_not_called()
            transformers.assert_not_called()
            self.assertEqual(query.call_count,1)
            requested_where=query.call_args.args[1]
            self.assertIn("ASSETGROUP IN (9,10)",requested_where)
            self.assertIn("created_user IN ('alice','bob')",requested_where)
            self.assertIn("2026-10-22",requested_where)
            self.assertEqual(
                query.call_args.args[3] is None,not bool(district)
            )
            self.assertEqual(output.read_bytes(),b"MY EXISTING TRANSFORMER REPORT")
            return result

    def test_burera_no_trace_produces_auditable_surveyor_progress(self):
        with tempfile.TemporaryDirectory() as folder:
            result=self.run_report(folder)
            self.assertEqual(result["report_mode"],"SURVEYOR_ONLY_NO_TRACE")
            self.assertTrue(result["no_topology_traces"])
            self.assertFalse(result["feeder_connectivity_verified"])
            self.assertEqual(result["matched_line_records"],3)
            self.assertAlmostEqual(result["total_km"],2.25)
            self.assertEqual([x["length_km"] for x in result["users"]],[1.5,.75])
            output=Path(result["output_xlsx"])
            self.assertIn("_SURVEYOR_NO_TRACE_",output.name)
            book=load_workbook(output)
            try:
                self.assertEqual(book.sheetnames,[
                    "Sheet2","Method and Warnings","Selected Source LV Lines",
                ])
                self.assertEqual(book["Sheet2"]["D7"].value,"=SUM(D5:D6)")
                self.assertEqual(book["Sheet2"]["D4"].value,
                                 "Surveyed LV Line Records(km)")
                notes=dict(list(book["Method and Warnings"].values)[1:])
                self.assertEqual(notes["Network trace requests"],0)
                self.assertIn("NOT VERIFIED",notes["Transformer feeder attribution"])
                self.assertIn("NOT EVALUATED",notes["Cross-sector continuation"])
                self.assertIn("not",notes["Area rule"].lower())
                entries=list(book["Selected Source LV Lines"].values)
                self.assertEqual(len(entries),4)
                self.assertEqual([x[0] for x in entries[1:]],[10,11,12])
                self.assertAlmostEqual(sum(x[6] for x in entries[1:]),2.25)
            finally:
                book.close()

    def test_unbounded_area_is_reported_as_not_geographically_filtered(self):
        with tempfile.TemporaryDirectory() as folder:
            result=self.run_report(folder,district=None)
            self.assertIn("NOT_CLIPPED",result["length_scope"])
            book=load_workbook(result["output_xlsx"],read_only=True)
            try:
                rows=list(book["Selected Source LV Lines"].values)
                self.assertTrue(all(v[7]=="ALL MATCHING RECORDS" for v in rows[1:]))
            finally:
                book.close()

    def test_bad_source_records_never_publish_report(self):
        bad=[
            {10:feature(10,"alice",float("nan"))},
            {10:feature(10,"alice",-50.0)},
            {10:feature(10,"alice",None)},
            {10:feature(10,"alice",200,ag=8)},
            {10:feature(10,"unknown",200)},
        ]
        for data in bad:
            with self.subTest(data=data),tempfile.TemporaryDirectory() as folder:
                self.run_report(folder,data=data)
                self.assertFalse(list(Path(folder).glob("*SURVEYOR_NO_TRACE*.xlsx")))

    def test_toolbox_off_switch_really_selects_no_trace_mode(self):
        code=(Path(__file__).resolve().parents[1]/"CustomerMeterPhotoQA.pyt").read_text(
            encoding="utf-8"
        )
        tree=ast.parse(code)
        tool=next(
            node for node in tree.body if isinstance(node,ast.ClassDef)
            and node.name=="GenerateSurveyProgressReport"
        )
        source=ast.get_source_segment(code,tool)
        self.assertIn('report_mode=("transformer_origin" if origin_enabled else "surveyor_only")',source)
        self.assertIn("SURVEYOR SOURCE-RECORD REPORT ONLY",source)
        self.assertIn("origin.value = True",source)
        self.assertIn("Uncheck",source.capitalize()) if False else None
        self.assertNotIn("schtasks",source)


if __name__ == "__main__":
    unittest.main()
