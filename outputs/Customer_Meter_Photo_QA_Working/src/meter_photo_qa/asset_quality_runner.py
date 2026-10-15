"""Read-only district-aware Quality Watcher for distribution transformers and LV poles."""
from __future__ import annotations

import os
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Sequence

from .asset_quality import (
    ASSETS, domain_map, evaluate_asset, resolve_fields, summarize_assets,
    annotate_duplicate_asset_globalids,
)
from .area_selection import selected_districts
from .progress_report import (
    DEFAULT_BOUNDARY_SERVICE_URL, DEFAULT_LINE_SERVICE_URL,
    _as_date, _get_gis, _get_sector_features, _query_with_retry, _layer_url,
    _sql_text,
)
from .quality_query_batches import query_attributes_in_batches


def _query_where(start, end, asset_group, usernames):
    end_next = end + timedelta(days=1)
    start_str = start.strftime("%Y-%m-%d 00:00:00")
    next_str = end_next.strftime("%Y-%m-%d 00:00:00")
    clauses = [
        f"ASSETGROUP = {int(asset_group)}",
        "((" +
        f"created_date >= TIMESTAMP '{start_str}' AND created_date < TIMESTAMP '{next_str}'"
        ") OR (" +
        f"last_edited_date >= TIMESTAMP '{start_str}' AND last_edited_date < TIMESTAMP '{next_str}'"
        "))",
    ]
    selected = [str(u).strip() for u in usernames or [] if str(u).strip()]
    if selected:
        quoted = ",".join(f"'{_sql_text(u)}'" for u in selected)
        clauses.append(f"created_user IN ({quoted})")
    return " AND ".join(clauses)


def _filters_for_area(boundary_features, custom_boundary_geometry):
    from arcgis.geometry.filters import intersects
    if custom_boundary_geometry is not None:
        return [(
            "Custom boundary",
            intersects(custom_boundary_geometry,
                       sr=custom_boundary_geometry.get("spatialReference", {})),
        )]
    if boundary_features is None:
        return [("All Rwanda", None)]
    result = []
    sr = boundary_features.spatial_reference
    for ix, feature in enumerate(boundary_features.features, 1):
        geom = feature.geometry
        if not geom or not geom.get("rings"):
            raise RuntimeError(f"Selected boundary polygon {ix} has invalid geometry.")
        attrs = dict(feature.attributes or {})
        desc = " / ".join(str(attrs.get(k) or "").strip()
                          for k in ("district", "sector", "cell")
                          if attrs.get(k)) or f"Boundary {ix}"
        result.append((desc, intersects(geom, sr=sr)))
    if not result:
        raise RuntimeError("No selected administrative boundaries were returned.")
    return result


def _report_xlsx(output_path, asset, rows, mapping, fields, summary, selected_areas, start, end, selected_users):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    overview = wb.active
    overview.title = "Summary"
    details = wb.create_sheet("Asset Details")
    surveyors = wb.create_sheet("Surveyor Gaps")
    schema = wb.create_sheet("Schema Verification")
    rules_sheet = wb.create_sheet("Rules")
    geo_sheet = wb.create_sheet("By Recorded Location")
    duplicate_groups_sheet = wb.create_sheet("Duplicate GlobalIDs")
    duplicate_records_sheet = wb.create_sheet("Duplicate Asset Records")

    overview.append([ASSETS[asset]["title"] + " - Field Data Quality"])
    overview.append(["Scope", "; ".join(selected_areas) if selected_areas else "All Rwanda"])
    overview.append(["From", start.isoformat()])
    overview.append(["To (inclusive)", end.isoformat()])
    overview.append(["Date basis", "Created or last edited"])
    overview.append(["Surveyor filter", "; ".join(selected_users) if selected_users else "ALL surveyors (no filter)"])
    overview.append(["Surveyor filter count", len(selected_users) if selected_users else "ALL"])
    overview.append(["Surveyor accounts returned", len({item["username"].casefold() for item in rows})])
    overview.append(["Feature Layer", ASSETS[asset]["layer_id"]])
    overview.append(["Asset Group", ASSETS[asset]["asset_group"]])
    for label, key in [
        ("Features examined", "total"), ("With outstanding gaps", "with_gaps"),
        ("No gaps in verified fields", "no_gaps_found"),
        ("Not verified (schema mismatch)", "not_verified"),
        ("Total individual issues (including duplicated identifiers)", "issue_count"),
        ("Original attribute-gap records (excluding duplicate identifiers)", "original_attribute_gap_records"),
        ("Original attribute issues (excluding duplicate identifiers)", "original_attribute_issue_count"),
        ("Duplicate GlobalID groups within this run", "duplicate_globalid_groups"),
        ("Feature records with reused GlobalIDs", "duplicate_feature_records"),
        ("Additional copies beyond one per GlobalID", "extra_records_with_reused_globalids"),
    ]:
        overview.append([label, summary[key]])
    overview.append(["Source", "REGUN FeatureServer (read-only)"])
    overview.append(["Interpretation", "Missing template fields are NOT checked; see Schema Verification."])
    if asset == "pole":
        overview.append(["Append Cable (mandatory) — missing/unknown",
                         sum("append_cable" in item["values"] and
                             item["values"].get("append_cable_status") == "MISSING_OR_UNKNOWN"
                             for item in rows)])
        overview.append(["Height > 12: MV/HV Notes required",
                         sum(item["values"].get("height_status") == "MV_HV_NOTES_REQUIRED"
                             for item in rows)])
        overview.append(["Height not recorded (informational, not an existing mandatory gap)",
                         sum(item["values"].get("height_status") == "UNRECORDED"
                             for item in rows)])
        overview.append(["Height check — source field not verified",
                         sum(item["values"].get("height_status") == "NOT_VERIFIED_FIELD"
                             for item in rows)])
        overview.append(["Append Cable — schema not verified",
                         sum(item["values"].get("append_cable_status") == "NOT_VERIFIED_FIELD"
                             for item in rows)])
        for key, label in (
            ("service_pole", "Service Pole"),
            ("structure", "Structure"),
            ("insulator_type", "Insulator Type"),
            ("number_of_circuit", "Number of Circuit"),
        ):
            overview.append([
                label + " — missing/unknown",
                sum(item["values"].get(key + "_status") == "MISSING_OR_UNKNOWN"
                    for item in rows),
            ])
            overview.append([
                label + " — field not verified",
                sum(item["values"].get(key + "_status") == "NOT_VERIFIED_FIELD"
                    for item in rows),
            ])
        overview.append(["Number of Circuit — invalid count",
                         sum(item["values"].get("number_of_circuit_status") == "INVALID_COUNT"
                             for item in rows)])
    overview.append(["Location source", "Recorded District/Sector/Cell fields; NOT spatially verified."])
    overview.append(["Duplicate scan coverage", "ONLY features matching this run\u2019s districts, dates and surveyor filters; counterparts outside scope may be absent."])
    overview.append(["Duplicate review", "Distinct OBJECTIDs with reused GlobalID need source investigation, not automatic deletion."])
    overview.append(["Records lacking all three recorded location fields",
                     sum(not (item["recorded_district"] and item["recorded_sector"] and item["recorded_cell"])
                         for item in rows)])
    overview.append(["Exclusion", "Transformer group 12 only; other Electric Device groups, including plant equipment, are excluded."] if asset == "transformer"
                    else ["Exclusion", "Structure Junction LV Pole group 7 only."])

    details.append([
        "Status", "OBJECTID", "GlobalID", "Asset Type", "Surveyor Username",
        "Recorded District", "Recorded Sector", "Recorded Cell",
        "Outstanding Issues", "Unverified Checks", "Other OBJECTIDs sharing GlobalID",
        "Owner Name", "Type of Installation", "Phase Normal",
        "Transformer Protection Switch", "Transformer mounting support",
        "Transformer Name", "Number of Low Voltage feeders",
        "Manufacturer", "Serial No", "Rated kVA",
        "Manufacturer Evidence", "Notes",
        *([
            "Service Pole", "Service Pole QA",
            "Structure", "Structure QA",
            "Insulator Type", "Insulator Type QA",
            "Number of Circuit", "Number of Circuit QA",
            "Pole Height", "Height QA", "Append Cable", "Append Cable QA",
        ] if asset == "pole" else []),
    ])
    for item in rows:
        vals = item["values"]
        details.append([
            item["status"], item["objectid"], item["globalid"],
            item["asset_type_label"], item["username"],
            item["recorded_district"], item["recorded_sector"],
            item["recorded_cell"],
            "; ".join(item["issues"]), "; ".join(item["unverified_checks"]),
            "; ".join(map(str, item.get("duplicate_peer_objectids", []))),
            vals.get("owner_name", ""), vals.get("installation_type", ""),
            vals.get("phase_normal", ""), vals.get("protection_switch", ""),
            vals.get("mounting_support", ""), vals.get("name", ""),
            vals.get("lv_feeder_count", ""), vals.get("manufacturer", ""),
            vals.get("serial", ""), vals.get("capacity", ""),
            vals.get("manufacturer_status", ""), item["notes"],
            *([
                vals.get("service_pole", ""), vals.get("service_pole_status", ""),
                vals.get("structure", ""), vals.get("structure_status", ""),
                vals.get("insulator_type", ""), vals.get("insulator_type_status", ""),
                vals.get("number_of_circuit", ""), vals.get("number_of_circuit_status", ""),
                vals.get("height", ""), vals.get("height_status", ""),
                vals.get("append_cable", ""), vals.get("append_cable_status", ""),
            ] if asset == "pole" else []),
        ])

    # Duplicates must not hide the seven original field-correction issues.
    duplicate_groups_sheet.append([
        "GlobalID", "Distinct OBJECTIDs", "Number of records",
        "Recorded District(s)", "Recorded Sector(s)", "Recorded Cell(s)",
        "Surveyor(s)", "Transformer/Pole name(s)",
        "Geometry comparison", "Correction advice",
    ])
    duplicate_records_sheet.append([
        "GlobalID", "OBJECTID", "Other OBJECTIDs With Same GlobalID",
        "Number In Duplicate Group", "Recorded District", "Recorded Sector",
        "Recorded Cell", "Created By", "Asset Type", "Transformer Name",
        "Other Outstanding Issues", "Source Confirmation",
    ])
    group_members = defaultdict(list)
    for item in rows:
        if item.get("duplicate_peer_objectids"):
            key = item["globalid"].strip().strip("{}").casefold()
            group_members[key].append(item)
    for key, group in sorted(group_members.items()):
        unique_oids = sorted({int(item["objectid"]) for item in group})
        if len(unique_oids) < 2:
            raise RuntimeError("Duplicate group lost an OBJECTID; refusing an incomplete register.")
        first_gid = group[0]["globalid"]
        def unique_text(field):
            return "; ".join(sorted({str(item.get(field) or "").strip()
                                     for item in group if str(item.get(field) or "").strip()},
                                    key=str.casefold))
        names = "; ".join(sorted({
            str(item["values"].get("name") or "").strip()
            for item in group if str(item["values"].get("name") or "").strip()
        }, key=str.casefold))
        duplicate_groups_sheet.append([
            first_gid, "; ".join(map(str, unique_oids)), len(group),
            unique_text("recorded_district"), unique_text("recorded_sector"),
            unique_text("recorded_cell"), unique_text("username"), names,
            "NOT CHECKED (attributes only)",
            "Verify whether both physical assets exist, their geometry, network associations and attachments. Do not delete automatically.",
        ])
        for item in sorted(group, key=lambda x: int(x["objectid"])):
            duplicate_records_sheet.append([
                item["globalid"], item["objectid"],
                "; ".join(map(str,item["duplicate_peer_objectids"])), len(group),
                item["recorded_district"], item["recorded_sector"],
                item["recorded_cell"], item["username"], item["asset_type_label"],
                item["values"].get("name", ""),
                "; ".join(issue for issue in item["issues"]
                          if issue not in item.get("identifier_issues", [])),
                "Source OBJECTID; spatial geometry and physical duplication NOT independently checked",
            ])

    surveyors.append(["Username", "Checked", "With Gaps", "No Gaps Found", "Not Verified", "Issue Count"])
    grouped = defaultdict(list)
    for item in rows:
        grouped[item["username"]].append(item)
    for username, records in sorted(grouped.items(), key=lambda pair: pair[0].casefold()):
        counts = summarize_assets(records)
        surveyors.append([
            username, counts["total"], counts["with_gaps"], counts["no_gaps_found"],
            counts["not_verified"], counts["issue_count"],
        ])
    geo_sheet.append([
        "Recorded District", "Recorded Sector", "Recorded Cell",
        "Assets Checked", "Assets with Gaps", "Individual Issues", "Not Verified",
    ])
    location_groups = defaultdict(list)
    for item in rows:
        key = (
            item["recorded_district"] or "(not recorded)",
            item["recorded_sector"] or "(not recorded)",
            item["recorded_cell"] or "(not recorded)",
        )
        location_groups[key].append(item)
    for (district_name, sector_name, cell_name), members in sorted(
        location_groups.items(), key=lambda entry: tuple(part.casefold() for part in entry[0])
    ):
        summary_by_area = summarize_assets(members)
        geo_sheet.append([
            district_name, sector_name, cell_name, summary_by_area["total"],
            summary_by_area["with_gaps"], summary_by_area["issue_count"],
            summary_by_area["not_verified"],
        ])
    schema.append(["Rule / Field Purpose", "Resolved Live Field", "Status", "Candidate Names"])
    for key, candidates in ASSETS[asset]["checks"].items():
        schema.append([
            key, mapping.get(key) or "", "VERIFIED" if mapping.get(key) else "NOT FOUND - NOT CHECKED",
            "; ".join(candidates),
        ])
    schema.append(["NOTE", "", "Branch is never scored or required." if asset == "transformer"
                   else "Manufacturer is flagged only if the manufacturer AND Notes are empty.", ""])
    schema.append(["NOTE", "", "Review unrecognized live schema fields before judging results.", ""])
    schema.append(["Live Feature Service Field", "Alias", "Type", ""])
    for field in fields:
        schema.append([field.get("name"), field.get("alias"), field.get("type"), ""])

    rules_sheet.append(["Type", "Rule", "Interpretation"])
    rules_sheet.append(["Both", "GlobalID", "Missing GlobalID is a gap."])
    rules_sheet.append(["Both", "Asset Type", "Unknown/unsupported asset type is a gap."])
    rules_sheet.append(["Transformer", "Seven original correction checks", "Owner Name, Type of Installation, Phase Normal, Protection Switch, mounting support, Transformer Name and LV feeder count; only live-mapped fields are evaluated."])
    rules_sheet.append(["Transformer", "Supplementary fields", "Manufacturer, Serial and Rated kVA are displayed where available but NOT scored in the original correction register."])
    rules_sheet.append(["Transformer", "Branch", "Explicitly excluded from mandatory checking."])
    rules_sheet.append([
        "LV Pole", "Manufacturer", "Flag only when Manufacturer is missing, Notes is empty, "
        "and PoleStructure is NOT Substandard. The Substandard option exempts this "
        "manufacturer check even with blank Notes. All other pole rules still apply."
    ])
    rules_sheet.append([
        "LV Pole", "Nameplate evidence", "An explicit no-nameplate statement is marked "
        "DOCUMENTED_NO_NAMEPLATE; any other nonempty Notes is labeled NOTES_PRESENT. "
        "The verified Substandard structure option is labeled SUBSTANDARD_EXEMPT, "
        "without inventing a Manufacturer value."
    ])
    rules_sheet.append(["LV Pole", "Service Pole", "Service_Pole: mandatory selected value; No is valid."])
    rules_sheet.append(["LV Pole", "Structure", "PoleStructure: mandatory known value, not PylonStructure or TowerStructure."])
    rules_sheet.append(["LV Pole", "Insulator Type", "InsulatorTypes: mandatory known value."])
    rules_sheet.append(["LV Pole", "Number of Circuit", "Number_Circuit: mandatory positive whole number."])
    rules_sheet.append([
        "LV Pole", "Append Cable — mandatory",
        "A selected, known value is required for every eligible pole. "
        "An explicit Yes or No is valid; blank, null, N/A and Unknown are gaps. "
        "If no matching live field is found, the check is NOT VERIFIED, "
        "never presumed complete.",
    ])
    rules_sheet.append(["LV Pole", "Height > 12",
        "If recorded height exceeds 12 in the source field, Notes must affirm the pole "
        "is for MV lines or HV lines/towers. Unknown field mapping = NOT_VERIFIED; "
        "the threshold uses stored field units pending live schema verification."])
    rules_sheet.append(["Both", "NOT_VERIFIED", "Unavailable live fields are not silently passed."])
    rules_sheet.append(["Both", "Source data", "Read-only. No records are edited."])
    rules_sheet.append(["Both", "Recorded District/Sector/Cell", "Exported from feature attributes for correction coordination; not proof of actual location."])
    rules_sheet.append(["Both", "Duplicate GlobalID", "Every repeated GlobalID within this selected run is retained by OBJECTID and marked as a data-integrity gap."])
    rules_sheet.append(["Both", "Duplicates outside selected area/date/user", "Not detected unless both source OBJECTIDs occur in the current report; full inventory needs a separate unrestricted scan."])
    rules_sheet.append(["Both", "Duplicate geometry", "Not compared. Duplicate identifiers do not prove identical physical equipment."])
    rules_sheet.append(["Both", "Changes", "No GIS data changed; manual engineering review required."])

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 30
        for col in sheet.columns:
            letter = col[0].column_letter
            max_len = max(len(str(c.value or "")) for c in list(col)[:1500])
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 13), 55)
    for sheet in (details, surveyors, geo_sheet, duplicate_groups_sheet, duplicate_records_sheet):
        if sheet.max_row > 1:
            table = Table(
                displayName=("DistributionTransformerDetails" if sheet == details and asset == "transformer"
                             else "LVPoleDetails" if sheet == details
                             else "AssetSurveyorGaps" if sheet == surveyors
                             else "RecordedAssetLocations" if sheet == geo_sheet
                             else "DuplicateAssetGroups" if sheet == duplicate_groups_sheet
                             else "DuplicateAssetRecords"),
                ref=f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    # Surveyor notes, names and GIS attributes are external text. Prevent a
    # string such as '=HYPERLINK(...)' from becoming an executable Excel formula.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"

    wb.calculation.fullCalcOnLoad = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing quality report: {output_path}")
    wb.save(output_path)


def run_asset_quality(
    asset: str,
    output_xlsx: str | Path,
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    district: str | Sequence[str] | None = None,
    sectors: Sequence[str] | None = None,
    cells: Sequence[str] | None = None,
    usernames: Sequence[str] | None = None,
    custom_boundary_geometry: dict | None = None,
    service_url: str | None = None,
    boundary_service_url: str | None = None,
    version: str = "sde.DEFAULT",
    gis: Any | None = None,
) -> dict[str, Any]:
    """Run one strict asset-class audit against live REGUN metadata."""
    if asset not in ASSETS:
        raise ValueError(f"Unsupported asset audit: {asset}")
    start, end = _as_date(start_date), _as_date(end_date)
    if end < start:
        raise ValueError("End date cannot be earlier than start date.")
    output = Path(output_xlsx).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Report must be an .xlsx file.")
    now = datetime.now()
    output = output.with_name(f"{output.stem}_{now:%Y%m%d_%H%M%S}{output.suffix}")
    if output.exists():
        raise FileExistsError(f"Report already exists: {output}")

    active_gis = _get_gis(gis)
    service = service_url or os.environ.get("UN_REPORT_SERVICE_URL", DEFAULT_LINE_SERVICE_URL)
    boundary = boundary_service_url or os.environ.get("UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL)
    from arcgis.features import FeatureLayer
    from .progress_report import SECTOR_LAYER_ID, CELL_LAYER_ID
    layer = FeatureLayer(_layer_url(service, ASSETS[asset]["layer_id"]), active_gis)
    sector_layer = FeatureLayer(_layer_url(boundary, SECTOR_LAYER_ID), active_gis)
    cell_layer = FeatureLayer(_layer_url(boundary, CELL_LAYER_ID), active_gis)
    boundary_features = None
    if custom_boundary_geometry is None:
        boundary_features = _get_sector_features(
            sector_layer, district, sectors, cell_layer, cells
        )
    selections = _filters_for_area(boundary_features, custom_boundary_geometry)
    fields = [dict(item) for item in layer.properties.fields]
    mapping = resolve_fields(fields, asset)
    # Do not invent physical data-collection field names. Flag unverified
    # schema coverage clearly in both the report and ArcGIS messages.
    core = ["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE","created_user","created_date","last_edited_date"]
    lookup = {field["name"].casefold(): field["name"] for field in fields}
    for name in core:
        if name.casefold() not in lookup:
            raise RuntimeError(f"Live {ASSETS[asset]['title']} service lacks {name}.")
    names = list(dict.fromkeys([
        *[lookup[name.casefold()] for name in core],
        *[lookup[field.casefold()] for field in ("District", "Sector", "Cell")
          if field.casefold() in lookup],
        *[resolved for resolved in mapping.values() if resolved],
    ]))
    def progress(message):
        try:
            import arcpy
            arcpy.AddMessage(message)
        except ImportError:
            pass
    attributes = query_attributes_in_batches(
        layer, _query_where(start, end, ASSETS[asset]["asset_group"], usernames),
        version, names, selections, _query_with_retry,
        progress=progress, asset_label="Distribution transformer" if asset == "transformer" else "LV pole",
        duplicate_globalid_policy="report",
    )
    domains = domain_map(layer.properties, ASSETS[asset]["asset_group"])
    results = [
        evaluate_asset(item, asset, mapping, domains)
        for item in attributes
    ]
    duplicate_groups = annotate_duplicate_asset_globalids(results)
    if duplicate_groups:
        progress(
            f"Found {len(duplicate_groups)} repeated {asset} GlobalID groups, "
            "all distinct OBJECTIDs retained for QA. See Excel Duplicate GlobalIDs."
        )
    results.sort(key=lambda x: (x["status"] != "GAP", x["username"].casefold(), x["objectid"]))
    counts = summarize_assets(results)
    surveyor_rows = defaultdict(list)
    for row in results:
        surveyor_rows[row["username"]].append(row)
    per_surveyor = [
        {"username": username, **summarize_assets(records)}
        for username, records in sorted(
            surveyor_rows.items(), key=lambda item: item[0].casefold()
        )
    ]
    selected = selected_districts(district)
    selected_users = sorted({str(user).strip() for user in (usernames or [])
                             if str(user).strip()}, key=str.casefold)
    _report_xlsx(output, asset, results, mapping, fields, counts, selected, start, end,
                 selected_users)
    return {
        "read_only": True, "asset": asset, "output_xlsx": str(output),
        "districts": selected, "date_start": start.isoformat(),
        "date_end": end.isoformat(), **counts,
        "surveyor_filter": selected_users,
        "returned_surveyor_count": len({item["username"].casefold() for item in results}),
        "per_surveyor": per_surveyor,
        "incomplete_recorded_locations": sum(
            not (item["recorded_district"] and item["recorded_sector"] and item["recorded_cell"])
            for item in results
        ),
        "unverified_live_fields": [key for key, found in mapping.items() if not found],
        "source_layer": _layer_url(service, ASSETS[asset]["layer_id"]),
    }
