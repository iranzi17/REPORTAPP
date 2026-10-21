"""Read-only LV survey progress reporting for the REG utility network."""

from __future__ import annotations

import os
import time
import json
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

from .connectivity_trace import TopologyDirtyAreaError, trace_connected_lv_lines
from .area_selection import district_filters, selected_districts


DEFAULT_LINE_SERVICE_URL = (
    "https://gisweb.reg.rw/server/rest/services/"
    "REGUN_Service__Versio2/FeatureServer"
)
DEFAULT_BOUNDARY_SERVICE_URL = (
    "https://gisweb.reg.rw/server/rest/services/Hosted/"
    "Sector_and_Cell_Boundaries/FeatureServer"
)
LINE_LAYER_ID = 10
TRANSFORMER_LAYER_ID = 7
DISTRIBUTION_TRANSFORMER_ASSETGROUP = 12
TRANSFORMER_POINT_NAME_FIELD = "Transfomer_name_"
SECTOR_LAYER_ID = 0
CELL_LAYER_ID = 1
DEFAULT_USER_NAME_MAP_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "report_user_names.json"
)
DEFAULT_BOUNDARY_INDEX_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "rwanda_district_sectors.json"
)


def load_province_index(path: str | Path | None = None) -> dict[str, list[str]]:
    """Load the complete Rwanda province-to-district selector index."""
    candidate = Path(path).expanduser() if path else DEFAULT_BOUNDARY_INDEX_PATH
    if not candidate.exists():
        return {}
    with candidate.open("r", encoding="utf-8") as stream:
        data = json.load(stream)
    raw_provinces = data.get("provinces", {}) if isinstance(data, dict) else {}
    if not isinstance(raw_provinces, dict):
        raise ValueError("The Rwanda boundary index must contain a provinces object.")
    result: dict[str, list[str]] = {}
    for province, districts in raw_provinces.items():
        clean_province = str(province or "").strip()
        if not clean_province or not isinstance(districts, list):
            continue
        clean_districts = sorted(
            {
                str(district or "").strip()
                for district in districts
                if str(district or "").strip()
            },
            key=str.casefold,
        )
        if clean_districts:
            result[clean_province] = clean_districts
    return dict(sorted(result.items(), key=lambda item: item[0].casefold()))


def load_boundary_index(path: str | Path | None = None) -> dict[str, list[str]]:
    """Load the complete Rwanda district-to-sector selector index."""
    candidate = Path(path).expanduser() if path else DEFAULT_BOUNDARY_INDEX_PATH
    if not candidate.exists():
        return {}
    with candidate.open("r", encoding="utf-8") as stream:
        data = json.load(stream)
    raw_districts = data.get("districts", data) if isinstance(data, dict) else {}
    if not isinstance(raw_districts, dict):
        raise ValueError("The Rwanda boundary index must contain a districts object.")
    result: dict[str, list[str]] = {}
    for district, sectors in raw_districts.items():
        clean_district = str(district or "").strip()
        if not clean_district or not isinstance(sectors, list):
            continue
        clean_sectors = sorted(
            {str(sector or "").strip() for sector in sectors if str(sector or "").strip()},
            key=str.casefold,
        )
        if clean_sectors:
            result[clean_district] = clean_sectors
    return dict(sorted(result.items(), key=lambda item: item[0].casefold()))


def load_cell_index(
    path: str | Path | None = None,
) -> dict[str, dict[str, list[str]]]:
    """Load the complete Rwanda district-to-sector-to-cell selector index."""
    candidate = Path(path).expanduser() if path else DEFAULT_BOUNDARY_INDEX_PATH
    if not candidate.exists():
        return {}
    with candidate.open("r", encoding="utf-8") as stream:
        data = json.load(stream)
    raw_cells = data.get("cells", {}) if isinstance(data, dict) else {}
    if not isinstance(raw_cells, dict):
        raise ValueError("The Rwanda boundary index must contain a cells object.")

    result: dict[str, dict[str, list[str]]] = {}
    for district, sectors in raw_cells.items():
        clean_district = str(district or "").strip()
        if not clean_district or not isinstance(sectors, dict):
            continue
        clean_sector_map: dict[str, list[str]] = {}
        for sector, cells in sectors.items():
            clean_sector = str(sector or "").strip()
            if not clean_sector or not isinstance(cells, list):
                continue
            clean_cells = sorted(
                {str(cell or "").strip() for cell in cells if str(cell or "").strip()},
                key=str.casefold,
            )
            if clean_cells:
                clean_sector_map[clean_sector] = clean_cells
        if clean_sector_map:
            result[clean_district] = dict(
                sorted(clean_sector_map.items(), key=lambda item: item[0].casefold())
            )
    return dict(sorted(result.items(), key=lambda item: item[0].casefold()))


def _fallback_sectors(district: str, index: dict[str, list[str]]) -> list[str]:
    requested = str(district or "").strip().casefold()
    for name, sectors in index.items():
        if name.casefold() == requested:
            return list(sectors)
    return []


def _fallback_cells(
    district: str,
    sectors: Sequence[str],
    index: dict[str, dict[str, list[str]]],
) -> list[str]:
    requested_district = str(district or "").strip().casefold()
    requested_sectors = {str(item).strip().casefold() for item in sectors if str(item).strip()}
    for district_name, sector_map in index.items():
        if district_name.casefold() != requested_district:
            continue
        cells = {
            cell
            for sector_name, sector_cells in sector_map.items()
            if sector_name.casefold() in requested_sectors
            for cell in sector_cells
        }
        return sorted(cells, key=str.casefold)
    return []


def _as_date(value: date | datetime | str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError("A date value is required.")
    for parser in (
        date.fromisoformat,
        lambda item: datetime.strptime(item, "%m/%d/%Y").date(),
        lambda item: datetime.strptime(item, "%d/%m/%Y").date(),
    ):
        try:
            return parser(text)
        except ValueError:
            continue
    raise ValueError(f"Unsupported date value: {value}")


def _sql_text(value: str) -> str:
    return str(value).replace("'", "''")


def _layer_url(service_url: str, layer_id: int) -> str:
    base = service_url.rstrip("/")
    tail = base.rsplit("/", 1)[-1]
    if tail.isdigit():
        return base
    return f"{base}/{layer_id}"


def _timestamp_sql(value: date) -> str:
    return value.strftime("%Y-%m-%d 00:00:00")


def _build_line_where(
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    usernames: Sequence[str] | None = None,
    asset_group_ids: Sequence[int] = (9, 10),
) -> str:
    start = _as_date(start_date)
    end = _as_date(end_date)
    if end < start:
        raise ValueError("The end date cannot be earlier than the start date.")
    following_day = end + timedelta(days=1)
    groups = ",".join(str(int(item)) for item in asset_group_ids)
    parts = [
        f"ASSETGROUP IN ({groups})",
        f"created_date >= TIMESTAMP '{_timestamp_sql(start)}'",
        f"created_date < TIMESTAMP '{_timestamp_sql(following_day)}'",
        "created_user IS NOT NULL",
    ]
    clean_users = [str(item).strip() for item in usernames or [] if str(item).strip()]
    if clean_users:
        quoted = ",".join(f"'{_sql_text(item)}'" for item in clean_users)
        parts.append(f"created_user IN ({quoted})")
    return " AND ".join(parts)


def _query_with_retry(layer: Any, attempts: int = 3, **kwargs: Any) -> Any:
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            return layer.query(**kwargs)
        except Exception as exc:  # pragma: no cover - depends on the live service
            last_error = exc
            if attempt < attempts:
                time.sleep(attempt * 2)
    raise RuntimeError(f"ArcGIS query failed after {attempts} attempts: {last_error}")


def _get_gis(gis: Any | None = None) -> Any:
    if gis is not None:
        return gis
    from arcgis.gis import GIS

    return GIS("pro")


def _feature_layers(
    gis: Any,
    service_url: str,
    boundary_service_url: str,
) -> tuple[Any, Any, Any]:
    from arcgis.features import FeatureLayer

    line_layer = FeatureLayer(_layer_url(service_url, LINE_LAYER_ID), gis)
    sector_layer = FeatureLayer(
        _layer_url(boundary_service_url, SECTOR_LAYER_ID), gis
    )
    cell_layer = FeatureLayer(
        _layer_url(boundary_service_url, CELL_LAYER_ID), gis
    )
    return line_layer, sector_layer, cell_layer


def _transformer_layer(gis: Any, service_url: str) -> Any:
    from arcgis.features import FeatureLayer

    return FeatureLayer(_layer_url(service_url, TRANSFORMER_LAYER_ID), gis)


def _distinct_text(features: Iterable[Any], field_name: str) -> list[str]:
    values = {
        str(feature.attributes.get(field_name)).strip()
        for feature in features
        if feature.attributes.get(field_name) not in (None, "")
    }
    return sorted(values, key=str.casefold)


def list_districts(
    boundary_service_url: str | None = None,
    gis: Any | None = None,
) -> list[str]:
    """Return available district names from the live boundary layer."""
    fallback = load_boundary_index()
    live_districts: list[str] = []
    try:
        from arcgis.features import FeatureLayer

        active_gis = _get_gis(gis)
        url = boundary_service_url or os.environ.get(
            "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
        )
        layer = FeatureLayer(_layer_url(url, SECTOR_LAYER_ID), active_gis)
        feature_set = _query_with_retry(
            layer,
            where="district IS NOT NULL",
            out_fields="district",
            return_geometry=False,
            return_distinct_values=True,
            order_by_fields="district",
        )
        live_districts = _distinct_text(feature_set.features, "district")
    except Exception:  # pragma: no cover - depends on live portal availability
        pass
    return sorted(set(live_districts) | set(fallback), key=str.casefold)


def list_sectors(
    district: str,
    boundary_service_url: str | None = None,
    gis: Any | None = None,
) -> list[str]:
    """Return available sector names for one district."""
    if not str(district or "").strip():
        return []
    fallback = load_boundary_index()
    live_sectors: list[str] = []
    try:
        from arcgis.features import FeatureLayer

        active_gis = _get_gis(gis)
        url = boundary_service_url or os.environ.get(
            "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
        )
        layer = FeatureLayer(_layer_url(url, SECTOR_LAYER_ID), active_gis)
        where = f"UPPER(district) = '{_sql_text(district).upper()}'"
        feature_set = _query_with_retry(
            layer,
            where=where,
            out_fields="sector",
            return_geometry=False,
            return_distinct_values=True,
            order_by_fields="sector",
        )
        live_sectors = _distinct_text(feature_set.features, "sector")
    except Exception:  # pragma: no cover - depends on live portal availability
        pass
    return sorted(
        set(live_sectors) | set(_fallback_sectors(district, fallback)),
        key=str.casefold,
    )


def list_cells(
    district: str,
    sectors: Sequence[str],
    boundary_service_url: str | None = None,
    gis: Any | None = None,
) -> list[str]:
    """Return available cell names for one district and selected sectors."""
    clean_district = str(district or "").strip()
    clean_sectors = [str(item).strip() for item in sectors if str(item).strip()]
    if not clean_district or not clean_sectors:
        return []
    fallback = load_cell_index()
    live_cells: list[str] = []
    try:
        from arcgis.features import FeatureLayer

        active_gis = _get_gis(gis)
        url = boundary_service_url or os.environ.get(
            "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
        )
        layer = FeatureLayer(_layer_url(url, CELL_LAYER_ID), active_gis)
        quoted = ",".join(f"'{_sql_text(item).upper()}'" for item in clean_sectors)
        where = (
            f"UPPER(district) = '{_sql_text(clean_district).upper()}' "
            f"AND UPPER(sector) IN ({quoted})"
        )
        feature_set = _query_with_retry(
            layer,
            where=where,
            out_fields="cell",
            return_geometry=False,
            return_distinct_values=True,
            order_by_fields="cell",
        )
        live_cells = _distinct_text(feature_set.features, "cell")
    except Exception:  # pragma: no cover - depends on live portal availability
        pass
    return sorted(
        set(live_cells) | set(_fallback_cells(clean_district, clean_sectors, fallback)),
        key=str.casefold,
    )


def _get_sector_features(
    sector_layer: Any,
    district: str | Sequence[str] | None,
    sectors: Sequence[str] | None,
    cell_layer: Any | None = None,
    cells: Sequence[str] | None = None,
) -> Any | None:
    """Query the exact union of selected district, sector and cell boundaries.

    Sector and cell names must be district-qualified for multi-district runs.
    """
    where = district_filters(district, sectors, cells)
    if where is None:
        return None
    clean_cells = [str(value).strip() for value in (cells or []) if str(value).strip()]
    if clean_cells and cell_layer is None:
        raise ValueError("The cell boundary layer is not available.")
    boundary_layer = cell_layer if clean_cells else sector_layer
    out_fields = (
        "objectid,district,sector,cell" if clean_cells
        else "objectid,district,sector"
    )
    feature_set = _query_with_retry(
        boundary_layer, where=where, out_fields=out_fields,
        return_geometry=True, return_all_records=True,
    )
    if not feature_set.features:
        raise ValueError("No administrative boundary was found for the selected districts/sectors/cells.")
    return feature_set


def _combined_boundary_geometry(
    sector_feature_set: Any | None,
    custom_boundary_geometry: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    """Return one REST polygon for either an administrative or drawn boundary."""
    if custom_boundary_geometry is not None:
        return custom_boundary_geometry
    if sector_feature_set is None:
        return None
    rings = []
    for feature in sector_feature_set.features:
        rings.extend(feature.geometry.get("rings", []))
    return {
        "rings": rings,
        "spatialReference": sector_feature_set.spatial_reference,
    }


def _geometry_filter(
    boundary_geometry: dict[str, Any],
    spatial_relationship: str,
) -> Any:
    from arcgis.geometry.filters import contains, intersects

    spatial_reference = boundary_geometry.get("spatialReference", {})
    if spatial_relationship == "contains":
        return contains(boundary_geometry, sr=spatial_reference)
    return intersects(boundary_geometry, sr=spatial_reference)


def _assert_complete_features(response: Any, label: str) -> list[Any]:
    """Reject incomplete live FeatureServer responses without another request.

    Silent truncation or repeated OBJECTIDs would overstate/understate surveyor
    production. In this situation, failing without an Excel report is safer.
    """
    if not hasattr(response, "features") or response.features is None:
        raise RuntimeError(f"{label}: FeatureServer response contains no feature list.")
    if getattr(response, "exceeded_transfer_limit", False):
        raise RuntimeError(
            f"{label}: FeatureServer reported a transfer limit; "
            "report is not published because records may be missing."
        )
    return response.features


def _query_lines(
    line_layer: Any,
    where: str,
    version: str,
    sector_feature_set: Any | None,
    custom_boundary_geometry: dict[str, Any] | None = None,
    spatial_relationship: str = "intersects",
    fields: str | None = None,
    return_geometry: bool = False,
) -> dict[int, dict[str, Any]]:
    active_fields = fields or (
        "OBJECTID,GLOBALID,created_user,created_date,ASSETGROUP,ASSETTYPE,"
        "Shape__Length,Transformer_Name,Sector"
    )
    common = {
        "where": where,
        "out_fields": active_fields,
        "return_geometry": return_geometry,
        "gdb_version": version,
        "return_all_records": True,
    }
    boundary_geometry = _combined_boundary_geometry(
        sector_feature_set, custom_boundary_geometry
    )
    if boundary_geometry is None:
        feature_set = _query_with_retry(line_layer, **common)
    else:
        feature_set = _query_with_retry(
            line_layer,
            geometry_filter=_geometry_filter(
                boundary_geometry, spatial_relationship
            ),
            **common,
        )

    spatial_reference = getattr(feature_set, "spatial_reference", None)
    if not spatial_reference and boundary_geometry:
        spatial_reference = boundary_geometry.get("spatialReference", {})
    spatial_reference = dict(spatial_reference or {})
    records = {}
    for feature in _assert_complete_features(feature_set, "LV line progress"):
        attributes = dict(feature.attributes or {})
        raw_id = attributes.get("OBJECTID")
        if raw_id is None:
            raise RuntimeError("LV line progress: missing source OBJECTID.")
        objectid = int(raw_id)
        if objectid in records:
            raise RuntimeError(
                f"LV line progress: OBJECTID {objectid} was returned twice. "
                "A deduplicated report could conceal server paging errors."
            )
        if return_geometry:
            if not feature.geometry:
                raise RuntimeError(f"LV line progress: missing geometry for OBJECTID {objectid}.")
            geometry = dict(feature.geometry)
            if spatial_reference:
                geometry["spatialReference"] = spatial_reference
            attributes["_geometry"] = geometry
        records[objectid] = attributes
    return records


def _normalize_transformer_name(value: Any) -> str:
    """Normalize case and whitespace for recorded-label review."""
    return " ".join(str(value or "").split()).casefold()


def _query_local_transformers(
    transformer_layer: Any,
    boundary_geometry: dict[str, Any],
    version: str,
) -> list[dict[str, Any]]:
    feature_set = _query_with_retry(
        transformer_layer,
        where=(f"ASSETGROUP = {DISTRIBUTION_TRANSFORMER_ASSETGROUP} "
               "AND ASSETTYPE IN (1,2,3,4,5,6)"),
        out_fields=(
            f"OBJECTID,GLOBALID,{TRANSFORMER_POINT_NAME_FIELD},ASSETGROUP,"
            "ASSETTYPE,created_user,created_date"
        ),
        return_geometry=False,
        geometry_filter=_geometry_filter(boundary_geometry, "intersects"),
        gdb_version=version,
        return_all_records=True,
    )
    seen = set()
    transformers = []
    for feature in _assert_complete_features(feature_set, "Localized transformers"):
        attributes = dict(feature.attributes or {})
        raw_id = attributes.get("OBJECTID")
        if raw_id is None:
            raise RuntimeError("Localized transformers: missing source OBJECTID.")
        oid = int(raw_id)
        if oid in seen:
            raise RuntimeError(f"Localized transformers: repeated OBJECTID {oid}.")
        seen.add(oid)
        transformers.append(attributes)
    return transformers


def _build_connected_transformer_groups(
    transformer_points: Sequence[dict[str, Any]],
    trace_results: dict[int, dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    """Keep each point asset separate, including points with repeated labels."""
    groups = []
    for point in transformer_points:
        object_id = int(point["OBJECTID"])
        result = trace_results[object_id]
        display_name = " ".join(str(point.get(TRANSFORMER_POINT_NAME_FIELD) or "").split())
        transformer_ids = result["connected_transformer_objectids"]
        safe = (transformer_ids == {object_id}) and not bool(result.get("warnings"))
        groups.append({
            "key": str(object_id),
            "transformer_name": display_name or "[Unnamed distribution transformer]",
            "objectids": [object_id],
            "trace_line_objectids": set(result["line_objectids"]),
            "trackable": safe,
            "match_basis": (
                "Connected trace; one transformer point; no warnings"
                if safe else
                "Connected trace reported warnings; review source"
                if result.get("warnings") else
                f"Connected trace reaches {len(transformer_ids)} transformer points; review"
            ),
            "network_element_count": result["network_element_count"],
        })
    groups.sort(key=lambda item: (item["transformer_name"].casefold(), item["objectids"][0]))
    for index, group in enumerate(groups, start=1):
        group["transformer_no"] = f"T-{index:03d}"
    return groups, {group["key"]: group for group in groups}


def _query_line_geometries(
    line_layer: Any,
    object_ids: Sequence[int],
    version: str,
) -> dict[int, dict[str, Any]]:
    geometries: dict[int, dict[str, Any]] = {}
    clean_ids = sorted({int(object_id) for object_id in object_ids})
    for start in range(0, len(clean_ids), 500):
        batch = clean_ids[start:start + 500]
        feature_set = _query_with_retry(
            line_layer,
            object_ids=",".join(str(item) for item in batch),
            out_fields="OBJECTID",
            return_geometry=True,
            gdb_version=version,
            return_all_records=True,
        )
        spatial_reference = dict(
            getattr(feature_set, "spatial_reference", None) or {}
        )
        expected_ids = set(batch)
        returned_ids = set()
        for feature in _assert_complete_features(feature_set, "LV line geometry"):
            attributes = dict(feature.attributes or {})
            raw_id = attributes.get("OBJECTID")
            if raw_id is None:
                raise RuntimeError("LV line geometry: missing source OBJECTID.")
            oid = int(raw_id)
            if oid not in expected_ids or oid in returned_ids or oid in geometries:
                raise RuntimeError(
                    f"LV line geometry: unexpected or repeated OBJECTID {oid}."
                )
            if not feature.geometry:
                raise RuntimeError(f"LV line geometry: missing source geometry for {oid}.")
            geometry = dict(feature.geometry)
            if spatial_reference:
                geometry["spatialReference"] = spatial_reference
            elif not geometry.get("spatialReference"):
                raise RuntimeError(f"LV line geometry: no verified CRS for OBJECTID {oid}.")
            geometries[oid] = geometry
            returned_ids.add(oid)
        missing = expected_ids - returned_ids
        if missing:
            raise RuntimeError(
                f"LV line geometry: {len(missing)} source OBJECTIDs were omitted "
                "from a requested batch; report is not published."
            )
    return geometries


def _geometry_extent(geometries: Iterable[dict[str, Any]]) -> dict[str, float] | None:
    coordinates = []
    for geometry in geometries:
        for path in geometry.get("paths", []):
            coordinates.extend(point[:2] for point in path if len(point) >= 2)
    if not coordinates:
        return None
    xs = [point[0] for point in coordinates]
    ys = [point[1] for point in coordinates]
    return {"xmin": min(xs), "ymin": min(ys), "xmax": max(xs), "ymax": max(ys)}


def _administrative_areas_for_lines(
    cell_layer: Any,
    line_geometries: dict[int, dict[str, Any]],
) -> dict[int, dict[str, list[str]]]:
    """Locate exception lines against national cells using one candidate query."""
    if not line_geometries:
        return {}
    first_geometry = next(iter(line_geometries.values()))
    spatial_reference = dict(first_geometry.get("spatialReference") or {})
    extent = _geometry_extent(line_geometries.values())
    if extent is None:
        return {}
    envelope = {**extent, "spatialReference": spatial_reference}
    candidates = _query_with_retry(
        cell_layer,
        where="1=1",
        out_fields="district,sector,cell",
        return_geometry=True,
        geometry_filter=_geometry_filter(envelope, "intersects"),
        return_all_records=True,
    )
    candidate_sr = dict(candidates.spatial_reference or spatial_reference)

    try:
        import arcpy

        candidate_shapes = []
        for feature in candidates.features:
            geometry = dict(feature.geometry)
            geometry["spatialReference"] = candidate_sr
            candidate_shapes.append((dict(feature.attributes), arcpy.AsShape(geometry, True)))
        result = {}
        for object_id, geometry in line_geometries.items():
            line_shape = arcpy.AsShape(geometry, True)
            matches = [
                attributes
                for attributes, area_shape in candidate_shapes
                if not line_shape.disjoint(area_shape)
            ]
            result[object_id] = {
                field: sorted({
                    str(item.get(field) or "").strip()
                    for item in matches
                    if str(item.get(field) or "").strip()
                }, key=str.casefold)
                for field in ("district", "sector", "cell")
            }
        return result
    except Exception:  # pragma: no cover - ArcGIS Pro supplies arcpy
        result = {}
        for object_id, geometry in line_geometries.items():
            feature_set = _query_with_retry(
                cell_layer,
                where="1=1",
                out_fields="district,sector,cell",
                return_geometry=False,
                geometry_filter=_geometry_filter(geometry, "intersects"),
                return_all_records=True,
            )
            result[object_id] = {
                field: _distinct_text(feature_set.features, field)
                for field in ("district", "sector", "cell")
            }
        return result


def _epoch_to_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, (int, float)):
        return datetime.utcfromtimestamp(float(value) / 1000.0)
    return None


def list_available_users(
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    district: str | Sequence[str] | None = None,
    sectors: Sequence[str] | None = None,
    cells: Sequence[str] | None = None,
    custom_boundary_geometry: dict[str, Any] | None = None,
    service_url: str | None = None,
    boundary_service_url: str | None = None,
    version: str = "sde.DEFAULT",
    gis: Any | None = None,
) -> list[str]:
    """Return creators with LV line records in the selected period and area."""
    active_gis = _get_gis(gis)
    active_service_url = service_url or os.environ.get(
        "UN_REPORT_SERVICE_URL", DEFAULT_LINE_SERVICE_URL
    )
    active_boundary_url = boundary_service_url or os.environ.get(
        "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
    )
    line_layer, sector_layer, cell_layer = _feature_layers(
        active_gis, active_service_url, active_boundary_url
    )
    sector_feature_set = None
    if custom_boundary_geometry is None:
        sector_feature_set = _get_sector_features(
            sector_layer, district, sectors, cell_layer, cells
        )
    where = _build_line_where(start_date, end_date)
    records = _query_lines(
        line_layer,
        where,
        version,
        sector_feature_set,
        custom_boundary_geometry,
    )
    return sorted(
        {
            str(attributes.get("created_user")).strip()
            for attributes in records.values()
            if attributes.get("created_user") not in (None, "")
        },
        key=str.casefold,
    )


def _ordinal(value: int) -> str:
    if 10 <= value % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(value % 10, "th")
    return f"{value}{suffix}"


def _date_range_label(start: date, end: date) -> str:
    if start.year == end.year:
        start_label = f"{_ordinal(start.day)} {start.strftime('%B')}"
    else:
        start_label = f"{_ordinal(start.day)} {start.strftime('%B %Y')}"
    end_label = f"{_ordinal(end.day)} {end.strftime('%B %Y')}"
    return f"From {start_label} to {end_label}"


def _natural_join(values: Sequence[str]) -> str:
    clean = [str(item).strip() for item in values if str(item).strip()]
    if not clean:
        return ""
    if len(clean) == 1:
        return clean[0]
    if len(clean) == 2:
        return f"{clean[0]} and {clean[1]}"
    return f"{', '.join(clean[:-1])} and {clean[-1]}"


def build_report_title(
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    district: str | Sequence[str] | None = None,
    sectors: Sequence[str] | None = None,
    cells: Sequence[str] | None = None,
    custom_boundary: bool = False,
) -> str:
    """Create a correct title for one or many district selections."""
    start = _as_date(start_date)
    end = _as_date(end_date)
    districts = selected_districts(district)
    clean_sectors = [str(item).strip() for item in sectors or [] if str(item).strip()]
    clean_cells = [str(item).strip() for item in cells or [] if str(item).strip()]
    if custom_boundary:
        area = " CUSTOM BOUNDARY"
    elif len(districts) > 1:
        area = f" {_natural_join([d.upper() for d in districts])} DISTRICTS"
        if clean_sectors:
            area += f": {_natural_join(clean_sectors)} " + ("Sector" if len(clean_sectors) == 1 else "Sectors")
        else:
            area += ": All Sectors"
    elif districts and clean_sectors:
        sector_word = "Sector" if len(clean_sectors) == 1 else "Sectors"
        area = f" {districts[0].upper()} DISTRICT: {_natural_join(clean_sectors)} {sector_word}"
    elif districts:
        area = f" {districts[0].upper()} DISTRICT: All Sectors"
    else:
        area = ""
    if clean_cells and not custom_boundary:
        cell_word = "Cell" if len(clean_cells) == 1 else "Cells"
        area += f"; {_natural_join(clean_cells)} {cell_word}"
    return f"Progress Report{area} ({_date_range_label(start, end)})"


def _load_name_overrides(path: str | Path | None) -> dict[str, str]:
    candidate = Path(path).expanduser() if path else DEFAULT_USER_NAME_MAP_PATH
    if not candidate.exists():
        return {}
    with candidate.open("r", encoding="utf-8") as stream:
        data = json.load(stream)
    if not isinstance(data, dict):
        raise ValueError("The report user-name mapping must be a JSON object.")
    return {
        str(username).strip(): str(full_name).strip()
        for username, full_name in data.items()
        if str(username).strip() and str(full_name).strip()
    }


def _resolve_full_names(
    gis: Any,
    usernames: Sequence[str],
    name_overrides: dict[str, str] | None = None,
) -> dict[str, str]:
    overrides = name_overrides or {}
    names: dict[str, str] = {}
    for username in usernames:
        if username in overrides:
            names[username] = overrides[username]
            continue
        full_name = username
        try:
            portal_user = gis.users.get(username)
            if portal_user is not None:
                full_name = str(
                    getattr(portal_user, "fullName", None)
                    or getattr(portal_user, "fullname", None)
                    or username
                ).strip()
        except Exception:  # pragma: no cover - depends on portal availability
            pass
        names[username] = full_name
    return names


def _create_transformer_tracking_sheet(
    workbook: Any,
    title: str,
    transformer_register: Sequence[dict[str, Any]],
    boundary_log: Sequence[dict[str, Any]],
    localization_warnings: Sequence[dict[str, Any]],
    tracking_context: dict[str, Any],
) -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    sheet = workbook.create_sheet("Transformer Tracking Log")
    sheet.sheet_view.showGridLines = False
    navy = "17365D"
    blue = "5B9BD5"
    light_blue = "D9EAF7"
    green = "70AD47"
    light_green = "E2F0D9"
    orange = "F4B183"
    light_orange = "FCE4D6"
    red = "C00000"
    light_red = "F4CCCC"
    gray = "E7E6E6"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="BFBFBF")
    section_border = Border(bottom=thin_gray)

    sheet.merge_cells("A1:M1")
    sheet["A1"] = f"Transformer Boundary Continuation Log - {title}"
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].font = Font(name="Calibri", size=14, bold=True, color=white)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 36

    sheet.merge_cells("A2:M2")
    sheet["A2"] = (
        "The selected boundary localizes distribution-transformer points. A validated "
        "Connected utility-network trace supplies distinct LV feature IDs for each point. "
        "Survey dates and users are applied afterward. Lines reached from multiple "
        "transformers are listed below for review and excluded from totals."
    )
    sheet["A2"].fill = PatternFill("solid", fgColor=light_blue)
    sheet["A2"].font = Font(name="Calibri", size=10, italic=True, color=navy)
    sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 34

    register_start = 8
    register_end = register_start + len(transformer_register) - 1
    summary_pairs = [
        ("Transformers", "B4"),
        ("Surveyors", "D4"),
        ("LV records", "F4"),
        ("LV km", "H4"),
        ("Cross / outside", "J4"),
        ("Exception km", "L4"),
    ]
    for column, (label, _) in zip((1, 3, 5, 7, 9, 11), summary_pairs):
        cell = sheet.cell(4, column, label)
        cell.fill = PatternFill("solid", fgColor=light_green)
        cell.font = Font(bold=True, color=navy)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        value = sheet.cell(4, column + 1)
        value.fill = PatternFill("solid", fgColor=green)
        value.font = Font(bold=True, color=white)
        value.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = value.border = section_border
    sheet.row_dimensions[4].height = 30
    if transformer_register:
        sheet["B4"] = f"=COUNTA(A{register_start}:A{register_end})"
        sheet["F4"] = f"=SUM(E{register_start}:E{register_end})"
        sheet["H4"] = f"=SUM(F{register_start}:F{register_end})"
        sheet["J4"] = f"=SUM(G{register_start}:H{register_end})"
        sheet["L4"] = f"=SUM(I{register_start}:I{register_end})"
    else:
        for cell_ref in ("B4", "F4", "H4", "J4", "L4"):
            sheet[cell_ref] = 0
    sheet["D4"] = int(tracking_context.get("surveyors_followed", 0))
    sheet["H4"].number_format = "0.00"
    sheet["L4"].number_format = "0.00"

    sheet["A5"] = "Warnings"
    sheet["A5"].fill = PatternFill("solid", fgColor=light_orange)
    sheet["A5"].font = Font(bold=True, color=red)
    sheet["B5"] = len(localization_warnings)
    sheet["B5"].fill = PatternFill("solid", fgColor=orange)
    sheet["B5"].font = Font(bold=True)
    sheet["B5"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("C5:M5")
    sheet["C5"] = f"Selected boundary: {tracking_context.get('boundary_label', '')}"
    sheet["C5"].fill = PatternFill("solid", fgColor=gray)
    sheet["C5"].font = Font(bold=True, color=navy)
    sheet["C5"].alignment = Alignment(vertical="center", wrap_text=True)

    register_headers = [
        "Transformer No.", "Transformer Name", "System OBJECTID(s)",
        "Surveyor(s) Working", "LV Records", "Total km", "Crossing Records",
        "Outside Records", "Exception km", "Match Evidence",
    ]
    sheet.merge_cells("A6:J6")
    sheet["A6"] = "Localized Transformer Register"
    sheet["A6"].fill = PatternFill("solid", fgColor=navy)
    sheet["A6"].font = Font(bold=True, color=white)
    sheet["A6"].alignment = Alignment(horizontal="left")
    for column, header in enumerate(register_headers, start=1):
        cell = sheet.cell(7, column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[7].height = 30
    for row_index, item in enumerate(transformer_register, start=register_start):
        values = [
            item["transformer_no"], item["transformer_name"], item["objectids"],
            item["surveyors"], item["line_records"], item["total_km"],
            item["crossing_records"], item["outside_records"], item["exception_km"],
            item.get("match_basis", "Connected trace"),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.border = section_border
            cell.alignment = Alignment(vertical="top", wrap_text=column in (2, 3, 4))
        for column in (6, 9):
            sheet.cell(row_index, column).number_format = "0.00"
    if transformer_register:
        table = Table(
            displayName="LocalizedTransformerRegister",
            ref=f"A7:J{register_end}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        sheet.add_table(table)

    detail_title_row = max(9, register_end + 2)
    detail_header_row = detail_title_row + 1
    detail_start = detail_header_row + 1
    detail_end = detail_start + len(boundary_log) - 1
    sheet.merge_cells(start_row=detail_title_row, start_column=1, end_row=detail_title_row, end_column=13)
    sheet.cell(detail_title_row, 1, "Boundary Continuation Detail")
    sheet.cell(detail_title_row, 1).fill = PatternFill("solid", fgColor=navy)
    sheet.cell(detail_title_row, 1).font = Font(bold=True, color=white)
    detail_headers = [
        "S/N", "Transformer No.", "Transformer Name", "Transformer OBJECTID(s)",
        "Creator Name", "User Name", "Line OBJECTID", "Boundary Status",
        "Line km", "District(s) Touched", "Sector(s) Touched", "Cell(s) Touched",
        "Created Date",
    ]
    for column, header in enumerate(detail_headers, start=1):
        cell = sheet.cell(detail_header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[detail_header_row].height = 34
    for row_index, item in enumerate(boundary_log, start=detail_start):
        values = [
            row_index - detail_start + 1,
            item["transformer_no"], item["transformer_name"], item["objectids"],
            item["full_name"], item["username"], item["line_objectid"],
            item["boundary_status"], item["length_km"], item["districts"],
            item["sectors"], item["cells"], item["created_date"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.border = section_border
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 5, 8, 10, 11, 12))
        sheet.cell(row_index, 9).number_format = "0.000"
        sheet.cell(row_index, 13).number_format = "yyyy-mm-dd hh:mm"
    if boundary_log:
        table = Table(
            displayName="TransformerBoundaryContinuation",
            ref=f"A{detail_header_row}:M{detail_end}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        sheet.add_table(table)
        status_range = f"H{detail_start}:H{detail_end}"
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'LEFT(H{detail_start},7)="CROSSES"'],
                fill=PatternFill("solid", fgColor=light_orange),
                font=Font(color="9C5700", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'LEFT(H{detail_start},8)="FOLLOWED"'],
                fill=PatternFill("solid", fgColor=light_red),
                font=Font(color=red, bold=True),
            ),
        )

    warning_title_row = max(detail_header_row + 2, detail_end + 2)
    warning_header_row = warning_title_row + 1
    warning_start = warning_header_row + 1
    warning_end = warning_start + len(localization_warnings) - 1
    sheet.merge_cells(start_row=warning_title_row, start_column=1, end_row=warning_title_row, end_column=9)
    sheet.cell(warning_title_row, 1, "Unassigned LV Candidates - Not Counted")
    sheet.cell(warning_title_row, 1).fill = PatternFill("solid", fgColor=orange)
    sheet.cell(warning_title_row, 1).font = Font(bold=True, color=red)
    warning_headers = [
        "S/N", "Line OBJECTID", "Creator Name", "User Name",
        "Entered Transformer Name", "Line km", "Recorded Sector", "Boundary Status", "Reason",
    ]
    for column, header in enumerate(warning_headers, start=1):
        cell = sheet.cell(warning_header_row, column, header)
        cell.fill = PatternFill("solid", fgColor=gray)
        cell.font = Font(bold=True, color=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, item in enumerate(localization_warnings, start=warning_start):
        values = [
            row_index - warning_start + 1, item["line_objectid"],
            item["full_name"], item["username"], item["transformer_name"],
            item["length_km"], item["recorded_sector"],
            item.get("boundary_status", ""), item["reason"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.border = section_border
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 5, 7, 8))
        sheet.cell(row_index, 6).number_format = "0.000"
    if localization_warnings:
        table = Table(
            displayName="TransformerLocalizationWarnings",
            ref=f"A{warning_header_row}:I{warning_end}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        sheet.add_table(table)

    widths = {
        "A": 12, "B": 19, "C": 28, "D": 24, "E": 30, "F": 25,
        "G": 16, "H": 28, "I": 12, "J": 23, "K": 24, "L": 27, "M": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A7:J{max(7, register_end)}"


def _create_connected_name_review_sheet(
    workbook: Any, items: Sequence[dict[str, Any]],
) -> None:
    """Show connected lines whose recorded transformer label needs correction."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    sheet = workbook.create_sheet("Connected Name Review")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Connected LV lines with a different recorded Transformer Name"
    sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells("A2:G2")
    sheet["A2"] = (
        "These lines are included in survey totals because the utility-network trace "
        "connected them to the listed transformer. Review and correct their recorded label."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[2].height = 32
    headers = [
        "Transformer OBJECTID", "Transformer Name", "Line OBJECTID",
        "Recorded Transformer Name", "User Name", "Boundary Status", "Line km",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[4].height = 30
    for row_index, item in enumerate(items, start=5):
        for column, value in enumerate((
            item["transformer_objectid"], item["transformer_name"],
            item["line_objectid"], item["recorded_transformer_name"],
            item["username"], item["boundary_status"], item["length_km"],
        ), start=1):
            sheet.cell(row_index, column, value)
        sheet.cell(row_index, 7).number_format = "0.000"
    table = Table(displayName="ConnectedNameReview", ref=f"A4:G{4 + len(items)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    for column, width in {"A": 24, "B": 28, "C": 20, "D": 32, "E": 26, "F": 30, "G": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"


def _create_workbook(
    output_path: Path,
    title: str,
    rows: Sequence[dict[str, Any]],
    transformer_register: Sequence[dict[str, Any]] | None = None,
    boundary_log: Sequence[dict[str, Any]] | None = None,
    localization_warnings: Sequence[dict[str, Any]] | None = None,
    connected_name_review: Sequence[dict[str, Any]] | None = None,
    tracking_context: dict[str, Any] | None = None,
    full_feeder_inventory: dict[str, Any] | None = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    # Keep the generated report identical to the approved field template.
    sheet.title = "Sheet2"
    sheet.sheet_format.defaultRowHeight = 15.75

    green = "C6E0B4"
    header_blue = "B4C6E7"
    total_yellow = "FFD966"
    black = "000000"
    thin = Side(style="thin", color=black)
    medium = Side(style="medium", color=black)
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.merge_cells("A2:D2")
    title_cell = sheet["A2"]
    title_cell.value = title
    title_cell.fill = PatternFill("solid", fgColor=green)
    title_cell.font = Font(name="Calibri", size=11, bold=True, color=black)
    title_cell.alignment = Alignment(horizontal="center")
    title_cell.border = Border(left=medium, right=medium, top=medium, bottom=medium)
    sheet.row_dimensions[2].height = 18

    headers = [
        "S/N",
        "Creator Name",
        "User Name",
        "Surveyed LV Line Records(km)",
    ]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=header_blue)
        cell.font = Font(name="Calibri", size=11, bold=True, color=black)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=column == 4,
        )
        cell.border = cell_border
    sheet.row_dimensions[4].height = 30

    first_data_row = 5
    for index, row in enumerate(rows, start=1):
        excel_row = first_data_row + index - 1
        values = [index, row["full_name"], row["username"], row["length_km"]]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="Calibri", size=11, color=black)
            cell.border = cell_border
        sheet.cell(excel_row, 1).fill = PatternFill("solid", fgColor=green)
        sheet.cell(excel_row, 1).font = Font(name="Calibri", size=11, color=black)

    total_row = first_data_row + len(rows)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    total_label = sheet.cell(total_row, 1, "Total")
    total_label.alignment = Alignment(horizontal="left")
    total_label.font = Font(name="Calibri", size=11, bold=True, color=black)
    total_label.fill = PatternFill("solid", fgColor=total_yellow)
    for column in range(1, 4):
        cell = sheet.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor=total_yellow)
        cell.border = cell_border
    total_value = sheet.cell(total_row, 4)
    if rows:
        total_value.value = f"=SUM(D{first_data_row}:D{total_row - 1})"
    else:
        total_value.value = 0
    total_value.number_format = "0.00"
    total_value.font = Font(name="Calibri", size=11, bold=True, color=black)
    total_value.fill = PatternFill("solid", fgColor=total_yellow)
    total_value.border = cell_border

    widths = {
        "A": 7.28515625,
        "B": 63.140625,
        "C": 37.42578125,
        "D": 18.28515625,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    if tracking_context is not None:
        _create_transformer_tracking_sheet(
            workbook,
            title,
            transformer_register or [],
            boundary_log or [],
            localization_warnings or [],
            tracking_context,
        )
    if connected_name_review:
        _create_connected_name_review_sheet(workbook, connected_name_review)
    if full_feeder_inventory is not None:
        from .feeder_inventory import add_feeder_inventory_sheets
        add_feeder_inventory_sheets(
            workbook,
            full_feeder_inventory["rows"],
            full_feeder_inventory["summaries"],
            full_feeder_inventory["scope"],
        )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_topology_blocker_diagnostic(
    requested_output: str | Path,
    error: TopologyDirtyAreaError,
    *,
    start_date: Any,
    end_date: Any,
    districts: Sequence[str],
    sectors: Sequence[str],
    cells: Sequence[str],
    usernames: Sequence[str],
) -> Path:
    """Write an unmistakably non-production Excel diagnostic, without GIS calls.

    The blocked progress XLSX is NEVER written or overwritten. This companion
    identifies the first failed topology trace and ALL ElectricLine GUIDs
    returned by REG for that trace. Previous successful traces are not
    treated as complete production totals. A later transformer may also be
    dirty; the operator must validate and rerun the entire report.
    """
    from uuid import uuid4
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    requested = Path(requested_output).expanduser().resolve()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    diagnostic = requested.with_name(
        requested.stem + "_TRACE_BLOCKED_" + stamp + "_" + uuid4().hex[:8] + ".xlsx"
    )
    if diagnostic.exists():
        raise FileExistsError("Refusing to overwrite an existing trace diagnostic.")

    def safe_text(value: Any) -> str:
        text = str(value if value is not None else "")
        return "'" + text if text.startswith(("=", "+", "-", "@")) else text

    workbook = Workbook()
    summary = workbook.active
    summary.title = "TRACE BLOCKED"
    summary.append(["STATUS: TRACE BLOCKED — NOT A PROGRESS REPORT", "Value"])
    entries = [
        ("Report generated", "NO — do not use this workbook for progress totals"),
        ("Requested report filename", requested.name),
        ("Original report untouched", "YES"),
        ("Starting district(s)", "; ".join(map(str, districts))),
        ("Starting sector(s)", "; ".join(map(str, sectors)) or "(not restricted)"),
        ("Starting cell(s)", "; ".join(map(str, cells)) or "(not restricted)"),
        ("From date", start_date),
        ("To date", end_date),
        ("Selected surveyor usernames", "; ".join(map(str, usernames))),
        ("REG source version", error.version),
        ("Failed transformer OBJECTID", error.transformer_objectid),
        ("Failed transformer GlobalID", error.transformer_globalid),
        ("Trace position", f"{error.trace_position} of {error.trace_total}"),
        ("Earlier traces", f"{error.trace_position - 1} returned but NOT published"),
        ("Remaining transformers", max(0, error.trace_total-error.trace_position)),
        ("Dirty ElectricLine GlobalIDs returned", len(error.dirty_line_globalids)),
        ("Next step", "Ask an authorized REG Utility Network editor to inspect the listed dirty features, "
                      "validate the affected topology and rerun the whole report."),
        ("Safety", "No REG edits made. validateConsistency remains TRUE. No partial kilometres."),
        ("Additional server message", error.original_error[:15000]),
    ]
    for key, value in entries:
        summary.append([safe_text(key),safe_text(value)])
    lines = workbook.create_sheet("Dirty Electric Lines")
    lines.append(["#", "Dirty ElectricLine GlobalID", "Reported at transformer OBJECTID",
                  "Topology version", "Action"])
    for index,guid in enumerate(error.dirty_line_globalids,1):
        lines.append([
            index, safe_text(guid), error.transformer_objectid, error.version,
            "Locate by GlobalID; inspect Dirty Areas; validate affected extent if authorized",
        ])
    if not error.dirty_line_globalids:
        lines.append([1, "REG returned no extractable ElectricLine GUID",
                      error.transformer_objectid,error.version,
                      "Inspect the original service error on TRACE BLOCKED"])
    notes = workbook.create_sheet("What this means")
    for row in [
        ["CHECK", "INTERPRETATION"],
        ["Trace consistency", "Dirty topology blocked connected trace. Do not switch validateConsistency off."],
        ["Feeder coverage", "No credible full transformer-origin kilometres can be published from this attempt."],
        ["Scope", "Only the first blocked trace is shown. Other transformers may have additional dirty areas."],
        ["Unrelated error", "'Class not registered' may require separate REG/ArcGIS Server investigation "
                            "if it persists after topology validation."],
        ["What to do", "Coordinate with the REG Utility Network administrator; validation may affect other editors."],
        ["Rerun", "After topology validation, run the same manual report again with unchanged parameters."],
    ]:
        notes.append([safe_text(value) for value in row])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="7F1D1D")
            cell.font = Font(bold=True, color="FFFFFF")
        sheet.column_dimensions["A"].width = 45 if sheet is summary else 30
        sheet.column_dimensions["B"].width = 100 if sheet is summary else 68
        if sheet is lines:
            sheet.column_dimensions["E"].width=82
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical="top",wrap_text=True)
    diagnostic.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(diagnostic)
    workbook.close()
    return diagnostic


def generate_progress_report(
    output_xlsx: str | Path,
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    usernames: Sequence[str],
    district: str | Sequence[str] | None = None,
    sectors: Sequence[str] | None = None,
    cells: Sequence[str] | None = None,
    custom_boundary_geometry: dict[str, Any] | None = None,
    service_url: str | None = None,
    boundary_service_url: str | None = None,
    version: str = "sde.DEFAULT",
    user_name_map_path: str | Path | None = None,
    gis: Any | None = None,
    include_full_feeder_inventory: bool = False,
) -> dict[str, Any]:
    """Create selected-surveyor progress and optionally full trace-connected feeder evidence."""
    start = _as_date(start_date)
    end = _as_date(end_date)
    clean_users = list(dict.fromkeys(
        str(item).strip() for item in usernames if str(item).strip()
    ))
    if not clean_users:
        raise ValueError("Select at least one surveyor.")
    if end < start:
        raise ValueError("The end date cannot be earlier than the start date.")
    if include_full_feeder_inventory and len(selected_districts(district)) != 1:
        raise ValueError(
            "A full-feeder report needs exactly one source-transformer district. "
            "Select a district; the report will follow its traced LV lines "
            "even when they cross outside that district."
        )

    active_gis = _get_gis(gis)
    active_service_url = service_url or os.environ.get(
        "UN_REPORT_SERVICE_URL", DEFAULT_LINE_SERVICE_URL
    )
    active_boundary_url = boundary_service_url or os.environ.get(
        "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
    )
    line_layer, sector_layer, cell_layer = _feature_layers(
        active_gis, active_service_url, active_boundary_url
    )
    sector_feature_set = None
    if custom_boundary_geometry is None:
        sector_feature_set = _get_sector_features(
            sector_layer, district, sectors, cell_layer, cells
        )
    where = _build_line_where(start, end, clean_users)
    boundary_geometry = _combined_boundary_geometry(
        sector_feature_set, custom_boundary_geometry
    )
    transformer_groups: list[dict[str, Any]] = []
    transformer_groups_by_key: dict[str, dict[str, Any]] = {}
    intersect_records: dict[int, dict[str, Any]] = {}
    contained_ids: set[int] = set()
    line_areas: dict[int, dict[str, list[str]]] = {}
    all_period_records: dict[int, dict[str, Any]] = {}
    review_reason_by_id: dict[int, str] = {}
    full_feeder_inventory = None

    if boundary_geometry is not None:
        transformer_layer = _transformer_layer(active_gis, active_service_url)
        transformer_points = _query_local_transformers(
            transformer_layer,
            boundary_geometry,
            version,
        )
        trace_results = trace_connected_lv_lines(
            active_gis, active_service_url, transformer_points, version,
            LINE_LAYER_ID, TRANSFORMER_LAYER_ID,
        )
        transformer_groups, transformer_groups_by_key = _build_connected_transformer_groups(
            transformer_points, trace_results
        )
        if include_full_feeder_inventory:
            # Query ALL traced LV sections by source OBJECTID, independently
            # of selected surveyors and report dates. Do not spatially clip
            # to the origin district; geometry is looked up afterwards.
            from .transformer_line_audit import _fetch_line_attributes_by_ids
            from .feeder_inventory import assess_source_feeder_inventory

            traced_line_ids = sorted({
                int(oid) for group in transformer_groups
                for oid in group["trace_line_objectids"]
            })
            all_connected_attributes = _fetch_line_attributes_by_ids(
                line_layer, traced_line_ids, version, [
                    "OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE",
                    "Transformer_Name","created_user","Shape__Length",
                ],
            )
            traced_geometries = _query_line_geometries(
                line_layer, traced_line_ids, version
            )
            if set(traced_geometries) != set(traced_line_ids):
                raise RuntimeError(
                    "Full-feeder geometry retrieval is incomplete. "
                    "Cannot publish a false cross-district report."
                )
            geographic_areas = _administrative_areas_for_lines(
                cell_layer, traced_geometries
            )
            full_feeder_inventory = {
                "all_connected_attributes": all_connected_attributes,
                "geographic_areas": geographic_areas,
            }
        intersect_records = _query_lines(
            line_layer,
            where,
            version,
            sector_feature_set,
            custom_boundary_geometry,
        )
        records = {}
        if transformer_groups:
            all_period_records = _query_lines(
                line_layer,
                where,
                version,
                None,
            )
            contained_ids = set(_query_lines(
                line_layer,
                where,
                version,
                sector_feature_set,
                custom_boundary_geometry,
                spatial_relationship="contains",
                fields="OBJECTID",
            ))
            intersect_ids = set(intersect_records)
            memberships: dict[int, list[dict[str, Any]]] = defaultdict(list)
            for group in transformer_groups:
                for line_id in group["trace_line_objectids"]:
                    memberships[line_id].append(group)
            for object_id, attributes in all_period_records.items():
                candidates = memberships.get(object_id, [])
                if not candidates:
                    continue
                if len(candidates) > 1:
                    review_reason_by_id[object_id] = (
                        "Connected trace reaches this LV line from multiple localized "
                        "transformer points; assignment needs review."
                    )
                    continue
                group = candidates[0]
                if not group["trackable"]:
                    review_reason_by_id[object_id] = (
                        "The connected trace reaches multiple distribution-transformer "
                        "points; assignment needs review."
                    )
                    continue
                attributes["_transformer_key"] = group["key"]
                attributes["_transformer_no"] = group["transformer_no"]
                attributes["_name_mismatch"] = (
                    _normalize_transformer_name(attributes.get("Transformer_Name"))
                    != _normalize_transformer_name(group["transformer_name"])
                )
                if object_id in contained_ids:
                    attributes["_boundary_status"] = "WITHIN SELECTED BOUNDARY"
                elif object_id in intersect_ids:
                    attributes["_boundary_status"] = "CROSSES SELECTED BOUNDARY"
                else:
                    attributes["_boundary_status"] = "FOLLOWED OUTSIDE BOUNDARY"
                records[object_id] = attributes

            exception_ids = [
                object_id
                for object_id, attributes in records.items()
                if attributes["_boundary_status"] != "WITHIN SELECTED BOUNDARY"
            ]
            exception_geometries = _query_line_geometries(
                line_layer, exception_ids, version
            )
            line_areas = _administrative_areas_for_lines(
                cell_layer, exception_geometries
            )
    else:
        records = _query_lines(line_layer, where, version, None)

    lengths_by_user: dict[str, float] = defaultdict(float)
    counts_by_user: dict[str, int] = defaultdict(int)
    for attributes in records.values():
        username = str(attributes.get("created_user") or "").strip()
        if username not in clean_users:
            continue
        lengths_by_user[username] += float(attributes.get("Shape__Length") or 0.0)
        counts_by_user[username] += 1

    reported_users = [
        username for username in clean_users if counts_by_user.get(username, 0) > 0
    ]
    localization_warning_records = []
    if boundary_geometry is not None:
        warning_by_id = {
            object_id: attributes
            for object_id, attributes in intersect_records.items()
            if object_id not in records
        }
        for object_id in review_reason_by_id:
            warning_by_id[object_id] = all_period_records[object_id]
        localization_warning_records = list(warning_by_id.values())
    name_overrides = _load_name_overrides(user_name_map_path)
    names_needed = list(dict.fromkeys(
        reported_users
        + [
            str(item.get("created_user") or "").strip()
            for item in localization_warning_records
            if str(item.get("created_user") or "").strip()
        ]
    ))
    full_names = _resolve_full_names(active_gis, names_needed, name_overrides)
    rows = [
        {
            "full_name": full_names.get(username, username),
            "username": username,
            "length_km": lengths_by_user[username] / 1000.0,
            "record_count": counts_by_user[username],
        }
        for username in reported_users
    ]
    clean_sectors = [str(item).strip() for item in sectors or [] if str(item).strip()]
    clean_cells = [str(item).strip() for item in cells or [] if str(item).strip()]
    districts_selected = selected_districts(district)
    if (
        custom_boundary_geometry is None
        and len(districts_selected) == 1
        and not clean_sectors
        and sector_feature_set is not None
    ):
        clean_sectors = _distinct_text(sector_feature_set.features, "sector")
    title = build_report_title(
        start,
        end,
        district,
        clean_sectors,
        clean_cells,
        custom_boundary=custom_boundary_geometry is not None,
    )

    transformer_register = []
    boundary_log = []
    localization_warnings = []
    connected_name_review = []
    tracking_context = None
    district_label = _natural_join(districts_selected)
    if boundary_geometry is not None:
        for group in transformer_groups:
            group_records = [
                attributes
                for attributes in records.values()
                if attributes.get("_transformer_key") == group["key"]
            ]
            usernames_working = sorted({
                str(item.get("created_user") or "").strip()
                for item in group_records
                if str(item.get("created_user") or "").strip()
            }, key=str.casefold)
            crossing_records = sum(
                item.get("_boundary_status") == "CROSSES SELECTED BOUNDARY"
                for item in group_records
            )
            outside_records = sum(
                item.get("_boundary_status") == "FOLLOWED OUTSIDE BOUNDARY"
                for item in group_records
            )
            exception_km = sum(
                float(item.get("Shape__Length") or 0.0)
                for item in group_records
                if item.get("_boundary_status") != "WITHIN SELECTED BOUNDARY"
            ) / 1000.0
            transformer_register.append({
                "transformer_no": group["transformer_no"],
                "transformer_name": group["transformer_name"],
                "objectids": ", ".join(str(item) for item in group["objectids"]),
                "surveyors": "; ".join(
                    f"{full_names.get(username, username)} ({username})"
                    for username in usernames_working
                ) or "No matching LV work in the selected period",
                "line_records": len(group_records),
                "total_km": sum(
                    float(item.get("Shape__Length") or 0.0)
                    for item in group_records
                ) / 1000.0,
                "crossing_records": crossing_records,
                "outside_records": outside_records,
                "exception_km": exception_km,
                "match_basis": (
                    group["match_basis"] +
                    (f"; {sum(bool(item.get('_name_mismatch')) for item in group_records)} label mismatch(es)"
                     if any(item.get("_name_mismatch") for item in group_records) else "")
                ),
            })

        for object_id, attributes in records.items():
            if attributes.get("_name_mismatch"):
                group = transformer_groups_by_key[attributes["_transformer_key"]]
                connected_name_review.append({
                    "transformer_objectid": group["objectids"][0],
                    "transformer_name": group["transformer_name"],
                    "line_objectid": object_id,
                    "recorded_transformer_name": str(attributes.get("Transformer_Name") or "[Blank]"),
                    "username": str(attributes.get("created_user") or "").strip(),
                    "boundary_status": attributes.get("_boundary_status", ""),
                    "length_km": float(attributes.get("Shape__Length") or 0.0) / 1000.0,
                })
            boundary_status = attributes.get("_boundary_status")
            if boundary_status == "WITHIN SELECTED BOUNDARY":
                continue
            group = transformer_groups_by_key[attributes["_transformer_key"]]
            username = str(attributes.get("created_user") or "").strip()
            areas = line_areas.get(object_id, {})
            boundary_log.append({
                "transformer_no": group["transformer_no"],
                "transformer_name": group["transformer_name"],
                "objectids": ", ".join(str(item) for item in group["objectids"]),
                "full_name": full_names.get(username, username),
                "username": username,
                "line_objectid": object_id,
                "boundary_status": boundary_status,
                "length_km": float(attributes.get("Shape__Length") or 0.0) / 1000.0,
                "districts": "; ".join(areas.get("district", [])),
                "sectors": "; ".join(areas.get("sector", [])),
                "cells": "; ".join(areas.get("cell", [])),
                "created_date": _epoch_to_datetime(attributes.get("created_date")),
            })
        boundary_log.sort(key=lambda item: (
            item["transformer_no"], item["username"].casefold(),
            item["boundary_status"], item["line_objectid"],
        ))
        connected_name_review.sort(key=lambda item: (
            item["transformer_objectid"], item["line_objectid"]
        ))

        for attributes in localization_warning_records:
            username = str(attributes.get("created_user") or "").strip()
            entered_name = " ".join(
                str(attributes.get("Transformer_Name") or "").split()
            )
            localization_warnings.append({
                "line_objectid": int(attributes["OBJECTID"]),
                "full_name": full_names.get(username, username),
                "username": username,
                "transformer_name": entered_name or "[Blank]",
                "length_km": float(attributes.get("Shape__Length") or 0.0) / 1000.0,
                "recorded_sector": str(attributes.get("Sector") or "").strip(),
                "boundary_status": (
                    "WITHIN SELECTED BOUNDARY" if int(attributes["OBJECTID"]) in contained_ids
                    else "INTERSECTS SELECTED BOUNDARY" if int(attributes["OBJECTID"]) in intersect_records
                    else "OUTSIDE SELECTED BOUNDARY"
                ),
                "reason": review_reason_by_id.get(
                    int(attributes["OBJECTID"]),
                    "This LV line was not reached by a connected trace from any "
                    "distribution transformer inside the selected boundary."
                ),
            })
        localization_warnings.sort(key=lambda item: (
            item["transformer_name"].casefold(), item["username"].casefold(),
            item["line_objectid"],
        ))

        if custom_boundary_geometry is not None:
            boundary_label = "Custom drawn polygon"
        elif clean_cells:
            boundary_label = (
                f"{district_label} / {', '.join(clean_sectors)} / "
                f"{', '.join(clean_cells)}"
            ).strip(" / ")
        elif clean_sectors:
            boundary_label = f"{district_label} / {', '.join(clean_sectors)}".strip(" / ")
        else:
            boundary_label = f"{district_label or 'Selected administrative area'} / All sectors"
        tracking_context = {
            "boundary_label": boundary_label,
            "surveyors_followed": len(reported_users),
        }

    if full_feeder_inventory is not None:
        # "records" is the surveyor/date-matched subset. The new feeder
        # sheet is complete across source transformer topology regardless.
        full_rows, full_summaries, full_stats = assess_source_feeder_inventory(
            transformer_groups,
            full_feeder_inventory["all_connected_attributes"],
            full_feeder_inventory["geographic_areas"],
            selected_districts(district)[0],
            set(records),
            selected_origin_sectors=(
                clean_sectors if custom_boundary_geometry is None else []
            ),
        )
        full_feeder_inventory = {
            "rows":full_rows,
            "summaries":full_summaries,
            "scope":{
                **full_stats,
                "district":selected_districts(district)[0],
                "version":version,
            },
        }
    output_path = Path(output_xlsx).expanduser().resolve()
    _create_workbook(
        output_path,
        title,
        rows,
        transformer_register=transformer_register,
        boundary_log=boundary_log,
        localization_warnings=localization_warnings,
        connected_name_review=connected_name_review,
        tracking_context=tracking_context,
        full_feeder_inventory=full_feeder_inventory,
    )

    return {
        "read_only": True,
        "output_xlsx": str(output_path),
        "title": title,
        "from_date": start.isoformat(),
        "to_date_inclusive": end.isoformat(),
        "district": district_label,
        "districts": districts_selected,
        "sectors": clean_sectors,
        "cells": clean_cells,
        "boundary_mode": (
            "custom" if custom_boundary_geometry is not None else "administrative"
        ),
        "selected_users": clean_users,
        "reported_users": reported_users,
        "full_feeder_inventory": (
            full_feeder_inventory["scope"] if full_feeder_inventory is not None
            else None
        ),
        "full_feeder_sheet": (
            "Full Feeder Segments" if full_feeder_inventory is not None else ""
        ),
        "matched_line_records": len(records),
        "localized_transformers": len(transformer_groups),
        "ambiguous_connected_transformers": sum(
            not group["trackable"]
            for group in transformer_groups
        ),
        "connected_lv_feature_ids": len({
            object_id for group in transformer_groups
            for object_id in group["trace_line_objectids"]
        }),
        "trace_network_elements": sum(group["network_element_count"] for group in transformer_groups),
        "name_mismatch_records": sum(bool(item.get("_name_mismatch")) for item in records.values()),
        "match_evidence": "Connected utility-network trace with consistency validation; traversability not evaluated",
        "boundary_crossing_records": sum(
            item.get("_boundary_status") == "CROSSES SELECTED BOUNDARY"
            for item in records.values()
        ),
        "followed_outside_records": sum(
            item.get("_boundary_status") == "FOLLOWED OUTSIDE BOUNDARY"
            for item in records.values()
        ),
        "localization_warnings": len(localization_warnings),
        "tracking_log_sheet": (
            "Transformer Tracking Log" if tracking_context is not None else ""
        ),
        "total_km": sum(item["length_km"] for item in rows),
        "users": rows,
        "line_service_url": _layer_url(active_service_url, LINE_LAYER_ID),
        "boundary_service_url": _layer_url(
            active_boundary_url, CELL_LAYER_ID if clean_cells else SECTOR_LAYER_ID
        ),
        "version": version,
    }
