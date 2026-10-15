"""Carry reviewer decisions between timestamped LV quality workbooks."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
from typing import Any, Sequence


REVIEW_HEADERS = [
    "Record Key", "GlobalID", "Source OBJECTID", "Username", "This Run Status",
    "Completion %", "Missing Attributes", "Invalid Attributes",
    "First Flagged UTC", "Last Flagged UTC", "Flagged Runs", "Last Seen UTC",
    "Review Status", "Assigned To", "Reviewed By", "Reviewed At UTC", "Resolution Notes",
]
MANUAL_FIELDS = ("review_status", "assigned_to", "reviewed_by", "reviewed_at_utc", "resolution_notes")


def find_previous_report(output_path: Path) -> Path | None:
    if not output_path.parent.is_dir():
        return None
    stem = re.sub(r"_\d{8}_\d{6}(?:_\d+)?$", "", output_path.stem)
    pattern = re.compile(rf"^{re.escape(stem)}_\d{{8}}_\d{{6}}(?:_\d+)?\.xlsx$", re.IGNORECASE)
    reports = sorted(
        path for path in output_path.parent.iterdir()
        if path != output_path and path.is_file() and pattern.fullmatch(path.name)
    )
    return reports[-1] if reports else None


def load_review_rows(path: str | Path | None) -> dict[str, dict[str, Any]]:
    if not path:
        return {}
    from openpyxl import load_workbook

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Previous QA report was not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if "Review Queue" not in workbook:
            return {}
        sheet = workbook["Review Queue"]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        header = next(rows, None)
        if list(header or [])[:len(REVIEW_HEADERS)] != REVIEW_HEADERS:
            raise ValueError(f"The Review Queue columns are not valid in: {source}")
        prior: dict[str, dict[str, Any]] = {}
        for values in rows:
            if not values or not values[0]:
                continue
            key = str(values[0]).strip()
            if key in prior:
                raise ValueError(f"Duplicate record key in the prior Review Queue: {key}")
            prior[key] = {
                "record_key": key,
                "globalid": values[1] or "",
                "source_objectid": values[2],
                "username": values[3] or "",
                "this_run_status": values[4] or "NOT_CHECKED",
                "completion_percent": float(values[5] or 0) * 100,
                "missing_attributes": values[6] or "",
                "invalid_attributes": values[7] or "",
                "first_flagged_utc": values[8] or "",
                "last_flagged_utc": values[9] or "",
                "flagged_runs": int(values[10] or 0),
                "last_seen_utc": values[11] or "",
                "review_status": str(values[12] or "OPEN").strip().upper(),
                "assigned_to": values[13] or "",
                "reviewed_by": values[14] or "",
                "reviewed_at_utc": values[15] or "",
                "resolution_notes": values[16] or "",
            }
        return prior
    finally:
        workbook.close()


def build_review_rows(
    results: Sequence[dict[str, Any]],
    previous: dict[str, dict[str, Any]],
    run_utc: datetime,
) -> list[dict[str, Any]]:
    now = run_utc.isoformat()
    current = {key: dict(value, this_run_status="NOT_CHECKED") for key, value in previous.items()}
    for result in results:
        globalid = str(result.get("globalid") or "").strip()
        # A shared GlobalID alone is an unsafe review key. Retain legacy
        # keys for normal rows; key only the colliding rows by OBJECTID too.
        conflict = bool(result.get("duplicate_peer_objectids"))
        key = globalid.casefold() if globalid else f"oid:{result['source_objectid']}"
        if conflict:
            key += f"|oid:{result['source_objectid']}"
        this_run = ("IDENTIFIER_CONFLICT"
                    if conflict and result["status"] != "FLAG"
                    else result["status"])
        needs_review = this_run in {"FLAG", "IDENTIFIER_CONFLICT"}
        if not needs_review and key not in current:
            continue
        old = current.get(key, {})
        status = str(old.get("review_status") or "OPEN").strip().upper()
        if needs_review and status in {"RESOLVED", "DISMISSED"}:
            status = "REOPENED"
        current[key] = {
            "record_key": key,
            "globalid": globalid,
            "source_objectid": result["source_objectid"],
            "username": result["username"],
            "this_run_status": this_run,
            "completion_percent": result["completion_percent"],
            "missing_attributes": "; ".join(result["missing_attributes"]),
            "invalid_attributes": "; ".join(
                [*result["invalid_attributes"], *result.get("identifier_issues", [])]
            ),
            "first_flagged_utc": old.get("first_flagged_utc") or now,
            "last_flagged_utc": now if needs_review else old.get("last_flagged_utc", ""),
            "flagged_runs": int(old.get("flagged_runs") or 0) + needs_review,
            "last_seen_utc": now,
            "review_status": status,
            "assigned_to": old.get("assigned_to", ""),
            "reviewed_by": old.get("reviewed_by", ""),
            "reviewed_at_utc": old.get("reviewed_at_utc", ""),
            "resolution_notes": old.get("resolution_notes", ""),
        }
    return sorted(
        current.values(),
        key=lambda row: (row["this_run_status"] != "FLAG", row["review_status"] == "RESOLVED", row["record_key"]),
    )
