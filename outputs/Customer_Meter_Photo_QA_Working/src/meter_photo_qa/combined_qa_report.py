"""One consolidated QA report: LV lines + transformers + LV poles.

Retains all three existing QA engines. Geometry continuity is a clearly
marked OPTIONAL appendix and never substitutes for in-district data quality.
The source Utility Network service is read-only.
"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date, datetime
from pathlib import Path
import re

from .progress_report import _as_date
from .quality_watcher import run_quality_watcher
from .asset_quality_runner import run_asset_quality
from .lv_feeder_report import run_lv_feeder_report

SHEET_PREFIXES = {
    "lv_line": "LV",
    "transformer": "DT",
    "pole": "Pole",
    "geometry": "Geo",
}
HIGHLIGHT = "16324E"
LIGHT = "E7F0F9"


def _safe_excel_text(value):
    # Avoid Excel formula injection from GIS text attributes. Ordinary
    # formulas in historic reports are preserved and their sheet refs
    # rewritten separately.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet_title(prefix, title):
    full=f"{prefix} {title}"
    if len(full)>31:
        full=full[:31].rstrip()
    return full


def _copy_source_workbook(target, source_path, prefix):
    """Copy source sheets with original data, basic styles, and formulas rebased.

    The original component workbooks remain available untouched for their
    native table formatting, interactive review queue and annotations.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb=load_workbook(source_path, read_only=False, data_only=False)
    old_to_new={s.title:_sheet_title(prefix,s.title) for s in wb.worksheets}
    try:
        for original in wb.worksheets:
            dest=target.create_sheet(old_to_new[original.title])
            dest.sheet_state = original.sheet_state
            dest.sheet_view.showGridLines=False
            dest.freeze_panes=original.freeze_panes
            for row in original.iter_rows():
                for src in row:
                    if src.value is None:
                        continue
                    value=src.value
                    if src.data_type=="f" and isinstance(value,str):
                        # Rebase only own-source refs, never try evaluating
                        # source formulas in Python.
                        for name, renamed in sorted(old_to_new.items(),key=lambda x:-len(x[0])):
                            value=value.replace(
                                "'" + name.replace("'","''") + "'!",
                                "'" + renamed.replace("'","''") + "'!"
                            )
                            value=re.sub(
                                rf"(?<![\w']){re.escape(name)}!",
                                "'" + renamed.replace("'","''") + "'!",
                                value
                            )
                    else:
                        value=_safe_excel_text(value)
                    dst=dest.cell(src.row,src.column,value)
                    if src.has_style:
                        dst.font=copy(src.font)
                        dst.fill=copy(src.fill)
                        dst.border=copy(src.border)
                        dst.alignment=copy(src.alignment)
                        dst.number_format=src.number_format
                        dst.protection=copy(src.protection)
                    if src.hyperlink:
                        dst.hyperlink=copy(src.hyperlink)
            # Header positions are sometimes below row one. Preserve
            # original merged cells, row heights, column widths and filters.
            for merged in original.merged_cells.ranges:
                dest.merge_cells(str(merged))
            for key,dim in original.column_dimensions.items():
                dest.column_dimensions[key].width=dim.width
                dest.column_dimensions[key].hidden=dim.hidden
            for index,dim in original.row_dimensions.items():
                if dim.height:
                    dest.row_dimensions[index].height=dim.height
            if original.auto_filter.ref:
                dest.auto_filter.ref=original.auto_filter.ref
            # Data validation/conditional formats/tables are retained in
            # original workbooks, avoiding unsafe cross-workbook references.
    finally:
        wb.close()
    return old_to_new


def _surveyor_rollup(line_result, transformer_result, pole_result):
    records=defaultdict(lambda: {
        "lv_checked":0,"lv_outstanding":0,"dt_checked":0,"dt_gaps":0,
        "pole_checked":0,"pole_gaps":0
    })
    for item in line_result.get("outstanding_lv_by_user", []):
        key=str(item.get("username") or "")
        if not key:
            continue
        # May have difference in shape across watcher versions;
        # detailed outstanding totals are derived from watcher report.
        for possible in ("affected_lines","outstanding_lines","with_gaps","lines_with_gaps"):
            if possible in item:
                records[key]["lv_outstanding"]=int(item[possible] or 0)
                break
    for item in line_result.get("user_summary", []):
        key=str(item.get("username") or "")
        if not key:
            continue
        records[key]["lv_checked"]=int(item.get("total_lines") or 0)
        records[key]["lv_outstanding"]=max(
            records[key]["lv_outstanding"],
            int(item.get("affected_lines") or item.get("outstanding_lines") or item.get("priority_gap_lines") or 0),
        )
    for source,prefix in ((transformer_result,"dt"),(pole_result,"pole")):
        for item in source.get("per_surveyor",[]):
            key=str(item.get("username") or "")
            if not key:
                continue
            records[key][f"{prefix}_checked"]=int(item.get("total") or 0)
            records[key][f"{prefix}_gaps"]=int(item.get("with_gaps") or 0)
    return records


def _build_consolidated_workbook(destination, reports, context, geometry):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb=Workbook()
    overview=wb.active
    overview.title="Combined Summary"
    overview.sheet_view.showGridLines=False
    overview.append(["REG LV DISTRIBUTION ASSET QUALITY — CONSOLIDATED"])
    overview.merge_cells("A1:G1")
    overview["A1"].font=Font(color="FFFFFF",bold=True,size=15)
    overview["A1"].fill=PatternFill("solid",fgColor=HIGHLIGHT)
    overview.row_dimensions[1].height=32
    overview.append(["Starting district",context["district"]])
    overview.append(["Selected accounts","; ".join(context["usernames"]) if context["usernames"] else "All surveyors"])
    overview.append(["From",context["date_start"]])
    overview.append(["To",context["date_end"]])
    overview.append(["Date rule","Created or last edited during review period"])
    overview.append(["QA boundary","Original QA features inside district; no assertion of downstream feeder source"])
    overview.append(["Continuity rule","Optional, separately labeled, provisional geometry beyond district borders"])
    overview.append(["Read-only","No source feature or network topology modified"])
    overview.append(["Live source validation","Use source workbooks and review evidence; geometry is not an energized feeder proof"])
    overview.append([])
    overview.append([
        "Asset class","Records examined","Original outstanding features",
        "Original issues","Duplicate GlobalID groups",
        "Further verification","Source report",
    ])
    lv=reports["lv_line"]
    dt=reports["transformer"]
    pole=reports["pole"]
    rows=[
        ("LV Lines",lv.get("total_lines",0),lv.get("outstanding_lv_lines",0),
         lv.get("outstanding_lv_issues",0), lv.get("duplicate_globalid_groups",0), 
         f"{lv.get('flagged_lines',0)} lines flagged by threshold",
         Path(lv["output_xlsx"]).name),
        ("Distribution Transformers",dt.get("total",0),
         dt.get("original_attribute_gap_records",dt.get("with_gaps",0)),
         dt.get("original_attribute_issue_count",dt.get("issue_count",0)),
         dt.get("duplicate_globalid_groups",0),
         f"{dt.get('not_verified',0)} records not fully verified",
         Path(dt["output_xlsx"]).name),
        ("LV Poles",pole.get("total",0),
         pole.get("original_attribute_gap_records",pole.get("with_gaps",0)),
         pole.get("original_attribute_issue_count",pole.get("issue_count",0)),
         pole.get("duplicate_globalid_groups",0),
         f"{pole.get('not_verified',0)} records not fully verified",
         Path(pole["output_xlsx"]).name),
    ]
    for row in rows:overview.append(list(row))
    overview.append([])
    overview.append(["ADDITIONAL LV SCORED-FIELD GAPS — SEPARATE FROM ORIGINAL PRIORITY COUNTS"])
    overview.append(["LV line records with missing/invalid scored fields",
                     lv.get("scored_field_gap_records", 0)])
    overview.append(["LV scored-field missing entries",
                     lv.get("scored_missing_entries", 0)])
    overview.append(["LV scored-field invalid entries",
                     lv.get("scored_invalid_entries", 0)])
    overview.append([])
    overview.append(["CROSS-DISTRICT PROVISIONAL GEOMETRY (SEPARATE FROM QA TOTALS)"])
    if geometry.get("status")=="COMPLETE":
        for k,label in [
            ("starting_transformers","Starting district transformers"),
            ("nationwide_candidate_lv_lines","LV candidate geometries for selected accounts (nationwide)"),
            ("connected_lv_line_ids","Provisional same-creator lines associated with district starts"),
            ("unassociated_lines","Candidate lines not associated with selected transformer starts"),
            ("suggested_gap_count","Suggested endpoint gaps (not auto-connected)"),
        ]:
            overview.append([label,geometry.get(k,0)])
        overview.append(["Geometry status","PROVISIONAL ONLY — not electrical tracing"])
    else:
        overview.append(["Geometry status",geometry.get("status","NOT_RUN")])
        overview.append(["Reason",geometry.get("reason","Not requested")])
    survey=wb.create_sheet("Surveyor Comparison")
    survey.append(["Username","LV Lines Checked","LV Priority Gaps",
                   "Transformers Checked","Transformer Gaps",
                   "LV Poles Checked","Pole Gaps"])
    rollup=_surveyor_rollup(lv,dt,pole)
    for username,stats in sorted(rollup.items(),key=lambda i:i[0].casefold()):
        survey.append([username,stats["lv_checked"],stats["lv_outstanding"],
                       stats["dt_checked"],stats["dt_gaps"],
                       stats["pole_checked"],stats["pole_gaps"]])
    notes=wb.create_sheet("Reporting Rules")
    for a,b in [
        ("Topic","Requirement and limitation"),
        ("Three original audits","All existing LV line, distribution-transformer and LV pole rules run unchanged."),
        ("LV gaps","Priority gap count is different from completion threshold FLAG count."),
        ("Additional scored-field gaps","Missing mandatory/scored LV attributes and invalid scored coded values are shown separately even when the completion threshold is met; supplying lines and the two years are informational."),
        ("Pole height not recorded","Reported as informational. Current mandatory height rule only evaluates recorded values exceeding 12; requirements must be confirmed before treating missing height as noncompliance."),
        ("LV duplicate GlobalIDs","All scoped source OBJECTIDs retained in separate duplicate registers; no geometry comparison or source deletion. Identifier conflicts are outside original attribute-gap totals."),
        ("Distribution transformers","Seven original gap categories; duplicate source GlobalID recorded separately."),
        ("Poles","Manufacturer missing is only a correction gap when Notes do not document an absent nameplate."),
        ("Date scope","Same created/edited date range and surveyor filter for all three asset QA checks."),
        ("District scope","For main QA, features intersecting district polygons."),
        ("Cross-border lines","Continuity supplement scans national created_user line geometries without a district clip."),
        ("Inference","Same creator + matching endpoints + transformer proximity is tentative only."),
        ("Geometry availability","Continuity failures do not invalidate completed independent QA results, but are prominently marked."),
        ("Prior account holders","For transferred portal accounts, historical personal authorship cannot be established from current labels alone."),
        ("Original reports","Source Excel workbooks remain in the Components folder and are not altered."),
        ("GIS safety","All service reads; no Utility Network edits."),
    ]: notes.append([a,b])
    contents=wb.create_sheet("Sheet Index")
    contents.append(["Source component","Original worksheet","Consolidated worksheet"])
    source_paths={
        "lv_line":reports["lv_line"]["output_xlsx"],
        "transformer":reports["transformer"]["output_xlsx"],
        "pole":reports["pole"]["output_xlsx"],
    }
    if geometry.get("status")=="COMPLETE":
        source_paths["geometry"]=geometry["output_xlsx"]
    for source,path in source_paths.items():
        names=_copy_source_workbook(wb,path,SHEET_PREFIXES[source])
        for old,new in names.items():
            contents.append([source,old,new])
    for sheet in (overview,survey,notes,contents):
        sheet.freeze_panes="A2" if sheet is not overview else "A12"
        sheet.sheet_view.showGridLines=False
        for idx,cell in enumerate(sheet[1],1):
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor=HIGHLIGHT)
        if sheet is overview:
            for cell in sheet[12]:
                cell.font=Font(bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor=HIGHLIGHT)
        widths={"A":38,"B":24,"C":28,"D":22,"E":24,"F":38,"G":46}
        for col,width in widths.items():
            sheet.column_dimensions[col].width=width
        if sheet is notes:
            sheet.column_dimensions["A"].width=30
            sheet.column_dimensions["B"].width=110
            for row in sheet.iter_rows(min_row=2,min_col=2,max_col=2):
                row[0].alignment=Alignment(wrap_text=True,vertical="top")
        if sheet is contents:
            sheet.column_dimensions["A"].width=22
            sheet.column_dimensions["B"].width=36
            sheet.column_dimensions["C"].width=40
    if destination.exists():
        raise FileExistsError(f"Consolidated report already exists: {destination}")
    wb.save(destination)
    wb.close()


def run_combined_qa_report(
    output_folder, district, usernames=None,
    start_date="2026-01-01",end_date=None,
    include_continuity=True,service_url=None,boundary_service_url=None,
    version="sde.DEFAULT",gis=None,progress=None,endpoint_tolerance=1.5,
):
    """Build one master workbook and retain the original QA reports.

    Every original audit is mandatory. Provisional connectivity is optional
    and its failure cannot quietly produce a 'verified' coverage assertion.
    """
    start=_as_date(start_date)
    end=_as_date(end_date or date.today())
    if end<start:
        raise ValueError("End date is earlier than start date.")
    if not district or not str(district).strip():
        raise ValueError("Report district is required.")
    users=sorted(set(str(u).strip() for u in (usernames or []) if str(u).strip()),
                 key=str.casefold)
    target=Path(output_folder).expanduser().resolve()
    if not target.is_dir():
        raise ValueError(f"Choose an existing destination folder: {target}")
    clean=re.sub(r"[^a-zA-Z0-9_-]+","_",str(district)).strip("_")
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_dir=target/f"REG_Combined_QA_{clean}_{stamp}"
    run_dir.mkdir(parents=False,exist_ok=False)
    component_dir=run_dir/"Components"
    component_dir.mkdir()
    manifest=run_dir/"Run_Manifest.txt"
    outcome={"status":"IN_PROGRESS","completed":[],"components":{}}
    def notify(message):
        if progress: progress(message)
    def checkpoint():
        manifest.write_text(
            "Status: "+outcome["status"]+"\n"+
            "QA scope: "+str(district)+" | "+start.isoformat()+" to "+end.isoformat()+"\n"+
            "Surveyor filter: "+("; ".join(users) if users else "ALL")+"\n"+
            "Completed components: "+", ".join(outcome["completed"])+"\n"+
            "Outputs: "+str(outcome["components"])+"\n",
            encoding="utf-8"
        )
    checkpoint()
    kwargs={
        "start_date":start,"end_date":end,
        "district":district,"usernames":users or None,
        "service_url":service_url,"boundary_service_url":boundary_service_url,
        "version":version,"gis":gis,
    }
    reports={}
    for key in ("lv_line","transformer","pole"):
        notify(f"Running original {key} QA for {district}, preserving its existing gap rules.")
        try:
            if key=="lv_line":
                result=run_quality_watcher(
                    output_xlsx=component_dir/"LV_Lines_QA.xlsx",
                    output_flagged_feature_class=None,**kwargs
                )
            else:
                result=run_asset_quality(
                    asset=key,output_xlsx=component_dir/f"{key}_QA.xlsx",**kwargs
                )
            reports[key]=result
            outcome["completed"].append(key)
            outcome["components"][key]=result["output_xlsx"]
            checkpoint()
        except Exception as exc:
            outcome["status"]="INCOMPLETE"
            outcome["failure"]=f"{key}: {exc}"
            checkpoint()
            raise RuntimeError(
                f"Consolidated QA failed during {key}: {exc}. "
                f"Completed original reports were kept under {component_dir}; "
                "a consolidated workbook was NOT claimed complete."
            ) from exc
    geometry={"status":"NOT_REQUESTED"}
    if include_continuity:
        if not users:
            geometry={"status":"SKIPPED_UNSCOPED",
                      "reason":"Select a surveyor to scope nationwide geometric continuity; full QA is unaffected."}
        else:
            try:
                notify(
                    "Running separately labeled, provisional cross-district "
                    "continuity. QA totals will NOT be replaced."
                )
                geometry=run_lv_feeder_report(
                    output_folder=component_dir,district=district,
                    usernames=users,service_url=service_url,
                    boundary_service_url=boundary_service_url,
                    version=version,gis=gis,progress=progress,
                    method="geometry",endpoint_tolerance=endpoint_tolerance,
                )
                geometry["status"]="COMPLETE"
                outcome["components"]["geometry"]=geometry["output_xlsx"]
                checkpoint()
            except Exception as exc:
                geometry={"status":"UNAVAILABLE","reason":str(exc)}
                notify(
                    f"WARNING: Provisional continuity unavailable ({exc}). "
                    "Original LV, transformer and pole QA remain complete."
                )
    output=run_dir/f"REG_Combined_LV_Quality_{clean}_{stamp}.xlsx"
    try:
        _build_consolidated_workbook(
            output,reports,{
                "district":district,"usernames":users,
                "date_start":start.isoformat(),"date_end":end.isoformat(),
            },geometry
        )
    except Exception as exc:
        outcome["status"]="INCOMPLETE"
        outcome["failure"]="Merge: "+str(exc)
        checkpoint()
        raise
    outcome["status"]="COMPLETE"
    outcome["combined_report"]=str(output)
    outcome["geometry"]=geometry.get("status")
    checkpoint()
    return {
        "status":"COMPLETE","read_only":True,
        "output_xlsx":str(output),"components_folder":str(component_dir),
        "manifest":str(manifest),"district":str(district),
        "start_date":start.isoformat(),"end_date":end.isoformat(),
        "surveyor_filter":users,"asset_reports":reports,
        "geometry":geometry,
    }
