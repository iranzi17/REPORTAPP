"""Network-wide read-only audit of repeated Electric Line GLOBALIDs.

Scans every Electric Line record using indexed OBJECTID keyset pagination,
without date/surveyor restrictions, then fetches only duplicate records'
geometry and intersects the official national cell boundary service.
No feature or topology edits are performed.
"""
from __future__ import annotations

import os
from collections import defaultdict, Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from .duplicate_line_inspector import _canonical_paths, canonical_globalid
from .progress_report import (
    DEFAULT_LINE_SERVICE_URL, DEFAULT_BOUNDARY_SERVICE_URL, LINE_LAYER_ID,
    CELL_LAYER_ID, _get_gis, _layer_url, _query_with_retry,
)


def _progress(message):
    try:
        import arcpy
        arcpy.AddMessage(message)
    except ImportError:
        pass


def collect_duplicate_ids(layer, version, where="1=1", page_size=500, progress=None):
    """Find repeated GlobalIDs across ALL matching OIDs; fail if scan incomplete.

    OBJECTID keyset paging avoids a huge return_ids_only response and the
    FeatureServer's maximum record count, which can cause 503 timeouts.
    Records must be ordered ASC by OBJECTID for complete pagination.
    """
    if page_size < 1:
        raise ValueError("page_size must be positive")
    report = progress or (lambda text: None)
    try:
        expected = int(_query_with_retry(
            layer, attempts=2, where=where, return_count_only=True,
            gdb_version=version,
        ))
    except (ValueError, TypeError) as exc:
        raise RuntimeError("Unable to read the Electric Line record count; scan aborted.") from exc
    if expected < 0:
        raise RuntimeError("FeatureServer returned negative record count")
    seen: dict[str, int] = {}
    duplicates: dict[str, list[int]] = {}
    skipped = 0
    scanned = 0
    last_id = -1
    while True:
        query_where = f"({where}) AND OBJECTID > {last_id}"
        response = _query_with_retry(
            layer, attempts=2, where=query_where,
            out_fields="OBJECTID,GLOBALID", return_geometry=False,
            order_by_fields="OBJECTID ASC",
            result_record_count=page_size, return_all_records=False,
            gdb_version=version,
        )
        if not hasattr(response, "features"):
            raise RuntimeError("Incomplete Electric Line page: no features returned.")
        if not response.features:
            break
        batch_max = last_id
        for feature in response.features:
            attrs = dict(feature.attributes)
            if attrs.get("OBJECTID") is None:
                raise RuntimeError("Electric Line query returned a missing OBJECTID")
            oid = int(attrs["OBJECTID"])
            if oid <= last_id or oid <= batch_max:
                raise RuntimeError(
                    "FeatureServer failed ordered OBJECTID pagination; "
                    "cannot guarantee a complete duplicate scan."
                )
            batch_max = oid
            try:
                gid = canonical_globalid(attrs.get("GLOBALID"))
            except ValueError:
                skipped += 1
                continue
            if gid in duplicates:
                duplicates[gid].append(oid)
            elif gid in seen:
                duplicates[gid] = [seen.pop(gid), oid]
            else:
                seen[gid] = oid
        scanned += len(response.features)
        last_id = batch_max
        if scanned == len(response.features) or scanned % 10000 < len(response.features):
            report(f"Scanned {scanned:,}/{expected:,} Electric Line records; "
                   f"{len(duplicates):,} duplicate GlobalID groups so far.")
        if scanned > expected:
            raise RuntimeError(
                f"Source dataset changed while scanning: scanned {scanned}, "
                f"initial count {expected}. No incomplete report created."
            )

    ending_count = int(_query_with_retry(
        layer, attempts=2, where=where, return_count_only=True,
        gdb_version=version,
    ))
    if scanned != expected or ending_count != expected:
        raise RuntimeError(
            f"FeatureServer count changed or rows were skipped: initial "
            f"{expected}, scanned {scanned}, final {ending_count}. "
            "Repeat after edits stabilize; no partial workbook was created."
        )
    report(f"Completed {scanned:,} records; {len(duplicates):,} duplicate groups, "
           f"{sum(map(len, duplicates.values())):,} affected records.")
    return duplicates, {"scanned": scanned, "invalid_or_missing_globalids": skipped}


def fetch_duplicate_features(layer, groups, version, batch_size=120, progress=None):
    """Fetch full source attributes/geometry only for duplicate OBJECTIDs."""
    ids = sorted({oid for group in groups.values() for oid in group})
    output = {}
    report = progress or (lambda text: None)
    if not ids:
        return output
    for offset in range(0, len(ids), batch_size):
        chunk = ids[offset:offset + batch_size]
        response = _query_with_retry(
            layer, attempts=2, where="1=1",
            object_ids=",".join(map(str, chunk)), out_fields="*",
            return_geometry=True, return_all_records=False,
            gdb_version=version,
        )
        if not hasattr(response, "features") or getattr(response, "exceeded_transfer_limit", False):
            raise RuntimeError("Incomplete duplicate geometry batch; aborting report.")
        sr = dict(getattr(response, "spatial_reference", None) or {})
        for feature in response.features:
            attrs = dict(feature.attributes)
            oid = int(attrs["OBJECTID"])
            if oid not in chunk or oid in output:
                raise RuntimeError(f"Unexpected/repeated Electric Line OBJECTID {oid} in geometry fetch.")
            geometry = dict(feature.geometry or {})
            if sr and not geometry.get("spatialReference"):
                geometry["spatialReference"] = sr
            output[oid] = {"attributes": attrs, "geometry": geometry}
        missing = set(chunk) - output.keys()
        if missing:
            raise RuntimeError(f"Could not fetch {len(missing)} duplicate Electric Line geometries.")
        if offset == 0 or offset + batch_size >= len(ids) or (offset//batch_size+1)%20 == 0:
            report(f"Loaded detailed geometries for {len(output):,}/{len(ids):,} duplicate features.")
    for gid, objects in groups.items():
        for oid in objects:
            actual = canonical_globalid(output[oid]["attributes"].get("GLOBALID"))
            if actual != gid:
                raise RuntimeError(f"Duplicate feature {oid} changed GlobalID during scan.")
    return output


def assign_admin_areas(cell_layer, duplicate_features, progress=None):
    """One true geometry-intersection lookup per distinct polyline geometry."""
    from arcgis.geometry.filters import intersects

    report = progress or (lambda text: None)
    cache = {}
    by_oid = {}
    for ix, (oid, record) in enumerate(duplicate_features.items(), 1):
        geom = record["geometry"]
        signature = _canonical_paths(geom)
        sr = geom.get("spatialReference") or {}
        key = (signature, repr(sorted(sr.items()))) if signature is not None else ("missing", oid)
        if key in cache:
            by_oid[oid] = cache[key]
            continue
        if signature is None or not sr:
            area = {"district": [], "sector": [], "cell": [],
                    "area_paths": [], "status": "GEOMETRY_OR_CRS_MISSING"}
        else:
            response = _query_with_retry(
                cell_layer, attempts=2, where="1=1",
                out_fields="district,sector,cell",
                geometry_filter=intersects(geom, sr=sr),
                return_geometry=False, return_all_records=True,
            )
            if not hasattr(response, "features") or getattr(response, "exceeded_transfer_limit", False):
                raise RuntimeError(
                    f"Administrative boundary lookup incomplete for Electric Line {oid}."
                )
            unique = set()
            for feature in response.features:
                attrs = dict(feature.attributes)
                parts = tuple(str(attrs.get(field) or "").strip()
                              for field in ("district", "sector", "cell"))
                if any(parts):
                    unique.add(parts)
            paths = sorted(unique, key=lambda values: tuple(s.casefold() for s in values))
            area = {
                "district": sorted({p[0] for p in paths if p[0]}, key=str.casefold),
                "sector": sorted({p[1] for p in paths if p[1]}, key=str.casefold),
                "cell": sorted({p[2] for p in paths if p[2]}, key=str.casefold),
                "area_paths": [" / ".join(p) for p in paths],
                "status": "MULTIPLE_CELLS" if len(paths) > 1 else
                          "MATCH" if paths else "NOT_FOUND",
            }
        cache[key] = area
        by_oid[oid] = area
        if ix == 1 or ix % 25 == 0 or ix == len(duplicate_features):
            report(f"Located {ix:,}/{len(duplicate_features):,} duplicate features "
                   f"against administrative cells ({len(cache):,} distinct geometries).")
    return by_oid


def _create_report(path, groups, records, areas, stats, source_url, boundary_url, version, classes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    group_sheet = wb.create_sheet("Duplicate GlobalIDs")
    details = wb.create_sheet("All Duplicate Records")
    area_summary = wb.create_sheet("By District Sector Cell")
    notes = wb.create_sheet("Method and Warnings")

    num_features = sum(len(x) for x in groups.values())
    overview = [
        ("Measure", "Value"), ("Scope", classes),
        ("FeatureServer", source_url), ("Admin boundary source", boundary_url),
        ("Version", version), ("Total electric lines scanned", stats["scanned"]),
        ("Duplicate GlobalID groups", len(groups)),
        ("Affected feature records", num_features),
        ("Additional copies beyond one per GlobalID", num_features - len(groups)),
        ("Missing/invalid GlobalIDs (not grouped)", stats["invalid_or_missing_globalids"]),
        ("Administrative locations", "Spatial intersection; more than one cell possible"),
        ("Data modification", "None: read-only diagnostic"),
    ]
    for row in overview:
        summary.append(row)
    group_sheet.append([
        "GlobalID", "Number of Records", "OBJECTIDs", "Same Exact XY Geometry",
        "District(s)", "Sector(s)", "Cell(s)", "Administrative Locations",
        "Admin Status",
    ])
    details.append([
        "GlobalID", "OBJECTID", "Other OBJECTIDs With Same GlobalID",
        "Duplicate Group Size", "District(s)", "Sector(s)", "Cell(s)",
        "District / Sector / Cell", "Admin Status",
        "Asset Group", "Asset Type", "SHAPE Length",
        "Created By", "Created Date", "Last Edited By", "Last Edited Date",
        "Transformer Name", "Same Exact XY Geometry",
    ])
    area_summary.append([
        "District", "Sector", "Cell", "Duplicate GlobalID Groups",
        "Duplicate Electric Line Records",
    ])
    geo_aggregate = defaultdict(lambda: [set(),set()])
    for gid, objects in sorted(groups.items()):
        recs = [records[oid] for oid in objects]
        signatures = [_canonical_paths(item["geometry"]) for item in recs]
        same_geom = bool(signatures) and all(x is not None for x in signatures) and len(set(signatures)) == 1
        area_paths = sorted({
            p for oid in objects for p in areas[oid]["area_paths"]
        }, key=str.casefold)
        all_districts = sorted({v for oid in objects for v in areas[oid]["district"]}, key=str.casefold)
        all_sectors = sorted({v for oid in objects for v in areas[oid]["sector"]}, key=str.casefold)
        all_cells = sorted({v for oid in objects for v in areas[oid]["cell"]}, key=str.casefold)
        statuses = sorted({areas[oid]["status"] for oid in objects})
        group_sheet.append([
            gid, len(objects), "; ".join(map(str, objects)),
            "YES" if same_geom else "NO / NEEDS SPATIAL REVIEW",
            "; ".join(all_districts), "; ".join(all_sectors),
            "; ".join(all_cells), "; ".join(area_paths),
            "; ".join(statuses),
        ])
        for oid in sorted(objects):
            row = records[oid]["attributes"]
            area = areas[oid]
            details.append([
                gid, oid, "; ".join(str(v) for v in objects if v != oid),
                len(objects), "; ".join(area["district"]),
                "; ".join(area["sector"]), "; ".join(area["cell"]),
                "; ".join(area["area_paths"]), area["status"],
                row.get("ASSETGROUP"), row.get("ASSETTYPE"),
                row.get("SHAPE__Length", row.get("SHAPE_Length", row.get("Shape__Length", ""))),
                row.get("created_user"), row.get("created_date"),
                row.get("last_edited_user"), row.get("last_edited_date"),
                row.get("Transformer_Name"), "YES" if same_geom else "NO / NEEDS SPATIAL REVIEW",
            ])
            for path_text in area["area_paths"]:
                components = tuple(path_text.split(" / ", 2))
                if len(components) == 3:
                    geo_aggregate[components][0].add(gid)
                    geo_aggregate[components][1].add(oid)

    for path_tuple, (unique_gids, objectids) in sorted(geo_aggregate.items()):
        area_summary.append([*path_tuple, len(unique_gids), len(objectids)])

    for row in (
        ("Rule", "Interpretation"),
        ("Duplicate GlobalID", "Two or more distinct Electric Line OBJECTIDs with the same normalized GUID."),
        ("All data included", "No survey dates or surveyors limit the all-network scan."),
        ("Exact geometry", "Identical XY vertex coordinates, allowing reversed direction and multipart ordering."),
        ("Same length only", "Matching lengths alone never establish duplicate geometry."),
        ("Administrative location", "All intersected national Cell polygons. Lines crossing borders may appear in several cells."),
        ("Group count by cell", "A duplicate GlobalID may appear in more than one administrative cell."),
        ("Unknown location", "NOT_FOUND or GEOMETRY_OR_CRS_MISSING must not be interpreted as a district."),
        ("Data changes", "If the number of source records changed during the scan, the tool aborts instead of publishing partial counts."),
        ("Next review", "Inspect connectivity, Utility Network associations and attachments; do not delete or rename automatically."),
        ("Dirty areas", "Duplicate GlobalIDs and dirty topology may coexist but causation is not established."),
        ("Version consistency", "The Electric Line is queried in the requested version; the hosted boundary service is not versioned."),
    ):
        notes.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
        for column in sheet.columns:
            cells = list(column)
            sheet.column_dimensions[cells[0].column_letter].width = min(
                85, max(15, max(len(str(c.value or "")) for c in cells[:1000]) + 2)
            )
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    for sheet, table_name in (
        (group_sheet,"DuplicateGlobalIDGroups"), (details,"AllDuplicateElectricLines"),
        (area_summary,"DuplicateCellTotals")
    ):
        if sheet.max_row > 1:
            tab = Table(
                displayName=table_name,
                ref=f"A1:{sheet.cell(1,sheet.max_column).column_letter}{sheet.max_row}",
            )
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(tab)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(f"Output already exists: {path}")
    wb.save(path)


def scan_electric_line_duplicate_globalids(
    output_xlsx: str | Path,
    classes: str = "all",
    service_url: str | None = None,
    boundary_service_url: str | None = None,
    version: str = "sde.DEFAULT",
    gis: Any | None = None,
    page_size: int = 500,
) -> dict:
    """Write a comprehensive source duplicate register with administrative areas."""
    if classes not in ("all", "lv"):
        raise ValueError("Scan classes must be all or lv.")
    output = Path(output_xlsx).expanduser().resolve()
    if output.suffix.casefold() != ".xlsx":
        raise ValueError("Output must be an Excel .xlsx file.")
    output = output.with_name(f"{output.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    active_gis = _get_gis(gis)
    from arcgis.features import FeatureLayer
    service = service_url or os.getenv("UN_REPORT_SERVICE_URL", DEFAULT_LINE_SERVICE_URL)
    boundary = boundary_service_url or os.getenv("UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL)
    electric_line = FeatureLayer(_layer_url(service, LINE_LAYER_ID), active_gis)
    cells = FeatureLayer(_layer_url(boundary, CELL_LAYER_ID), active_gis)
    scan_where = "1=1" if classes == "all" else "ASSETGROUP IN (9,10)"
    _progress("Scanning all Electric Line GlobalIDs using OBJECTID paging. "
              "Large country-wide datasets may take several minutes.")
    groups, stats = collect_duplicate_ids(
        electric_line, version, where=scan_where, page_size=page_size, progress=_progress,
    )
    records = fetch_duplicate_features(
        electric_line, groups, version, progress=_progress,
    )
    areas = assign_admin_areas(cells, records, progress=_progress)
    _create_report(
        output, groups, records, areas, stats, _layer_url(service, LINE_LAYER_ID),
        _layer_url(boundary, CELL_LAYER_ID), version, classes,
    )
    return {
        "read_only": True, "output_xlsx": str(output),
        "scope": classes, "version": version, **stats,
        "duplicate_globalid_groups": len(groups),
        "duplicate_feature_records": sum(map(len, groups.values())),
        "additional_copies": sum(len(v)-1 for v in groups.values()),
        "records_without_admin_location": sum(
            1 for x in areas.values() if x["status"] in (
                "NOT_FOUND","GEOMETRY_OR_CRS_MISSING",
            )
        ),
    }
