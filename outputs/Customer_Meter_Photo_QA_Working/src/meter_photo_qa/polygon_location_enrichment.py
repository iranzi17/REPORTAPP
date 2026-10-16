"""Read-only polygon location enrichment for existing LV Pole / transformer QA Excel.

Copies the workbook, never edits the original workbook or the REG FeatureServer.
Location columns are spreadsheet values only; each inferred value is audited.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import json
import os
from pathlib import Path

LOCATION_FIELDS=("Recorded District","Recorded Sector","Recorded Cell")
PROVENANCE="Location Provenance"
AUDIT_SHEET="Polygon Location Audit"
MAX_RECORDS=5000


def _text(value):
    return str(value or "").strip()


def _safe_text(value):
    s=_text(value)
    # Excel text from GIS services is untrusted.
    if s and s[0] in ("=","+","-","@"):
        return "'"+s
    return s


def _gid(value):
    return _text(value).strip("{}").casefold()


def _headers(sheet):
    return {str(cell.value).strip():cell.column for cell in sheet[1] if cell.value is not None}


def _prepare(workbook):
    if AUDIT_SHEET in workbook.sheetnames:
        raise ValueError("This report was already polygon-enriched; run the tool on the original report.")
    if "Asset Details" not in workbook.sheetnames or "Summary" not in workbook.sheetnames:
        raise ValueError("Expected an original REPORTAPP asset QA workbook with Asset Details and Summary.")
    summary=workbook["Summary"]
    meta={str(row[0].value):row[1].value
          for row in summary.iter_rows(min_row=1,max_col=2)
          if row[0].value is not None}
    group=int(meta.get("Asset Group") or 0)
    if group not in (7,12):
        raise ValueError("This tool is for REG LV Pole (group 7) and distribution transformer (group 12) point assets only.")
    details=workbook["Asset Details"]
    heads=_headers(details)
    mandatory=("OBJECTID","GlobalID","Status",*LOCATION_FIELDS)
    missing=[x for x in mandatory if x not in heads]
    if missing:
        raise ValueError("REPORTAPP asset QA workbook is missing columns: "+", ".join(missing))
    entries=[]
    for row in range(2,details.max_row+1):
        oid=details.cell(row,heads["OBJECTID"]).value
        if oid is None:
            continue
        if isinstance(oid,bool) or not str(oid).strip().isdigit() or int(oid)<=0:
            raise ValueError("Invalid source OBJECTID in report row "+str(row))
        original=tuple(_text(details.cell(row,heads[col]).value) for col in LOCATION_FIELDS)
        if any(not item for item in original):
            entries.append({"row":row,"objectid":int(oid),
                            "globalid":_gid(details.cell(row,heads["GlobalID"]).value),
                            "original":original})
    if len(entries)>MAX_RECORDS:
        raise ValueError("More than 5,000 incomplete features; use a scoped report first.")
    return details,group,heads,entries


def _make_live_resolver(entries,group,*,gis=None,service_url=None,boundary_url=None,version="sde.DEFAULT"):
    from arcgis.features import FeatureLayer
    from arcgis.geometry.filters import intersects
    from .progress_report import (
        DEFAULT_LINE_SERVICE_URL,DEFAULT_BOUNDARY_SERVICE_URL,
        _get_gis,_layer_url,CELL_LAYER_ID,
    )
    from .asset_quality import ASSETS
    active=_get_gis(gis)
    root=(service_url or os.environ.get("UN_REPORT_SERVICE_URL") or DEFAULT_LINE_SERVICE_URL).rstrip("/")
    boundaries=(boundary_url or os.environ.get("UN_SECTOR_BOUNDARY_URL") or DEFAULT_BOUNDARY_SERVICE_URL).rstrip("/")
    source=FeatureLayer(_layer_url(root,ASSETS["pole" if group==7 else "transformer"]["layer_id"]),active)
    cells=FeatureLayer(_layer_url(boundaries,CELL_LAYER_ID),active)
    # Query geometry by exact OBJECTIDs; GlobalIDs are not unique in REG data.
    point_by_oid={}
    ids=sorted({item["objectid"] for item in entries})
    for offset in range(0,len(ids),100):
        batch=ids[offset:offset+100]
        response=source.query(
            object_ids=",".join(map(str,batch)),out_fields="OBJECTID,GLOBALID,ASSETGROUP",
            return_geometry=True,gdb_version=version,out_sr=4326,
            return_all_records=True,
        )
        for item in response.features:
            attrs=dict(item.attributes or {})
            oid=attrs.get("OBJECTID")
            if oid is not None and int(oid) in batch:
                point_by_oid[int(oid)]=(attrs,dict(item.geometry or {}))
    cache={}
    def resolve(oid,expected_globalid):
        identity=(oid,expected_globalid)
        if identity in cache:
            return cache[identity]
        record=point_by_oid.get(oid)
        if record is None:
            result=("FEATURE_NOT_FOUND",None)
        else:
            attrs,geometry=record
            if int(attrs.get("ASSETGROUP") or -1)!=group:
                result=("WRONG_ASSET_GROUP",None)
            elif expected_globalid and _gid(attrs.get("GLOBALID"))!=expected_globalid:
                result=("SOURCE_GLOBALID_MISMATCH",None)
            elif "x" not in geometry or "y" not in geometry:
                result=("INVALID_POINT_GEOMETRY",None)
            else:
                point={"x":geometry["x"],"y":geometry["y"],
                       "spatialReference":geometry.get("spatialReference") or {"wkid":4326}}
                located=cells.query(
                    where="1=1",geometry_filter=intersects(
                        point,sr=point["spatialReference"]),
                    out_fields="district,sector,cell",return_geometry=False,
                    return_all_records=True,
                )
                candidates=set()
                for boundary in located.features:
                    attrs2={str(k).casefold():_text(v) for k,v in dict(boundary.attributes or {}).items()}
                    match=(attrs2.get("district",""),attrs2.get("sector",""),attrs2.get("cell",""))
                    if all(match):
                        candidates.add(match)
                if not candidates:
                    result=("NO_CELL_POLYGON_MATCH",None)
                elif len(candidates)>1:
                    result=("AMBIGUOUS_POLYGON_MATCH",None)
                else:
                    result=("POLYGON_MATCH",next(iter(candidates)))
        cache[identity]=result
        return result
    return resolve


def enrich_report(report_path,*,output_path=None,apply=False,resolver=None,
                  gis=None,service_url=None,boundary_url=None,version="sde.DEFAULT",
                  limit=None):
    from openpyxl import load_workbook
    source=Path(report_path).expanduser().resolve()
    if source.suffix.casefold()!=".xlsx" or not source.is_file():
        raise FileNotFoundError("Provide an existing REPORTAPP .xlsx quality report.")
    wb=load_workbook(source)
    try:
        details,group,heads,entries=_prepare(wb)
        total_incomplete=len(entries)
        if limit is not None:
            if isinstance(limit,bool) or not 1<=int(limit)<=MAX_RECORDS:
                raise ValueError("Limit must be an integer from 1 through 5,000.")
            entries=entries[:int(limit)]
        if resolver is None and entries:
            resolver=_make_live_resolver(
                entries,group,gis=gis,service_url=service_url,
                boundary_url=boundary_url,version=version,
            )
        outcomes=[]
        for entry in entries:
            status,location=resolver(entry["objectid"],entry["globalid"])
            previous=entry["original"]
            filled=previous
            if status=="POLYGON_MATCH":
                if not location or len(location)!=3 or not all(_text(x) for x in location):
                    raise ValueError("Polygon resolver returned incomplete administrative values.")
                location=tuple(map(_text,location))
                conflicts=[label for label,recorded,spatial in zip(LOCATION_FIELDS,previous,location)
                           if recorded and recorded.casefold()!=spatial.casefold()]
                if conflicts:
                    status="LOCATION_CONFLICT_REVIEW"
                else:
                    filled=tuple(recorded or spatial
                                 for recorded,spatial in zip(previous,location))
            outcomes.append({**entry,"status":status,"spatial":location,
                             "filled":filled})
        stats=Counter(row["status"] for row in outcomes)
        result={"status":"DRY_RUN" if not apply else "OUTPUT_CREATED",
                "source_report":str(source),"version":version,"asset_group":group,
                "rows_missing_location":total_incomplete,
                "rows_evaluated":len(entries),
                "rows_not_evaluated":total_incomplete-len(entries),
                "records_filled":sum(x["status"]=="POLYGON_MATCH" for x in outcomes),
                "cells_filled":sum(sum(not old and bool(new) for old,new in zip(x["original"],x["filled"]))
                                   for x in outcomes),
                "review_counts":dict(stats),"read_only_reg":True,"source_workbook_preserved":True}
        if not apply:
            return result
        output=(Path(output_path).expanduser().resolve() if output_path else
                source.with_name(source.stem+"_polygon_enriched_"+
                                 datetime.now().strftime("%Y%m%d_%H%M%S_%f")+".xlsx"))
        if output == source or output.exists():
            raise FileExistsError("Never overwrite the original report or an existing output.")
        # Make original source columns easy to use; keep a row-by-row audit of
        # spatially filled values and do not silently overwrite conflicts.
        col=heads.get(PROVENANCE)
        if col is None:
            col=details.max_column+1
            details.cell(1,col,PROVENANCE)
        if AUDIT_SHEET in wb.sheetnames:
            raise ValueError("This report was already enriched; use the original workbook.")
        audit=wb.create_sheet(AUDIT_SHEET)
        audit.append(["Source OBJECTID","Source GlobalID",
                      "Original District","Original Sector","Original Cell",
                      "Polygon District","Polygon Sector","Polygon Cell",
                      "Final District","Final Sector","Final Cell","Resolution"])
        for row in outcomes:
            idx=row["row"]
            if row["status"]=="POLYGON_MATCH":
                for label,value in zip(LOCATION_FIELDS,row["filled"]):
                    details.cell(idx,heads[label],_safe_text(value))
            details.cell(idx,col,row["status"])
            audit.append([row["objectid"],row["globalid"],
                          *(_safe_text(x) for x in row["original"]),
                          *(_safe_text(x) for x in (row["spatial"] or ("","",""))),
                          *(_safe_text(x) for x in row["filled"]),row["status"]])
        # Existing other summary sheets reflect original source values and
        # remain unchanged for audit. The new audit provides all enrichment.
        summary=wb["Summary"]
        summary.append(["Polygon enrichment note",
                        "Asset Details location blanks filled from verified cell polygons; original values and conflicts in Polygon Location Audit. REG data unchanged."])
        summary.append(["Polygon-enriched records",result["records_filled"]])
        summary.append(["Polygon-enriched cells",result["cells_filled"]])
        # Keep existing Excel table definitions unchanged. The audit/provenance
        # column sits immediately to the right of the original table, avoiding
        # invalid tableColumn metadata in a protected source report.
        output.parent.mkdir(parents=True,exist_ok=True)
        wb.save(output)
        result["output_xlsx"]=str(output)
        return result
    finally:
        wb.close()


def main(argv=None):
    parser=argparse.ArgumentParser(description="Fill empty REG QA Excel location cells using hosted Rwanda cell polygons; no GIS writes.")
    parser.add_argument("--report",required=True,help="Original LV Pole/Transformer asset QA workbook")
    parser.add_argument("--output",help="New output workbook path; never overwrites existing")
    parser.add_argument("--apply",action="store_true",help="Explicitly create a new report copy; default is dry run")
    parser.add_argument("--version",default="sde.DEFAULT",help="REG source geodatabase version")
    parser.add_argument("--limit",type=int,help="Pilot only the first N incomplete records; remaining rows stay unchanged")
    args=parser.parse_args(argv)
    print(json.dumps(enrich_report(args.report,output_path=args.output,apply=args.apply,
                                   version=args.version,limit=args.limit),indent=2,ensure_ascii=False))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
