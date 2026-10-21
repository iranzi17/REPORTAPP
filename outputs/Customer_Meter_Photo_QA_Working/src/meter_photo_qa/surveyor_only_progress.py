"""Manual, read-only LV progress without any Utility Network topology trace.

The report is evidence of selected source feature records, not energized
feeder attribution and not physically clipped kilometres inside polygons.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from math import isfinite
from pathlib import Path
from uuid import uuid4


def run_surveyor_only_progress(
    *, output_xlsx, start, end, usernames, district, sectors, cells,
    custom_boundary_geometry, service_url, boundary_service_url, version,
    user_name_map_path, gis=None,
):
    from . import progress_report as p
    selected=list(dict.fromkeys(str(u).strip() for u in usernames if str(u).strip()))
    if not selected:
        raise ValueError("At least one surveyor is required.")
    if end < start:
        raise ValueError("The end date cannot be earlier than the start date.")
    active=p._get_gis(gis)
    lines,sectors_layer,cells_layer=p._feature_layers(
        active,service_url,boundary_service_url,
    )
    boundary_set=(
        None if custom_boundary_geometry is not None else
        p._get_sector_features(sectors_layer,district,sectors,cells_layer,cells)
    )
    boundary=p._combined_boundary_geometry(boundary_set,custom_boundary_geometry)
    where=p._build_line_where(start,end,selected)
    records=p._query_lines(
        lines,where,version,boundary_set,custom_boundary_geometry,
    )
    lengths=defaultdict(float)
    counts=defaultdict(int)
    for oid,attrs in records.items():
        if int(attrs.get("ASSETGROUP") or -1) not in (9,10):
            raise RuntimeError(f"Non-LV record returned in progress: {oid}.")
        creator=str(attrs.get("created_user") or "").strip()
        if creator not in selected:
            raise RuntimeError(f"Unselected creator returned in progress: {oid}.")
        raw=attrs.get("Shape__Length")
        if raw is None or str(raw).strip()=="":
            raise RuntimeError(f"Missing source length for LV ObjectID {oid}.")
        length=float(raw)
        if not isfinite(length) or length<0:
            raise RuntimeError(f"Invalid source length for LV ObjectID {oid}.")
        lengths[creator]+=length
        counts[creator]+=1
    active_users=[u for u in selected if counts[u]]
    names=p._resolve_full_names(
        active,active_users,p._load_name_overrides(user_name_map_path),
    )
    rows=[{
        "username":u,"full_name":names.get(u,u),
        "length_km":lengths[u]/1000.0,"record_count":counts[u],
    } for u in active_users]
    district_names=p.selected_districts(district)
    sectors=[str(v).strip() for v in sectors or [] if str(v).strip()]
    cells=[str(v).strip() for v in cells or [] if str(v).strip()]
    title=p.build_report_title(
        start,end,district,sectors,cells,
        custom_boundary=custom_boundary_geometry is not None,
    )
    requested=Path(output_xlsx).expanduser().resolve()
    output=requested.with_name(
        requested.stem+"_SURVEYOR_NO_TRACE_"+
        datetime.now().strftime("%Y%m%d_%H%M%S_")+
        uuid4().hex[:8]+".xlsx"
    )
    staged=output.with_name(output.stem+".pending.xlsx")
    p._create_workbook(staged,title,rows)
    from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=load_workbook(staged)
    notes=book.create_sheet("Method and Warnings")
    notes.append(["SURVEYOR PROGRESS ONLY — NO NETWORK TRACE","Evidence"])
    notes_rows=[
        ("Report mode","SURVEYOR_ONLY_NO_TRACE"),
        ("Transformer feeder attribution","NOT VERIFIED; no connected trace was run"),
        ("Source data","REG ElectricLine FeatureServer ASSETGROUP IN (9,10)"),
        ("Time rule","created_date in inclusive selected date window"),
        ("Surveyor rule","created_user matches one of selected accounts"),
        ("Area rule","Intersect selected area if supplied; no spatial clipping"),
        ("Kilometre rule","FULL recorded Shape__Length for each intersecting line"),
        ("Cross-sector continuation","NOT EVALUATED; origin transformer not traced"),
        ("Recorded Transformer_Name","SOURCE LABEL ONLY; does not prove supply"),
        ("Selected districts","; ".join(district_names)),
        ("Selected sectors","; ".join(sectors)),
        ("Selected cells","; ".join(cells)),
        ("From date",start.isoformat()),
        ("To date",end.isoformat()),
        ("Source version",version),
        ("Included LV line OBJECTIDs",len(records)),
        ("Network trace requests",0),
    ]
    for entry in notes_rows:
        notes.append(entry)
    detail=book.create_sheet("Selected Source LV Lines")
    detail.append([
        "Line OBJECTID","GlobalID","Created User","Created Date",
        "Recorded Transformer Name (UNVERIFIED)","Asset Group",
        "Full Source Length (km)","Spatial Selection",
    ])
    for oid in sorted(records):
        record=records[oid]
        detail.append([
            oid,str(record.get("GLOBALID") or ""),
            str(record.get("created_user") or ""),
            p._epoch_to_datetime(record.get("created_date")),
            str(record.get("Transformer_Name") or ""),
            int(record["ASSETGROUP"]),
            float(record["Shape__Length"])/1000.0,
            "INTERSECTS AREA; NOT CLIPPED" if boundary else "ALL MATCHING RECORDS",
        ])
        detail.cell(detail.max_row,4).number_format="yyyy-mm-dd hh:mm"
        detail.cell(detail.max_row,7).number_format="0.000000"
    for sheet in (notes,detail):
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes="A2"
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="17365D")
        for col in range(1,sheet.max_column+1):
            sheet.column_dimensions[sheet.cell(1,col).column_letter].width=(
                70 if sheet==notes and col==2 else 40 if col<6 else 32
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical="top")
                if isinstance(cell.value,str) and cell.value.startswith(("=","+","-","@")):
                    cell.data_type="s"
    book.save(staged)
    book.close()
    if output.exists():
        raise FileExistsError("Surveyor-only output already exists; refusing overwrite.")
    staged.replace(output)
    return {
        "read_only":True,"report_mode":"SURVEYOR_ONLY_NO_TRACE",
        "no_topology_traces":True,"output_xlsx":str(output),
        "title":title,"from_date":start.isoformat(),
        "to_date_inclusive":end.isoformat(),
        "districts":district_names,"sectors":sectors,"cells":cells,
        "selected_users":selected,"reported_users":active_users,
        "matched_line_records":len(records),"users":rows,
        "total_km":sum(row["length_km"] for row in rows),
        "full_feeder_inventory":None,"tracking_log_sheet":"",
        "feeder_connectivity_verified":False,
        "length_scope":"FULL_FEATURE_LENGTH_FOR_INTERSECTING_LINES_NOT_CLIPPED",
        "warning":"No topology trace, no verified transformer feeder association; "
                  "full intersecting features are counted without polygon clipping.",
        "line_service_url":p._layer_url(service_url,p.LINE_LAYER_ID),
        "version":version,
    }
