"""Unvalidated geometry-based LV feeder review; no Utility Network trace call.

Source district locates transformer starting points only. Candidate LV records
are retrieved *nationally* before applying an endpoint graph, so geometrically
connected same-creator lines are not clipped by administrative boundaries.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from .asset_quality import resolve_fields
from .asset_quality_runner import _filters_for_area
from .geometry_continuity import group_line_continuity
from .progress_report import (
    DEFAULT_BOUNDARY_SERVICE_URL, DEFAULT_LINE_SERVICE_URL,
    LINE_LAYER_ID, TRANSFORMER_LAYER_ID, SECTOR_LAYER_ID, CELL_LAYER_ID,
    _get_gis, _get_sector_features, _layer_url, _query_with_retry,
)
from .quality_query_batches import (
    _objectids_from_query, _confirm_empty_id_response,
    query_attributes_in_batches,
)

WEB_MERCATOR_SR = 3857
MAX_GLOBAL_LINE_CANDIDATES = 25000


def _ensure_projected_geometry(geometry, response, kind, oid):
    if not geometry:
        raise RuntimeError(f"{kind} ObjectID {oid} lacks geometry.")
    spatial_ref = geometry.get("spatialReference") or (
        getattr(response, "spatial_reference", None) or {}
    )
    # A response converted by the service to EPSG:3857 must be explicit.
    try:
        wkid = int(spatial_ref.get("latestWkid") or spatial_ref.get("wkid") or 0)
    except (TypeError, ValueError, AttributeError):
        wkid = 0
    if wkid not in (3857, 102100, 102113):
        raise RuntimeError(
            f"{kind} ObjectID {oid} lacks verified projected Web Mercator "
            f"coordinates (received WKID {wkid}). "
            "Cannot measure metre tolerances safely; source layer must "
            "support out_sr=3857."
        )
    return dict(geometry)


def _fetch_geometry_by_ids(layer, objectids, version, fields, label,
                           progress=None, batch_size=120):
    ids=sorted(set(map(int,objectids)))
    output={}
    for offset in range(0,len(ids),batch_size):
        part=ids[offset:offset+batch_size]
        response=_query_with_retry(
            layer,attempts=2,
            where="1=1",object_ids=",".join(map(str,part)),
            out_fields=",".join(fields),return_geometry=True,
            out_sr=WEB_MERCATOR_SR,return_all_records=False,
            gdb_version=version,
        )
        if (not hasattr(response,"features") or
                getattr(response,"exceeded_transfer_limit",False)):
            raise RuntimeError(f"Incomplete {label} geometry batch; report cancelled.")
        expected=set(part)
        local=set()
        for item in response.features:
            attrs=dict(item.attributes)
            oid=int(attrs["OBJECTID"])
            if oid not in expected or oid in output or oid in local:
                raise RuntimeError(f"Unexpected or repeated {label} ObjectID {oid}.")
            local.add(oid)
            output[oid]={
                "attributes":attrs,
                "geometry":_ensure_projected_geometry(
                    dict(item.geometry or {}),response,label,oid
                ),
            }
        if local!=expected:
            raise RuntimeError(
                f"{label} service omitted {len(expected-local)} source "
                "ObjectIDs; refusing to publish a partial report."
            )
        if progress and (offset == 0 or offset +batch_size>=len(ids)
                         or offset//batch_size%20 == 0):
            progress(f"Loaded geometry for {len(output):,}/{len(ids):,} {label} features.")
    return [output[k] for k in sorted(output)]


def _get_nationwide_line_ids(layer, usernames, version, max_candidates):
    # Created-user comparison is applied to the national candidate pool,
    # never spatially clips it to the district where source transformers sit.
    users=[str(u).strip() for u in (usernames or []) if str(u).strip()]
    where="ASSETGROUP IN (9,10)"
    if users:
        quoted=",".join("'"+u.replace("'","''")+"'" for u in sorted(set(users)))
        where+=" AND created_user IN ("+quoted+")"
    options={
        "where":where,"return_count_only":True,
        "gdb_version":version,
    }
    count=_query_with_retry(layer,attempts=2,**options)
    if isinstance(count,dict):
        count=count.get("count")
    if isinstance(count,bool):
        raise RuntimeError("Service returned an invalid candidate count.")
    try:
        total=int(count)
    except (ValueError,TypeError) as exc:
        raise RuntimeError("Cannot establish nationwide candidate count.") from exc
    if total<0:
        raise RuntimeError("Invalid nationwide candidate count.")
    if total>max_candidates:
        raise RuntimeError(
            f"{total:,} LV line candidates across Rwanda exceed the "
            f"{max_candidates:,} safe geometry-processing limit. "
            "Select a surveyor (full name or username), or request "
            "a smaller dataset. No incomplete report generated."
        )
    id_options={
        "where":where,"return_ids_only":True,
        "return_geometry":False,"gdb_version":version,
    }
    ids_response=_query_with_retry(layer,attempts=2,**id_options)
    ids=_objectids_from_query(ids_response,"all Rwanda LV line candidates")
    if ids is None:
        ids=_confirm_empty_id_response(
            layer,_query_with_retry,id_options,"all Rwanda LV line candidates"
        )
    if len(ids)!=total:
        raise RuntimeError(
            f"Nationwide LV lines returned {len(ids)} unique OBJECTIDs "
            f"but service count is {total}; cannot trust continuity inventory."
        )
    return ids


def _write_geometry_report(
    output, district, usernames, lines, checked, version, source_service
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book=Workbook()
    summary=book.active;summary.title="Summary"
    connected=book.create_sheet("Provisional Connected Lines")
    unattached=book.create_sheet("Unassociated Candidate Lines")
    sources=book.create_sheet("Starting Transformers")
    gaps=book.create_sheet("Suggested Endpoint Gaps")
    cautions=book.create_sheet("Method and Limitations")
    flags=Counter(row["status"] for row in checked["results"])
    in_scope=[r for r in checked["results"] if r["source_count"]]
    out_scope=[r for r in checked["results"] if not r["source_count"]]
    for key,value in [
        ("Metric","Value"),
        ("Mode","PROVISIONAL GEOMETRY - NO UTILITY NETWORK TRACE"),
        ("Transformer starting district",district),
        ("Surveyor account(s)","; ".join(usernames) if usernames else "All (within maximum record limit)"),
        ("National LV line candidates scanned",len(lines)),
        ("Provisional connected line OBJECTIDs",len(in_scope)),
        ("Nationwide candidate lines not linked to local transformer starts",len(out_scope)),
        ("Geometric line components",checked["components"]),
        ("Potential endpoint gaps for review",len(checked["suggested_gaps"])),
        ("Source transformers evaluated",len(checked["source_details"])),
        ("Endpoint tolerance (metres)",checked["endpoint_tolerance_m"]),
        ("Transformer proximity search (metres)",checked["seed_tolerance_m"]),
        ("Name differences for review",flags["NAME_DIFFERENCE_PROVISIONAL"]),
        ("Multiple nearby source transformer cases",flags["MULTIPLE_TRANSFORMERS_NEAR_COMPONENT"]),
        ("Matched recorded labels (still provisional)",flags["PROVISIONAL_GEOMETRY_MATCH"]),
        ("Source","Utility Network FeatureServer Electric Line and Electric Device geometry; no topology validation"),
        ("Version",version),
        ("FeatureServer",source_service),
    ]:summary.append([key,value])
    hdr=[
        "Review Status","Line OBJECTID","Line GlobalID","Created User",
        "Recorded Transformer Name","Provisional Nearby Transformer Name",
        "Nearby Transformer OBJECTID","Nearby Transformer GlobalID",
        "Number of Nearby Source Transformers","Candidate Source OBJECTIDs",
        "Geometric Component ID","Component Size","Multiple Geometry Paths",
        "Evidence",
        "Phases","Voltage","Material","Conductor Size","Rated Current",
        "Insulation",
    ]
    connected.append(hdr);unattached.append(hdr)
    by_oid={int(x["attributes"]["OBJECTID"]):x["attributes"] for x in lines}
    for r in checked["results"]:
        row=by_oid[r["objectid"]]
        content=[
            r["status"],r["objectid"],r["globalid"],r["created_user"],
            row.get("Transformer_Name"),r["connected_transformer_name"],
            r["transformer_objectid"],r["transformer_globalid"],
            r["source_count"],"; ".join(map(str,r["source_objectids"])),
            r["component_id"],r["component_size"],
            "YES" if r["multipart"] else "NO",r["evidence"],
            row.get("Phases"),row.get("Voltage_Level_kV"),
            row.get("Material"),row.get("CondustorSize"),
            row.get("Rated_Current"),row.get("ConductorInsulation"),
        ]
        (connected if r["source_count"] else unattached).append(content)
    sources.append([
        "Transformer OBJECTID","Transformer GlobalID","Recorded Name",
        "Geometric Component Candidates","Nearest Endpoint (m)","Status",
    ])
    for r in checked["source_details"]:
        sources.append([
            r["objectid"],r["globalid"],r["name"],
            r["candidate_components"],r["nearest_endpoint_m"],r["status"],
        ])
    gaps.append([
        "LV Line OBJECTID A","LV Line OBJECTID B",
        "Endpoint Separation (m)","Interpretation",
    ])
    for r in checked["suggested_gaps"]:
        gaps.append([
            r["line_oid_a"],r["line_oid_b"],r["distance_m"],
            "REVIEW ONLY: not connected by this algorithm",
        ])
    for row in [
        ("Rule","Limitation"),
        ("Trace","No Utility Network topology trace or dirty-area validation is required."),
        ("Starts","Starting transformer points are spatially selected INSIDE requested district."),
        ("Nationwide scope","Matching creator LV line geometries are retrieved across Rwanda, even beyond district borders."),
        ("Continuity","Line graph joins ONLY same-created_user geometries whose endpoints are <= tolerance apart."),
        ("T-junctions","Endpoint touching middle of another line is NOT automatically traversed; may undercount branches."),
        ("Cross creator","A line made by a different user may interrupt a real feeder; no automatic bridging."),
        ("Proximity","Source association is possible geometry proximity, NOT electrical connectivity."),
        ("Name","Recorded transformer names cannot prove source identity and are not silently corrected."),
        ("Multiple sources","Several transformers near the same component -> ambiguous, engineering review required."),
        ("Switching","Open switches, deenergized sections, phase errors and actual electrical flow NOT verified."),
        ("Projected CRS","EPSG:3857 metre distances are approximate around Rwanda, not survey-grade measurements."),
        ("Suggested gaps","Near endpoints beyond tolerance are listed but NOT traversed automatically."),
        ("Other records","Not linked to one of the district's transformer starts; included for completeness review."),
        ("No source edits","Read-only source data, Excel report only."),
    ]:cautions.append(row)
    for sheet in book.worksheets:
        sheet.freeze_panes="A2"
        sheet.sheet_view.showGridLines=False
        for c in sheet[1]:
            c.font=Font(bold=True,color="FFFFFF")
            c.fill=PatternFill("solid",fgColor="17365D")
        for col in sheet.columns:
            first=next(iter(col))
            width=min(70,max(13,max(len(str(c.value or "")) for c in list(col)[:500])+2))
            sheet.column_dimensions[first.column_letter].width=width
        for row in sheet.iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value[:1] in ("=","+","-","@"):
                    c.data_type="s"
    if output.exists():
        raise FileExistsError(f"Report already exists: {output}")
    book.save(output)
    return {
        "connected":len(in_scope),
        "not_associated":len(out_scope),
        "review":len([r for r in in_scope if r["status"]!="PROVISIONAL_GEOMETRY_MATCH"]),
    }


def run_provisional_geometry_feeder_report(
    output_folder, district, usernames=None, service_url=None,
    boundary_service_url=None, version="sde.DEFAULT", gis=None,
    progress=None, endpoint_tolerance=1.5, seed_tolerance=7.0,
    suggestion_tolerance=3.0, max_candidates=MAX_GLOBAL_LINE_CANDIDATES,
):
    if not district:
        raise ValueError("A starting-transformer district is required.")
    dest=Path(output_folder).expanduser().resolve()
    if not dest.is_dir():
        raise ValueError(f"Choose an existing output folder: {dest}")
    from arcgis.features import FeatureLayer
    active=_get_gis(gis)
    service=str(service_url or DEFAULT_LINE_SERVICE_URL).rstrip("/")
    boundary=str(boundary_service_url or DEFAULT_BOUNDARY_SERVICE_URL).rstrip("/")
    line_layer=FeatureLayer(_layer_url(service,LINE_LAYER_ID),active)
    transformer_layer=FeatureLayer(_layer_url(service,TRANSFORMER_LAYER_ID),active)
    sector_layer=FeatureLayer(_layer_url(boundary,SECTOR_LAYER_ID),active)
    cell_layer=FeatureLayer(_layer_url(boundary,CELL_LAYER_ID),active)
    boundary_features=_get_sector_features(
        sector_layer,district,None,cell_layer,None
    )
    if not boundary_features:
        raise RuntimeError(
            "Cannot verify starting-transformer district polygon; no report."
        )
    area_filters=_filters_for_area(boundary_features,None)
    name_field=resolve_fields(
        [dict(field) for field in transformer_layer.properties.fields],
        "transformer",
    ).get("name")
    if not name_field:
        raise RuntimeError("Cannot identify distribution-transformer name in source schema.")
    def notify(message):
        if progress: progress(message)
    transformer_attrs=query_attributes_in_batches(
        transformer_layer,
        "ASSETGROUP = 12 AND ASSETTYPE IN (1,2,3,4,5,6)",
        version,["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE",name_field],
        area_filters,_query_with_retry,duplicate_globalid_policy="report",
        asset_label="Starting distribution transformer",progress=notify,
    )
    transformer_ids=[v["OBJECTID"] for v in transformer_attrs]
    transformers=_fetch_geometry_by_ids(
        transformer_layer,transformer_ids,version,
        ["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE",name_field],
        "starting transformer",notify,
    )
    for v in transformers:
        v["attributes"]["Transfomer_name_"]=v["attributes"].get(name_field)
    if not transformers:
        raise RuntimeError(
            f"No distribution transformers were found inside the {district} "
            "starting boundary. No nationwide continuity scan was attempted."
        )
    notify(
        f"Selected {len(transformers)} starting distribution transformers in {district}. "
        "No Utility Network traces or dirty-area checks will be performed."
    )
    line_fields={
        field["name"].casefold():field["name"]
        for field in line_layer.properties.fields
    }
    must=["OBJECTID","GLOBALID","ASSETGROUP","ASSETTYPE","created_user"]
    if any(x.casefold() not in line_fields for x in must):
        raise RuntimeError("Essential live LV line fields missing.")
    optional=[
        "Transformer_Name","Phases","Voltage_Level_kV","Material",
        "CondustorSize","Rated_Current","ConductorInsulation",
    ]
    columns=[line_fields[c.casefold()] for c in must+optional
             if c.casefold() in line_fields]
    objectids=_get_nationwide_line_ids(
        line_layer,usernames,version,max_candidates
    )
    notify(f"Scanning {len(objectids):,} LV line geometries across all Rwanda; "
           "the district boundary does not clip the line data.")
    line_records=_fetch_geometry_by_ids(
        line_layer,objectids,version,columns,"LV line",notify,
    )
    checked=group_line_continuity(
        line_records,transformers,endpoint_tolerance,seed_tolerance,
        suggestion_tolerance,
    )
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    clean_district="".join(c if c.isalnum() else "_" for c in district)
    output=dest/f"LV_Geometry_Continuity_{clean_district}_{stamp}.xlsx"
    sheet_stats=_write_geometry_report(
        output,district,usernames or [],line_records,checked,version,service
    )
    return {
        "output_xlsx":str(output),
        "read_only":True,
        "mode":"PROVISIONAL_GEOMETRY_NO_TRACE",
        "starting_transformers":len(transformers),
        "nationwide_candidate_lv_lines":len(line_records),
        "connected_lv_line_ids":sheet_stats["connected"],
        "surveyed_connected_lines":sheet_stats["connected"],
        "engineering_review_lines":sheet_stats["review"],
        "unassociated_lines":sheet_stats["not_associated"],
        "suggested_gap_count":len(checked["suggested_gaps"]),
        "boundary_rule":"District selects transformer start only; geometry can cross borders",
        "caution":"Physical geometry, creator and nearby transformer do not prove energized electrical connectivity.",
    }
