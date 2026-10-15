"""Read-only data-completeness watcher for newly surveyed LV lines."""

from __future__ import annotations

import json
import math
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from .progress_report import (
    CELL_LAYER_ID,
    DEFAULT_BOUNDARY_SERVICE_URL,
    DEFAULT_LINE_SERVICE_URL,
    DEFAULT_USER_NAME_MAP_PATH,
    LINE_LAYER_ID,
    _as_date,
    _feature_layers,
    _get_gis,
    _get_sector_features,
    _layer_url,
    _load_name_overrides,
    _query_with_retry,
    _resolve_full_names,
    _sql_text,
)
from .review_register import REVIEW_HEADERS, build_review_rows, find_previous_report, load_review_rows
from .phase_consistency import phase_type_conflict
from .area_selection import selected_districts
from .quality_query_batches import query_attributes_in_batches
from .outstanding_gaps import GAP_CATEGORIES, priority_gaps, summarize_priority_gaps


DEFAULT_QUALITY_RULES_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "lv_line_quality_rules.json"
)


def load_quality_rules(path: str | Path | None = None) -> dict[str, Any]:
    candidate = Path(path).expanduser() if path else DEFAULT_QUALITY_RULES_PATH
    with candidate.open("r", encoding="utf-8") as stream:
        rules = json.load(stream)
    if not isinstance(rules, dict) or not isinstance(rules.get("scored_attributes"), list):
        raise ValueError("The LV line quality rules file is not valid.")
    if not rules["scored_attributes"]:
        raise ValueError("At least one scored attribute is required.")
    threshold = rules.get("minimum_completion_percent")
    if not _is_finite_number(threshold) or not 0 <= float(threshold) <= 100:
        raise ValueError("The default completion percentage must be from 0 to 100.")
    keys = set()
    for rule in rules["scored_attributes"]:
        if not all(rule.get(field) for field in ("key", "label", "field", "validation")):
            raise ValueError("Each scored attribute needs a key, label, field, and validation.")
        if rule["key"] in keys:
            raise ValueError(f"Duplicate quality rule key: {rule['key']}")
        keys.add(rule["key"])
    if not any(rule.get("mandatory") for rule in rules["scored_attributes"]):
        raise ValueError("At least one quality rule must be mandatory.")
    return rules


def _timestamp_sql(value: date) -> str:
    return value.strftime("%Y-%m-%d 00:00:00")


def build_quality_where(
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    usernames: Sequence[str] | None = None,
) -> str:
    start = _as_date(start_date)
    end = _as_date(end_date)
    if end < start:
        raise ValueError("The end date cannot be earlier than the start date.")
    parts = [
        "ASSETGROUP IN (9,10)",
        "((created_date >= TIMESTAMP '"
        f"{_timestamp_sql(start)}' AND created_date < TIMESTAMP '"
        f"{_timestamp_sql(end + timedelta(days=1))}') OR "
        "(last_edited_date >= TIMESTAMP '"
        f"{_timestamp_sql(start)}' AND last_edited_date < TIMESTAMP '"
        f"{_timestamp_sql(end + timedelta(days=1))}'))",
    ]
    clean_users = [str(item).strip() for item in usernames or [] if str(item).strip()]
    if clean_users:
        quoted = ",".join(f"'{_sql_text(item)}'" for item in clean_users)
        parts.append(f"created_user IN ({quoted})")
    return " AND ".join(parts)


def _domain_maps(line_layer: Any) -> dict[int, dict[str, dict[Any, str]]]:
    maps: dict[int, dict[str, dict[Any, str]]] = {}
    for subtype in line_layer.properties.get("types", []):
        subtype_id = int(subtype.get("id"))
        if subtype_id not in (9, 10):
            continue
        field_maps: dict[str, dict[Any, str]] = {}
        for field_name, domain in subtype.get("domains", {}).items():
            if not isinstance(domain, dict):
                continue
            coded_values = domain.get("codedValues")
            if coded_values:
                field_maps[field_name] = {
                    item.get("code"): str(item.get("name"))
                    for item in coded_values
                }
        maps[subtype_id] = field_maps
    return maps



def _verified_phase_expectations(domains: dict) -> dict[str, int]:
    """Use only explicit single/three-phase labels in live ASSETTYPE domains."""
    expected = {}
    for group in (9, 10):
        labels = domains.get(group, {}).get("ASSETTYPE", {})
        for code, label in labels.items():
            normalized = " ".join(str(label).casefold().replace("-", " ").split())
            is_single = "single phase" in normalized
            is_three = "three phase" in normalized or "3 phase" in normalized
            if is_single != is_three:
                expected[f"{group}:{code}"] = 1 if is_single else 3
    return expected


def _validate_service_schema(line_layer: Any, rules: dict[str, Any], domains: dict) -> None:
    fields = {str(item.get("name", "")).casefold() for item in line_layer.properties.get("fields", [])}
    expected = {"OBJECTID", "GLOBALID", "created_user", "created_date", "last_edited_user", "last_edited_date", "ASSETGROUP", "ASSETTYPE", "Notes"}
    expected.update(str(item["field"]) for item in rules["scored_attributes"])
    expected.update(str(item["field"]) for item in rules.get("optional_attributes", []))
    missing_fields = sorted(name for name in expected if name.casefold() not in fields)
    if missing_fields:
        raise ValueError("The LV feature service is missing required fields: " + ", ".join(missing_fields))
    missing_domains = []
    for group in (9, 10):
        for rule in rules["scored_attributes"]:
            if str(rule["validation"]).startswith("domain_") and not domains.get(group, {}).get(rule["field"]):
                missing_domains.append(f"{rule['field']} (asset group {group})")
    if missing_domains:
        raise ValueError("The LV feature service is missing coded domains: " + ", ".join(missing_domains))


def _combined_geometry_filter(
    sector_feature_set: Any | None,
    custom_boundary_geometry: dict[str, Any] | None = None,
) -> Any | None:
    if sector_feature_set is None and custom_boundary_geometry is None:
        return None
    from arcgis.geometry.filters import intersects

    if custom_boundary_geometry is not None:
        geometry = custom_boundary_geometry
        spatial_reference = geometry.get("spatialReference", {})
    else:
        rings = []
        for feature in sector_feature_set.features:
            rings.extend(feature.geometry.get("rings", []))
        spatial_reference = sector_feature_set.spatial_reference
        geometry = {
            "rings": rings,
            "spatialReference": spatial_reference,
        }
    return intersects(geometry, sr=spatial_reference)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_finite_number(value: Any) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _decode_domain_value(
    attributes: dict[str, Any],
    field_name: str,
    domains: dict[int, dict[str, dict[Any, str]]],
) -> Any:
    value = attributes.get(field_name)
    asset_group = int(attributes.get("ASSETGROUP") or 0)
    return domains.get(asset_group, {}).get(field_name, {}).get(value, value)


def _assess_rule(
    attributes: dict[str, Any],
    rule: dict[str, Any],
    domains: dict[int, dict[str, dict[Any, str]]],
    placeholder_strings: set[str],
    current_year: int,
) -> dict[str, Any]:
    field_name = str(rule["field"])
    label = str(rule["label"])
    validation = str(rule["validation"])
    value = attributes.get(field_name)
    decoded = _decode_domain_value(attributes, field_name, domains)
    issue_type = ""
    issue_detail = ""
    valid = False

    if validation == "asset_group":
        valid = value in (9, 10)
        decoded = {9: "Overhead", 10: "Underground"}.get(value, value)
        if not valid:
            issue_type = "invalid"
            issue_detail = "Not an LV overhead or underground line"
    elif _is_blank(value):
        issue_type = "missing"
        issue_detail = "Blank"
    elif validation == "required_text":
        normalized = str(value).strip().casefold()
        valid = normalized not in placeholder_strings
        if not valid:
            issue_type = "invalid"
            issue_detail = f"Placeholder value: {value}"
    elif validation in ("domain_not_placeholder", "domain_any"):
        asset_group = int(attributes.get("ASSETGROUP") or 0)
        domain = domains.get(asset_group, {}).get(field_name, {})
        if value not in domain:
            issue_type = "invalid"
            issue_detail = f"Value is outside the live domain: {value}"
        else:
            label_normalized = str(domain[value]).strip().casefold()
            invalid_labels = {
                str(item).strip().casefold()
                for item in rule.get("invalid_domain_labels", [])
            }
            valid = validation == "domain_any" or label_normalized not in invalid_labels
            if not valid:
                issue_type = "invalid"
                issue_detail = f"Placeholder domain value: {domain[value]}"
    elif validation == "positive_number":
        valid = _is_finite_number(value) and float(value) > 0
        if not valid:
            issue_type = "invalid"
            issue_detail = f"Must be greater than zero: {value}"
    elif validation == "year":
        valid = _is_finite_number(value) and float(value).is_integer() and 1900 <= int(value) <= current_year + 1
        if not valid:
            issue_type = "invalid"
            issue_detail = f"Expected a year from 1900 to {current_year + 1}: {value}"
    else:
        raise ValueError(f"Unsupported quality validation rule: {validation}")

    return {
        "key": str(rule["key"]),
        "label": label,
        "field": field_name,
        "valid": bool(valid),
        "issue_type": issue_type,
        "issue_detail": issue_detail,
        "raw_value": value,
        "display_value": decoded,
    }


def assess_line_quality(
    attributes: dict[str, Any],
    rules: dict[str, Any],
    domains: dict[int, dict[str, dict[Any, str]]],
    minimum_completion_percent: float = 60.0,
    current_year: int | None = None,
) -> dict[str, Any]:
    """Score one LV line against the attached-template attributes."""
    if not 0 <= float(minimum_completion_percent) <= 100:
        raise ValueError("The minimum completion percentage must be from 0 to 100.")
    active_year = current_year or datetime.now().year
    placeholders = {
        str(item).strip().casefold()
        for item in rules.get("placeholder_strings", [])
    }
    checks = [
        _assess_rule(attributes, rule, domains, placeholders, active_year)
        for rule in rules["scored_attributes"]
    ]

    construction = next((item for item in checks if item["key"] == "year_construction"), None)
    commissioning = next((item for item in checks if item["key"] == "year_commissioning"), None)
    if construction and commissioning and construction["valid"] and commissioning["valid"]:
        if int(commissioning["raw_value"]) < int(construction["raw_value"]):
            commissioning["valid"] = False
            commissioning["issue_type"] = "invalid"
            commissioning["issue_detail"] = "Earlier than the construction year"

    completed = sum(1 for item in checks if item["valid"])
    required = len(checks)
    completion_percent = 100.0 * completed / required
    missing = [item["label"] for item in checks if item["issue_type"] == "missing"]
    invalid = [
        f"{item['label']} ({item['issue_detail']})"
        for item in checks
        if item["issue_type"] == "invalid"
    ]
    mandatory_failures = [item["label"] for item, rule in zip(checks, rules["scored_attributes"]) if rule.get("mandatory") and not item["valid"]]
    status = "MEETS_THRESHOLD" if completion_percent >= float(minimum_completion_percent) and not mandatory_failures else "FLAG"
    display_values = {item["key"]: item["display_value"] for item in checks}
    display_values.update(
        {
            item["key"]: attributes.get(item["field"])
            for item in rules.get("optional_attributes", [])
        }
    )
    return {
        "status": status,
        "completion_percent": completion_percent,
        "completed_fields": completed,
        "required_fields": required,
        "missing_attributes": missing,
        "invalid_attributes": invalid,
        "mandatory_failures": mandatory_failures,
        "checks": checks,
        "display_values": display_values,
    }


def _epoch_to_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, (int, float)):
        return datetime.fromtimestamp(float(value) / 1000.0, tz=timezone.utc).replace(tzinfo=None)
    return None


def _query_quality_attributes(
    line_layer: Any,
    where: str,
    version: str,
    geometry_filter: Any | None,
    rules: dict[str, Any],
    sector_feature_set: Any | None = None,
    custom_boundary_geometry: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Search each selected administrative polygon, then fetch ID batches.

    A multi-district polygon union generated expensive spatial-attribute
    requests that timed out with ArcGIS Server 503. Querying IDs per
    sector/cell avoids that union, and each subsequent attribute request
    uses a small indexed OBJECTID batch with no geometry filter.
    """
    fields = [
        "OBJECTID", "GLOBALID", "created_user", "created_date",
        "last_edited_user", "last_edited_date", "ASSETGROUP",
        "ASSETTYPE", "Notes",
    ]
    fields.extend(str(item["field"]) for item in rules["scored_attributes"])
    fields.extend(str(item["field"]) for item in rules.get("optional_attributes", []))
    unique_fields = list(dict.fromkeys(fields))

    filters: list[tuple[str, Any]] = []
    if sector_feature_set is not None and custom_boundary_geometry is None:
        from arcgis.geometry.filters import intersects

        reference = sector_feature_set.spatial_reference
        for index, feature in enumerate(sector_feature_set.features, 1):
            attrs = dict(getattr(feature, "attributes", {}) or {})
            name = " / ".join(
                str(attrs.get(field) or "").strip()
                for field in ("district", "sector", "cell")
                if str(attrs.get(field) or "").strip()
            ) or f"Boundary polygon {index}"
            geometry = getattr(feature, "geometry", None)
            if not geometry or not geometry.get("rings"):
                raise RuntimeError(f"Selected boundary has no polygon geometry: {name}")
            filters.append((name, intersects(geometry, sr=reference)))
    else:
        filters.append((
            "Custom boundary" if custom_boundary_geometry is not None else "All Rwanda",
            geometry_filter,
        ))

    def report_progress(message: str) -> None:
        try:
            import arcpy
            arcpy.AddMessage(message)
        except ImportError:  # tests outside ArcGIS Pro
            pass

    return query_attributes_in_batches(
        layer=line_layer,
        where=where,
        version=version,
        fields=unique_fields,
        geometry_filters=filters,
        query_fn=_query_with_retry,
        progress=report_progress,
        duplicate_globalid_policy="report",
    )


def _annotate_duplicate_line_globalids(results):
    """Preserve all distinct OBJECTIDs sharing a GlobalID, within report scope."""
    groups = defaultdict(list)
    unique_ids = set()
    for result in results:
        oid = int(result["source_objectid"])
        if oid in unique_ids:
            raise RuntimeError(f"Repeated source OBJECTID {oid} in LV QA results.")
        unique_ids.add(oid)
        result["duplicate_peer_objectids"] = []
        result["identifier_issues"] = []
        gid = str(result.get("globalid") or "").strip().strip("{}").casefold()
        if gid:
            groups[gid].append(result)
    duplicates = {}
    for gid, entries in sorted(groups.items()):
        if len(entries) < 2:
            continue
        oids = sorted(int(row["source_objectid"]) for row in entries)
        duplicates[gid] = oids
        for row in entries:
            peers = [oid for oid in oids if oid != row["source_objectid"]]
            row["duplicate_peer_objectids"] = peers
            row["identifier_issues"] = [
                "Duplicate GlobalID shared with OBJECTID(s): " +
                ", ".join(map(str, peers))
            ]
    return duplicates


def _summarize_users(
    results: Sequence[dict[str, Any]],
    full_names: dict[str, str],
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for result in results:
        grouped[result["username"]].append(result)
    rows = []
    for username, items in grouped.items():
        flagged = sum(item["status"] == "FLAG" for item in items)
        gap_counts = Counter()
        for item in items:
            gap_counts.update(item["missing_attributes"])
            gap_counts.update(text.split(" (", 1)[0] for text in item["invalid_attributes"])
        common_gaps = ", ".join(
            f"{label} ({count})" for label, count in gap_counts.most_common(3)
        )
        rows.append({
            "full_name": full_names.get(username, username),
            "username": username,
            "total_lines": len(items),
            "met_threshold_lines": len(items) - flagged,
            "flagged_lines": flagged,
            "flag_rate": flagged / len(items) if items else 0.0,
            "average_completion": sum(item["completion_percent"] for item in items) / len(items),
            "most_common_gaps": common_gaps,
        })
    return sorted(rows, key=lambda item: (-item["flag_rate"], item["username"].casefold()))


def _set_column_widths(sheet: Any, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _keep_source_text_literal(sheet: Any, first_row: int, last_row: int) -> None:
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"


def _create_quality_workbook(
    output_path: Path,
    results: Sequence[dict[str, Any]],
    user_summary: Sequence[dict[str, Any]],
    rules: dict[str, Any],
    threshold: float,
    title: str,
    metadata: Sequence[tuple[str, Any]],
    review_rows: Sequence[dict[str, Any]],
    duplicate_groups: dict[str, list[int]] | None = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Linesman Summary"
    priority_summary_sheet = workbook.create_sheet("Outstanding Gap Summary", 0)
    priority_details_sheet = workbook.create_sheet("Outstanding LV Gaps", 1)
    details_sheet = workbook.create_sheet("Line Details")
    rules_sheet = workbook.create_sheet("Rules")
    review_sheet = workbook.create_sheet("Review Queue")
    duplicate_group_sheet = workbook.create_sheet("Duplicate GlobalIDs")
    duplicate_records_sheet = workbook.create_sheet("Duplicate Line Records")
    metadata_sheet = workbook.create_sheet("Metadata")
    metadata_sheet.sheet_state = "hidden"

    navy = "17365D"
    blue = "5B9BD5"
    light_blue = "D9EAF7"
    green = "70AD47"
    light_green = "E2F0D9"
    red = "C00000"
    light_red = "F4CCCC"
    yellow = "FFD966"
    light_gray = "E7E6E6"
    white = "FFFFFF"
    border = Border(bottom=Side(style="thin", color="BFBFBF"))

    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.merge_cells("A1:H1")
    summary_sheet["A1"] = title
    summary_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    summary_sheet["A1"].font = Font(color=white, bold=True, size=15)
    summary_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary_sheet.row_dimensions[1].height = 40

    detail_end = max(2, len(results) + 1)
    summary_sheet["A3"] = "Total Lines"
    summary_sheet["B3"] = f"=COUNTA('Line Details'!E2:E{detail_end})"
    summary_sheet["C3"] = "Meets Threshold"
    summary_sheet["D3"] = f'=COUNTIF(\'Line Details\'!A2:A{detail_end},"MEETS_THRESHOLD")'
    summary_sheet["E3"] = "Flagged"
    summary_sheet["F3"] = f'=COUNTIF(\'Line Details\'!A2:A{detail_end},"FLAG")'
    summary_sheet["G3"] = "Threshold Rate"
    summary_sheet["H3"] = "=IF(B3=0,0,D3/B3)"
    summary_sheet["A4"] = "Average Completion"
    summary_sheet["B4"] = f"=IFERROR(AVERAGE('Line Details'!B2:B{detail_end}),0)"
    summary_sheet["C4"] = "Minimum Completion"
    summary_sheet["D4"] = "='Rules'!B2"
    summary_sheet["E4"] = "Flag Rule"
    summary_sheet.merge_cells("F4:H4")
    summary_sheet["F4"] = "Completion below the minimum is FLAGGED"
    summary_sheet["A5"] = "Review Queue Items"
    summary_sheet["B5"] = len(review_rows)
    summary_sheet["C5"] = "Active Reviews"
    summary_sheet["D5"] = sum(row["review_status"] in {"OPEN", "REOPENED", "IN_REVIEW"} for row in review_rows)
    summary_sheet["E5"] = "Duplicate GlobalID groups"
    summary_sheet["F5"] = len(duplicate_groups or {})
    summary_sheet["G5"] = "Affected source OBJECTIDs"
    summary_sheet["H5"] = sum(len(x) for x in (duplicate_groups or {}).values())
    for row in summary_sheet["3:4"]:
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in (1, 3, 5, 7):
        summary_sheet.cell(3, column).fill = PatternFill("solid", fgColor=light_blue)
        summary_sheet.cell(3, column).font = Font(bold=True, color=navy)
    for cell_ref in ("A4", "C4", "E4"):
        summary_sheet[cell_ref].fill = PatternFill("solid", fgColor=light_blue)
        summary_sheet[cell_ref].font = Font(bold=True, color=navy)
    summary_sheet["D3"].fill = PatternFill("solid", fgColor=light_green)
    summary_sheet["F3"].fill = PatternFill("solid", fgColor=light_red)
    summary_sheet["D4"].fill = PatternFill("solid", fgColor=yellow)
    for cell_ref in ("H3", "B4", "D4"):
        summary_sheet[cell_ref].number_format = "0.0%"

    summary_headers = [
        "Linesman Name",
        "Username",
        "Total Lines",
        "Meets Threshold",
        "Flagged",
        "Flag Rate",
        "Average Completion",
        "Most Common Gaps",
    ]
    header_row = 7
    for column, value in enumerate(summary_headers, start=1):
        summary_sheet.cell(header_row, column, value)
    for row_index, item in enumerate(user_summary, start=header_row + 1):
        row_values = [
            item["full_name"],
            item["username"],
            item["total_lines"],
            item["met_threshold_lines"],
            item["flagged_lines"],
            item["flag_rate"],
            item["average_completion"] / 100.0,
            item["most_common_gaps"],
        ]
        for column, value in enumerate(row_values, start=1):
            summary_sheet.cell(row_index, column, value)
        summary_sheet.cell(row_index, 6).number_format = "0.0%"
        summary_sheet.cell(row_index, 7).number_format = "0.0%"
    summary_end = header_row + max(1, len(user_summary))
    if user_summary:
        summary_table = Table(displayName="LinesmanQualitySummary", ref=f"A{header_row}:H{summary_end}")
        summary_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        summary_sheet.add_table(summary_table)
        summary_sheet.conditional_formatting.add(
            f"F{header_row + 1}:F{summary_end}",
            ColorScaleRule(start_type="min", start_color=light_green, mid_type="percentile",
                           mid_value=50, mid_color=yellow, end_type="max", end_color=light_red),
        )
        _keep_source_text_literal(summary_sheet, header_row + 1, summary_end)
    summary_sheet.freeze_panes = "A8"
    _set_column_widths(summary_sheet, {
        "A": 30, "B": 25, "C": 13, "D": 12, "E": 12,
        "F": 13, "G": 20, "H": 54,
    })

    detail_headers = [
        "Status", "Completion %", "Completed Fields", "Required Fields",
        "Source OBJECTID", "GlobalID", "Linesman Name", "Username",
        "Created Date", "Last Edited Date", "Last Edited User", "Construction Type",
        "Transformer Name", "LV Feeder Name", "Voltage Level", "Conductor Size",
        "Conductor Material", "Rated Current (A)", "Conductor Insulation", "Phases",
        "Supplying Lines", "Year of Construction", "Year of Commissioning",
        "Missing Attributes", "Invalid Attributes", "Notes",
    ]
    details_sheet.append(detail_headers)
    # Excel does not accept a direct cross-sheet reference inside a
    # conditional-formatting rule. Keep a same-sheet helper formula instead;
    # otherwise Excel repairs the workbook every time it is opened.
    details_sheet["AA1"] = "='Rules'!$B$2"
    details_sheet["AA1"].number_format = "0.0%"
    details_sheet.column_dimensions["AA"].hidden = True
    for result in results:
        values = result["display_values"]
        details_sheet.append([
            result["status"],
            result["completion_percent"] / 100.0,
            result["completed_fields"],
            result["required_fields"],
            result["source_objectid"],
            result["globalid"],
            result["full_name"],
            result["username"],
            result["created_date"],
            result["last_edited_date"],
            result["last_edited_user"],
            values.get("line_construction_type"),
            values.get("transformer_name"),
            values.get("lv_feeder_name"),
            values.get("voltage"),
            values.get("conductor_size"),
            values.get("conductor_material"),
            values.get("rated_current"),
            values.get("conductor_insulation"),
            values.get("phases"),
            values.get("supplying_lines"),
            values.get("year_construction"),
            values.get("year_commissioning"),
            "; ".join(result["missing_attributes"]),
            "; ".join(result["invalid_attributes"]),
            result["notes"],
        ])
    details_end = max(1, len(results) + 1)
    if results:
        details_table = Table(displayName="LVLineQualityDetails", ref=f"A1:Z{details_end}")
        details_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        details_sheet.add_table(details_table)
        details_sheet.conditional_formatting.add(
            f"A2:A{details_end}",
            FormulaRule(formula=['A2="FLAG"'], fill=PatternFill("solid", fgColor=light_red),
                        font=Font(color=red, bold=True)),
        )
        _keep_source_text_literal(details_sheet, 2, details_end)
        details_sheet.conditional_formatting.add(
            f"A2:A{details_end}",
            FormulaRule(formula=['A2="MEETS_THRESHOLD"'], fill=PatternFill("solid", fgColor=light_green),
                        font=Font(color="006100", bold=True)),
        )
        details_sheet.conditional_formatting.add(
            f"B2:B{details_end}",
            CellIsRule(operator="lessThan", formula=["$AA$1"],
                       fill=PatternFill("solid", fgColor=light_red)),
        )
    for row in details_sheet.iter_rows(min_row=2, max_row=details_end):
        row[1].number_format = "0.0%"
        row[8].number_format = "yyyy-mm-dd hh:mm"
        row[9].number_format = "yyyy-mm-dd hh:mm"
    details_sheet.freeze_panes = "A2"
    details_sheet.sheet_view.showGridLines = False
    details_sheet.auto_filter.ref = f"A1:Z{details_end}"
    _set_column_widths(details_sheet, {
        "A": 21, "B": 15, "C": 16, "D": 15, "E": 17, "F": 39,
        "G": 31, "H": 25, "I": 19, "J": 19, "K": 25, "L": 18,
        "M": 28, "N": 24, "O": 15, "P": 26, "Q": 23, "R": 18,
        "S": 22, "T": 11, "U": 22, "V": 20, "W": 21, "X": 48,
        "Y": 58, "Z": 55,
    })

    rules_sheet.sheet_view.showGridLines = False
    rules_sheet["A1"] = "LV Line Data-Completeness Rules"
    rules_sheet.merge_cells("A1:G1")
    rules_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    rules_sheet["A1"].font = Font(color=white, bold=True, size=14)
    rules_sheet["A1"].alignment = Alignment(horizontal="center")
    rules_sheet["A2"] = "Minimum completion threshold"
    rules_sheet["B2"] = threshold / 100.0
    rules_sheet["B2"].number_format = "0.0%"
    rules_sheet["B2"].fill = PatternFill("solid", fgColor=yellow)
    rule_headers = [
        "Attribute", "Live Field", "Validation", "Template Domain / Guidance",
        "Score Weight", "Rule Role", "Notes",
    ]
    for column, value in enumerate(rule_headers, start=1):
        rules_sheet.cell(4, column, value)
    for cell in rules_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for rule in rules["scored_attributes"]:
        rules_sheet.append([
            rule["label"], rule["field"], rule["validation"],
            rule.get("template_domain", ""), 1, "Mandatory" if rule.get("mandatory") else "Scored", rule.get("notes", ""),
        ])
    for optional in rules.get("optional_attributes", []):
        rules_sheet.append([
            optional["label"], optional["field"], "optional",
            optional.get("template_domain", ""), 0, "Optional", optional.get("notes", ""),
        ])
    rules_end = rules_sheet.max_row
    rules_table = Table(displayName="LVLineQualityRules", ref=f"A4:G{rules_end}")
    rules_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
        showLastColumn=False, showColumnStripes=False,
    )
    rules_sheet.add_table(rules_table)
    rules_sheet.freeze_panes = "A5"
    _set_column_widths(rules_sheet, {
        "A": 28, "B": 27, "C": 24, "D": 55, "E": 14, "F": 12, "G": 55,
    })

    review_sheet.sheet_view.showGridLines = False
    review_sheet.merge_cells("A1:Q1")
    review_sheet["A1"] = "LV Line Review Queue"
    review_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    review_sheet["A1"].font = Font(color=white, bold=True, size=14)
    review_sheet.merge_cells("A2:Q2")
    review_sheet["A2"] = "Edit the yellow review columns and save this workbook. The next run carries your decisions forward."
    review_sheet["A2"].font = Font(color=navy, italic=True)
    review_sheet.append([])
    review_sheet.append(REVIEW_HEADERS)
    for cell in review_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for row in review_rows:
        review_sheet.append([
            row["record_key"], row["globalid"], row["source_objectid"], row["username"],
            row["this_run_status"], row["completion_percent"] / 100.0,
            row["missing_attributes"], row["invalid_attributes"], row["first_flagged_utc"],
            row["last_flagged_utc"], row["flagged_runs"], row["last_seen_utc"],
            row["review_status"], row["assigned_to"], row["reviewed_by"],
            row["reviewed_at_utc"], row["resolution_notes"],
        ])
        review_sheet.cell(review_sheet.max_row, 6).number_format = "0.0%"
        for column in range(13, 18):
            review_sheet.cell(review_sheet.max_row, column).fill = PatternFill("solid", fgColor=yellow)
    if review_rows:
        review_table = Table(displayName="LVLineReviewQueue", ref=f"A4:Q{review_sheet.max_row}")
        review_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        review_sheet.add_table(review_table)
        _keep_source_text_literal(review_sheet, 5, review_sheet.max_row)
    status_validation = DataValidation(type="list", formula1='"OPEN,IN_REVIEW,RESOLVED,DISMISSED,REOPENED"')
    review_sheet.add_data_validation(status_validation)
    status_validation.add("M5:M1048576")
    review_sheet.freeze_panes = "A5"
    review_sheet.auto_filter.ref = f"A4:Q{max(4, review_sheet.max_row)}"
    _set_column_widths(review_sheet, {
        "A": 40, "B": 38, "C": 18, "D": 24, "E": 22, "F": 16,
        "G": 44, "H": 54, "I": 28, "J": 28, "K": 16, "L": 28,
        "M": 18, "N": 25, "O": 25, "P": 28, "Q": 60,
    })

    # Priority correction register: independent of the percentage-based QA result.
    # A line can MEET_THRESHOLD and still have an outstanding correction.
    priority_summary_sheet.sheet_view.showGridLines = False
    priority_summary_sheet.merge_cells("A1:L1")
    priority_summary_sheet["A1"] = "LV LINES - OUTSTANDING FIELD GAPS"
    priority_summary_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    priority_summary_sheet["A1"].font = Font(color=white, bold=True, size=14)
    priority_summary_sheet.merge_cells("A2:L2")
    priority_summary_sheet["A2"] = (
        "Only original correction priorities. Counts show issues, not distinct features; "
        "construction years and supplying lines remain in Line Details."
    )
    priority_summary_sheet["A2"].alignment = Alignment(wrap_text=True)
    priority_summary_sheet.row_dimensions[2].height = 29
    priority_headers = [
        "Username", "Surveyor", "Lines Checked", "With Outstanding Gaps",
        "Gap Rate", *GAP_CATEGORIES, "Total Issues",
    ]
    priority_summary_sheet.append([])
    priority_summary_sheet.append(priority_headers)
    priority_user_rows = summarize_priority_gaps(results)
    for user in priority_user_rows:
        priority_summary_sheet.append([
            user["username"], user["full_name"], user["lines_checked"],
            user["affected_lines"],
            user["affected_lines"] / user["lines_checked"] if user["lines_checked"] else 0,
            *(user[category] for category in GAP_CATEGORIES),
            user["total_issues"],
        ])
        priority_summary_sheet.cell(priority_summary_sheet.max_row, 5).number_format = "0.0%"
    total_checked = sum(item["lines_checked"] for item in priority_user_rows)
    total_affected = sum(item["affected_lines"] for item in priority_user_rows)
    total_issues = sum(item["total_issues"] for item in priority_user_rows)
    total_cells = [
        "TOTAL", "", total_checked, total_affected,
        total_affected / total_checked if total_checked else 0,
        *(sum(item[category] for item in priority_user_rows) for category in GAP_CATEGORIES),
        total_issues,
    ]
    priority_summary_sheet.append(total_cells)
    total_row = priority_summary_sheet.max_row
    priority_summary_sheet.cell(total_row, 5).number_format = "0.0%"
    for cell in priority_summary_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for cell in priority_summary_sheet[total_row]:
        cell.fill = PatternFill("solid", fgColor=light_blue)
        cell.font = Font(color=navy, bold=True)
    priority_summary_sheet.row_dimensions[4].height = 32
    priority_summary_sheet.freeze_panes = "C5"
    priority_summary_sheet.auto_filter.ref = f"A4:L{max(4, total_row - 1)}"
    _set_column_widths(priority_summary_sheet, {
        "A": 28, "B": 28, "C": 17, "D": 23, "E": 14,
        "F": 19, "G": 25, "H": 24, "I": 26, "J": 22,
        "K": 24, "L": 18,
    })
    _keep_source_text_literal(priority_summary_sheet, 5, total_row)

    priority_details_sheet.sheet_view.showGridLines = False
    priority_detail_headers = [
        "Username", "Surveyor", "GlobalID", "Source OBJECTID",
        "Asset Group", "Asset Type", "Transformer Name", "Phases",
        "Issue Count", "Outstanding Issues", "Issue Details",
        "Created Date", "Last Edited Date", "Threshold Status",
        "Completion %", "Notes",
    ]
    priority_details_sheet.append(priority_detail_headers)
    for result in results:
        issues = result.get("priority_gaps", [])
        if not issues:
            continue
        display_values = result["display_values"]
        priority_details_sheet.append([
            result["username"], result["full_name"], result["globalid"],
            result["source_objectid"], result.get("asset_group"),
            result.get("asset_type"), display_values.get("transformer_name"),
            display_values.get("phases"), len(issues),
            "; ".join(issue["category"] for issue in issues),
            "; ".join(issue["category"] + ": " + issue["detail"] for issue in issues),
            result["created_date"], result["last_edited_date"],
            result["status"], result["completion_percent"] / 100.0,
            result["notes"],
        ])
        row_no = priority_details_sheet.max_row
        priority_details_sheet.cell(row_no, 15).number_format = "0.0%"
        for date_column in (12, 13):
            priority_details_sheet.cell(row_no, date_column).number_format = "yyyy-mm-dd hh:mm"
    for cell in priority_details_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(wrap_text=True)
    if priority_details_sheet.max_row > 1:
        gap_table = Table(
            displayName="LVOutstandingGaps",
            ref=f"A1:P{priority_details_sheet.max_row}",
        )
        gap_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        priority_details_sheet.add_table(gap_table)
        _keep_source_text_literal(priority_details_sheet, 2, priority_details_sheet.max_row)
    priority_details_sheet.freeze_panes = "C2"
    priority_details_sheet.auto_filter.ref = f"A1:P{priority_details_sheet.max_row}"
    _set_column_widths(priority_details_sheet, {
        "A": 26, "B": 28, "C": 40, "D": 19, "E": 16, "F": 16,
        "G": 30, "H": 15, "I": 15, "J": 70, "K": 85,
        "L": 21, "M": 21, "N": 23, "O": 15, "P": 45,
    })

    # A separate identifier-integrity register must not alter the original
    # field-completeness scores or silently discard either physical record.
    duplicate_group_sheet.append([
        "Normalized GlobalID", "All OBJECTIDs", "Number of source records",
        "Geometry comparison", "Verification required",
    ])
    for gid, objectids in sorted((duplicate_groups or {}).items()):
        duplicate_group_sheet.append([
            gid, "; ".join(map(str, objectids)), len(objectids),
            "NOT VERIFIED",
            "Inspect geometry, attachments, connectivity and ownership; do not auto-delete.",
        ])
    duplicate_records_sheet.append([
        "GlobalID", "Source OBJECTID", "Other OBJECTIDs", "Surveyor Username",
        "Surveyor Name", "Asset Group", "Asset Type", "Created Date",
        "Last Edited Date", "Last Edited User", "Notes", "Status",
    ])
    for item in sorted(results, key=lambda x: (
        str(x["globalid"]).strip().strip("{}").casefold(),
        int(x["source_objectid"]),
    )):
        peers = item.get("duplicate_peer_objectids", [])
        if not peers:
            continue
        duplicate_records_sheet.append([
            item["globalid"], item["source_objectid"],
            "; ".join(map(str, peers)), item["username"], item["full_name"],
            item.get("asset_group"), item.get("asset_type"),
            item["created_date"], item["last_edited_date"],
            item["last_edited_user"], item["notes"], "IDENTIFIER_CONFLICT",
        ])
        for column in (8, 9):
            duplicate_records_sheet.cell(duplicate_records_sheet.max_row, column).number_format = "yyyy-mm-dd hh:mm"
    for sheet, widths, lastcol in (
        (duplicate_group_sheet, {"A": 44, "B": 43, "C": 24, "D": 20, "E": 87}, "E"),
        (duplicate_records_sheet, {"A": 44, "B": 21, "C": 28, "D": 24,
             "E": 28, "F": 15, "G": 15, "H": 22, "I": 22, "J": 26,
             "K": 65, "L": 26}, "L"),
    ):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{lastcol}{sheet.max_row}"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(wrap_text=True)
        _set_column_widths(sheet, widths)
        if sheet.max_row > 1:
            _keep_source_text_literal(sheet, 2, sheet.max_row)

    metadata_sheet.append(["Field", "Value"])
    for key, value in metadata:
        metadata_sheet.append([key, "" if value is None else str(value)])
    metadata_sheet.column_dimensions["A"].width = 30
    metadata_sheet.column_dimensions["B"].width = 120

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        raise FileExistsError(f"The QA report already exists: {output_path}")
    workbook.save(output_path)


def _spatial_reference_from_json(arcpy: Any, spatial_reference: dict[str, Any]) -> Any:
    if spatial_reference.get("latestWkid") or spatial_reference.get("wkid"):
        return arcpy.SpatialReference(
            int(spatial_reference.get("latestWkid") or spatial_reference.get("wkid"))
        )
    reference = arcpy.SpatialReference()
    reference.loadFromString(spatial_reference.get("wkt") or spatial_reference.get("wkt2"))
    return reference


def _write_flagged_feature_class(
    line_layer: Any,
    flagged_results: Sequence[dict[str, Any]],
    output_feature_class: str | Path,
    version: str,
) -> str:
    import arcpy

    output_path = str(Path(output_feature_class))
    parent, name = os.path.split(output_path)
    if not parent.lower().endswith(".gdb"):
        raise ValueError("The flagged-line output must be inside a file geodatabase.")
    if not arcpy.Exists(parent):
        gdb_parent, gdb_name = os.path.split(parent)
        arcpy.management.CreateFileGDB(gdb_parent, gdb_name)
    if arcpy.Exists(output_path):
        raise FileExistsError(f"The flagged-line output already exists: {output_path}")

    flagged_by_oid = {int(item["source_objectid"]): item for item in flagged_results}
    geometry_features = []
    spatial_reference = dict(line_layer.properties.extent.spatialReference)
    object_ids = sorted(flagged_by_oid)
    for start in range(0, len(object_ids), 500):
        batch = object_ids[start:start + 500]
        feature_set = _query_with_retry(
            line_layer,
            object_ids=",".join(str(item) for item in batch),
            out_fields="OBJECTID",
            return_geometry=True,
            gdb_version=version,
            return_all_records=True,
        )
        spatial_reference = dict(feature_set.spatial_reference or spatial_reference)
        geometry_features.extend(feature_set.features)

    arcpy.management.CreateFeatureclass(
        parent,
        name,
        "POLYLINE",
        spatial_reference=_spatial_reference_from_json(arcpy, spatial_reference),
    )
    fields = [
        ("SOURCE_OID", "LONG", None),
        ("SOURCE_GID", "TEXT", 38),
        ("CREATOR", "TEXT", 255),
        ("CREATED_AT", "DATE", None),
        ("COMPLETE_PCT", "DOUBLE", None),
        ("QA_STATUS", "TEXT", 24),
        ("MISSING", "TEXT", 2000),
        ("INVALID", "TEXT", 2000),
    ]
    for field_name, field_type, length in fields:
        kwargs = {"field_length": length} if length else {}
        arcpy.management.AddField(output_path, field_name, field_type, **kwargs)
    cursor_fields = [
        "SHAPE@", "SOURCE_OID", "SOURCE_GID", "CREATOR", "CREATED_AT",
        "COMPLETE_PCT", "QA_STATUS", "MISSING", "INVALID",
    ]
    with arcpy.da.InsertCursor(output_path, cursor_fields) as cursor:
        for feature in geometry_features:
            oid = int(feature.attributes["OBJECTID"])
            result = flagged_by_oid.get(oid)
            if not result:
                continue
            geometry_json = dict(feature.geometry)
            geometry_json["spatialReference"] = spatial_reference
            cursor.insertRow([
                arcpy.AsShape(geometry_json, True),
                oid,
                result["globalid"],
                result["username"],
                result["created_date"],
                result["completion_percent"],
                result["status"],
                "; ".join(result["missing_attributes"]),
                "; ".join(result["invalid_attributes"]),
            ])
    return str(Path(output_path).resolve())


def _timestamped_output(path: str | Path, run_time: datetime, collision: int = 0) -> Path:
    base = Path(path).expanduser().resolve()
    stem = re.sub(r"_\d{8}_\d{6}(?:_\d+)?$", "", base.stem)
    stamp = run_time.strftime("%Y%m%d_%H%M%S")
    extra = f"_{collision:02d}" if collision else ""
    return base.with_name(f"{stem}_{stamp}{extra}{base.suffix}")


def run_quality_watcher(
    output_xlsx: str | Path,
    start_date: date | datetime | str,
    end_date: date | datetime | str,
    usernames: Sequence[str] | None = None,
    district: str | Sequence[str] | None = None,
    sectors: Sequence[str] | None = None,
    cells: Sequence[str] | None = None,
    custom_boundary_geometry: dict[str, Any] | None = None,
    minimum_completion_percent: float = 60.0,
    output_flagged_feature_class: str | Path | None = None,
    service_url: str | None = None,
    boundary_service_url: str | None = None,
    version: str = "sde.DEFAULT",
    user_name_map_path: str | Path | None = None,
    rules_path: str | Path | None = None,
    previous_review_report: str | Path | None = None,
    gis: Any | None = None,
) -> dict[str, Any]:
    """Scan newly created LV lines, report completeness, and export flagged geometry."""
    start = _as_date(start_date)
    end = _as_date(end_date)
    threshold = float(minimum_completion_percent)
    if end < start:
        raise ValueError("The end date cannot be earlier than the start date.")
    if not 0 <= threshold <= 100:
        raise ValueError("The minimum completion percentage must be from 0 to 100.")

    run_local = datetime.now().astimezone()
    run_utc = run_local.astimezone(timezone.utc)
    output_base = Path(output_xlsx).expanduser().resolve()
    if output_base.suffix.casefold() != ".xlsx":
        raise ValueError("The QA report output must be an .xlsx workbook.")
    flagged_base = Path(output_flagged_feature_class).expanduser().resolve() if output_flagged_feature_class else None
    for collision in range(1000):
        output_path = _timestamped_output(output_base, run_local, collision)
        flagged_path = _timestamped_output(flagged_base, run_local, collision) if flagged_base else None
        if output_path.exists():
            continue
        if flagged_path:
            import arcpy
            if arcpy.Exists(str(flagged_path)):
                continue
        break
    else:
        raise FileExistsError("Could not create a unique name for this QA run.")
    prior_path = Path(previous_review_report).expanduser().resolve() if previous_review_report else find_previous_report(output_path)
    prior_reviews = load_review_rows(prior_path)

    active_gis = _get_gis(gis)
    active_service_url = service_url or os.environ.get(
        "UN_REPORT_SERVICE_URL", DEFAULT_LINE_SERVICE_URL
    )
    active_boundary_url = boundary_service_url or os.environ.get(
        "UN_SECTOR_BOUNDARY_URL", DEFAULT_BOUNDARY_SERVICE_URL
    )
    line_layer, sector_layer, cell_layer = _feature_layers(
        active_gis, active_service_url, active_boundary_url
    )
    sector_feature_set = None
    if custom_boundary_geometry is None:
        sector_feature_set = _get_sector_features(
            sector_layer, district, sectors, cell_layer, cells
        )
    # Do not build one huge union of every sector's geometry: large combined
    # searches cause ArcGIS Server 503 timeouts. The query helper searches
    # each polygon for IDs separately, followed by small attribute batches.
    geometry_filter = (
        _combined_geometry_filter(None, custom_boundary_geometry)
        if custom_boundary_geometry is not None else None
    )
    rules = load_quality_rules(rules_path)
    domains = _domain_maps(line_layer)
    _validate_service_schema(line_layer, rules, domains)
    phase_expectations = _verified_phase_expectations(domains)
    clean_users = list(dict.fromkeys(
        str(item).strip() for item in usernames or [] if str(item).strip()
    ))
    where = build_quality_where(start, end, clean_users)
    attributes_list = _query_quality_attributes(
        line_layer, where, version, geometry_filter, rules,
        sector_feature_set=sector_feature_set,
        custom_boundary_geometry=custom_boundary_geometry,
    )

    usernames_found = sorted({
        str(attributes.get("created_user") or "").strip()
        for attributes in attributes_list
        if str(attributes.get("created_user") or "").strip()
    }, key=str.casefold)
    name_overrides = _load_name_overrides(user_name_map_path)
    full_names = _resolve_full_names(active_gis, usernames_found, name_overrides)

    results = []
    for attributes in attributes_list:
        username = str(attributes.get("created_user") or "").strip() or "(unattributed)"
        quality = assess_line_quality(
            attributes,
            rules,
            domains,
            minimum_completion_percent=threshold,
            current_year=end.year,
        )
        conflict = phase_type_conflict(
            attributes,
            domains.get(int(attributes.get("ASSETGROUP") or 0), {}).get("Phases", {}),
            phase_expectations,
        )
        if conflict:
            quality["invalid_attributes"].append(conflict)
            quality["status"] = "FLAG"
        result = {
            **quality,
            "priority_gaps": priority_gaps(quality, conflict),
            "asset_group": attributes.get("ASSETGROUP"),
            "asset_type": attributes.get("ASSETTYPE"),
            "source_objectid": int(attributes["OBJECTID"]),
            "globalid": str(attributes.get("GLOBALID") or ""),
            "username": username,
            "full_name": full_names.get(username, username),
            "created_date": _epoch_to_datetime(attributes.get("created_date")),
            "last_edited_date": _epoch_to_datetime(attributes.get("last_edited_date")),
            "last_edited_user": str(attributes.get("last_edited_user") or ""),
            "notes": str(attributes.get("Notes") or ""),
        }
        results.append(result)
    duplicate_groups = _annotate_duplicate_line_globalids(results)
    if duplicate_groups:
        try:
            import arcpy
            arcpy.AddWarning(
                f"Found {len(duplicate_groups)} LV GlobalID conflict group(s). "
                "Both OBJECTIDs are retained; see Duplicate GlobalIDs and "
                "Duplicate Line Records in the Excel workbook."
            )
        except ImportError:
            pass
    results.sort(key=lambda item: (
        0 if item["status"] == "FLAG" else 1,
        item["completion_percent"],
        item["username"].casefold(),
        item["source_objectid"],
    ))
    user_summary = _summarize_users(results, full_names)
    priority_user_summary = summarize_priority_gaps(results)
    outstanding_results = [item for item in results if item["priority_gaps"]]
    flagged_results = [item for item in results if item["status"] == "FLAG"]
    met_threshold_count = len(results) - len(flagged_results)
    review_rows = build_review_rows(results, prior_reviews, run_utc)

    clean_sectors = [str(item).strip() for item in sectors or [] if str(item).strip()]
    clean_cells = [str(item).strip() for item in cells or [] if str(item).strip()]
    selected_areas = selected_districts(district)
    district_label = ", ".join(selected_areas)
    area_label = (
        "Custom boundary" if custom_boundary_geometry is not None
        else (district_label or "All Areas")
    )
    if clean_sectors and custom_boundary_geometry is None:
        area_label = f"{area_label}: {', '.join(clean_sectors)}"
    if clean_cells and custom_boundary_geometry is None:
        area_label = f"{area_label}; {', '.join(clean_cells)} cells"
    title = (
        f"LV Line Data-Quality Watcher — {area_label} "
        f"({start.isoformat()} to {end.isoformat()})"
    )
    active_rules_path = Path(rules_path).expanduser() if rules_path else DEFAULT_QUALITY_RULES_PATH
    active_name_map = (
        Path(user_name_map_path).expanduser()
        if user_name_map_path else DEFAULT_USER_NAME_MAP_PATH
    )
    metadata = [
        ("Generated local", run_local.isoformat()),
        ("Generated UTC", run_utc.isoformat()),
        ("From date", start.isoformat()),
        ("To date (inclusive)", end.isoformat()),
        ("Date basis", "Created or last edited"),
        ("District(s)", district_label),
        ("Sectors", "; ".join(clean_sectors)),
        ("Cells", "; ".join(clean_cells)),
        (
            "Boundary mode",
            "Custom drawn polygon" if custom_boundary_geometry is not None else "Administrative",
        ),
        ("Selected users", "; ".join(clean_users) if clean_users else "All matching users"),
        ("Minimum completion percent", threshold),
        ("LV outstanding lines", len(outstanding_results)),
        ("LV outstanding issues", sum(len(item["priority_gaps"]) for item in outstanding_results)),
        ("Reused GlobalID groups (this report only)", len(duplicate_groups)),
        ("Source records with reused GlobalIDs",
         sum(len(v) for v in duplicate_groups.values())),
        ("Geometry comparison for duplicates", "NOT VERIFIED"),
        ("Query", where),
        ("Version", version),
        ("Line feature service", _layer_url(active_service_url, LINE_LAYER_ID)),
        (
            "Boundary service",
            _layer_url(active_boundary_url, CELL_LAYER_ID if clean_cells else 0),
        ),
        ("Quality rules", str(active_rules_path.resolve())),
        ("User-name mapping", str(active_name_map.resolve())),
        ("Previous review report", str(prior_path) if prior_path else ""),
    ]
    _create_quality_workbook(
        output_path, results, user_summary, rules, threshold, title, metadata,
        review_rows, duplicate_groups
    )

    flagged_feature_class = ""
    if flagged_path:
        flagged_feature_class = _write_flagged_feature_class(
            line_layer,
            flagged_results,
            flagged_path,
            version,
        )

    return {
        "read_only_source": True,
        "run_local": run_local.isoformat(),
        "output_xlsx": str(output_path),
        "output_flagged_feature_class": flagged_feature_class,
        "from_date": start.isoformat(),
        "to_date_inclusive": end.isoformat(),
        "minimum_completion_percent": threshold,
        "district": district_label,
        "districts": selected_areas,
        "sectors": clean_sectors,
        "cells": clean_cells,
        "boundary_mode": (
            "custom" if custom_boundary_geometry is not None else "administrative"
        ),
        "total_lines": len(results),
        "met_threshold_lines": met_threshold_count,
        "flagged_lines": len(flagged_results),
        "threshold_rate": met_threshold_count / len(results) if results else 0.0,
        "review_queue_items": len(review_rows),
        "outstanding_lv_lines": len(outstanding_results),
        "outstanding_lv_issues": sum(len(item["priority_gaps"]) for item in outstanding_results),
        "duplicate_globalid_groups": len(duplicate_groups),
        "duplicate_feature_records": sum(len(v) for v in duplicate_groups.values()),
        "duplicate_source_objectids": duplicate_groups,
        "outstanding_lv_by_user": priority_user_summary,
        "previous_review_report": str(prior_path) if prior_path else "",
        "average_completion_percent": (
            sum(item["completion_percent"] for item in results) / len(results)
            if results else 0.0
        ),
        "flagged_objectids": [item["source_objectid"] for item in flagged_results],
        "user_summary": user_summary,
        "where": where,
        "version": version,
    }
