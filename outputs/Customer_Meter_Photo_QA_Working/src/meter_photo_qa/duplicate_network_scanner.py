"""Read-only nationwide GlobalID duplicate audit across REG distribution assets.

One explicit operator action. Pages only OBJECTID/GLOBALID, loads geometry
solely for duplicate records, then writes an Excel register and matching KMZ.
No scheduled jobs, database writes or Utility Network edits.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
import os
import time

from .duplicate_line_inspector import canonical_globalid, _canonical_paths
from .duplicate_line_scanner import fetch_duplicate_features, assign_admin_areas
from .gap_kmz import _add, _n, _write_geometry, _lines, _xy, COLORS
from .progress_report import (
    DEFAULT_LINE_SERVICE_URL, DEFAULT_BOUNDARY_SERVICE_URL,
    CELL_LAYER_ID, _get_gis, _layer_url, _query_with_retry,
)

SCOPES = ("all", "lv")
SOURCES = (
    ("line", 10, "1=1", "Electric Lines (LV, MV, HV)"),
    ("pole", 12, "ASSETGROUP = 7 AND ASSETTYPE IN (1,2,3)", "LV Poles"),
    ("transformer", 7, "ASSETGROUP = 12 AND ASSETTYPE IN (1,2,3,4,5,6)",
     "Distribution Transformers"),
)
TITLES = {"line": "Electric Lines", "pole": "LV Poles",
          "transformer": "Distribution Transformers"}


def _all_ids(layers, version, scope, page_size=500, pause_seconds=0.15,
             progress=None):
    """Collect normalized IDs across all three layers with bounded keyset pages.

    Includes cross-layer repeated GUIDs; OBJECTIDs remain namespaced by asset.
    Aborts if any layer's initial/final count differs or paging is incomplete.
    """
    if scope not in SCOPES:
        raise ValueError("Scope must be 'all' or 'lv'.")
    if not 1 <= page_size <= 1000:
        raise ValueError("Page size must be 1-1000.")
    if not 0 <= pause_seconds <= 3:
        raise ValueError("Request pause must be between 0 and 3 seconds.")
    report = progress or (lambda message: None)
    seen = {}
    dup = {}
    counts = {}
    invalid = {}
    for kind, _id, default_where, label in SOURCES:
        layer = layers[kind]
        where = ("ASSETGROUP IN (9,10)" if kind == "line" and scope == "lv"
                 else default_where)
        expected = int(_query_with_retry(
            layer, where=where, return_count_only=True,
            gdb_version=version, attempts=2,
        ))
        if expected < 0:
            raise RuntimeError(f"{label}: negative server record count.")
        last = -1
        scanned = 0
        skipped = 0
        while True:
            response = _query_with_retry(
                layer, where=f"({where}) AND OBJECTID > {last}",
                out_fields="OBJECTID,GLOBALID", return_geometry=False,
                order_by_fields="OBJECTID ASC", result_record_count=page_size,
                return_all_records=False, gdb_version=version, attempts=2,
            )
            if (not hasattr(response, "features") or
                getattr(response, "exceeded_transfer_limit", False)):
                raise RuntimeError(f"{label}: incomplete identity page; no outputs.")
            if not response.features:
                break
            high = last
            for feature in response.features:
                attributes = dict(feature.attributes)
                raw_oid = attributes.get("OBJECTID")
                if raw_oid is None:
                    raise RuntimeError(f"{label}: missing OBJECTID.")
                oid = int(raw_oid)
                if oid <= high:
                    raise RuntimeError(f"{label}: unordered or repeated OBJECTID.")
                high = oid
                try:
                    gid = canonical_globalid(attributes.get("GLOBALID"))
                except ValueError:
                    skipped += 1
                    continue
                identity = (kind, oid)
                if gid in dup:
                    dup[gid].append(identity)
                elif gid in seen:
                    dup[gid] = [seen.pop(gid), identity]
                else:
                    seen[gid] = identity
            scanned += len(response.features)
            last = high
            if scanned > expected:
                raise RuntimeError(f"{label}: row count increased mid-scan; retry later.")
            if scanned == len(response.features) or scanned % 10000 < len(response.features):
                report(f"{label}: scanned {scanned:,}/{expected:,} identity records.")
            if pause_seconds:
                time.sleep(pause_seconds)
        ending = int(_query_with_retry(
            layer, where=where, return_count_only=True,
            gdb_version=version, attempts=2,
        ))
        if scanned != expected or ending != expected:
            raise RuntimeError(
                f"{label}: count changed/incomplete ({expected} initial, "
                f"{scanned} read, {ending} final); no report is published."
            )
        counts[kind] = scanned
        invalid[kind] = skipped
    return dup, counts, invalid


def _fetch_selected(layers, groups, version, progress=None):
    """Retrieve each duplicate source feature once, verifying its stored GUID."""
    report = progress or (lambda message: None)
    result = {}
    for kind in TITLES:
        subset = {}
        for gid, members in groups.items():
            relevant = [oid for asset, oid in members if asset == kind]
            if relevant:
                subset[gid] = relevant
        if not subset:
            continue
        raw = fetch_duplicate_features(
            layers[kind], subset, version, progress=report, pause_seconds=0.1,
        )
        for oid, item in raw.items():
            key = (kind, oid)
            if key in result:
                raise RuntimeError(f"Duplicate record returned twice: {key}")
            result[key] = item
    return result


def _geometry_signature(kind, geometry):
    if kind == "line":
        return _canonical_paths(geometry)
    if "x" not in geometry or "y" not in geometry:
        return None
    return (float(geometry["x"]), float(geometry["y"]))


def _fetch_wgs84(layers, groups, originals, version):
    """Request WGS84 geometries only for duplicates; verify source ID and CRS.

    This separate bounded duplicate-only fetch is required because the
    original XY geometries used for administrative intersection can be in
    Rwanda TM, which is not legal to write unchanged as KML coordinates.
    """
    result = {}
    for kind in TITLES:
        expected = {
            oid: gid for gid, members in groups.items()
            for asset, oid in members if asset == kind
        }
        if not expected:
            continue
        ordered = sorted(expected)
        for start in range(0, len(ordered), 120):
            chunk = ordered[start:start+120]
            response = _query_with_retry(
                layers[kind], object_ids=",".join(map(str, chunk)),
                where="1=1", out_fields="OBJECTID,GLOBALID,ASSETGROUP",
                out_sr=4326, return_geometry=True, return_all_records=False,
                gdb_version=version, attempts=2,
            )
            if (not hasattr(response, "features") or
                getattr(response, "exceeded_transfer_limit", False)):
                raise RuntimeError(f"{kind}: incomplete KMZ geometry response.")
            shared_sr = dict(getattr(response, "spatial_reference", None) or {})
            for feature in response.features:
                attrs = dict(feature.attributes)
                oid = int(attrs["OBJECTID"])
                key = (kind, oid)
                if oid not in expected or key in result:
                    raise RuntimeError(f"{kind}: unexpected/duplicate KMZ OBJECTID {oid}.")
                if canonical_globalid(attrs.get("GLOBALID")) != expected[oid]:
                    raise RuntimeError(f"{kind}/{oid}: GlobalID changed during KMZ fetch.")
                source_group=originals[key]["attributes"].get("ASSETGROUP")
                if (source_group is None or attrs.get("ASSETGROUP") is None or
                    int(attrs["ASSETGROUP"]) != int(source_group)):
                    raise RuntimeError(
                        f"{kind}/{oid}: asset group changed during duplicate review."
                    )
                if kind == "pole" and int(source_group) != 7:
                    raise RuntimeError(f"{kind}/{oid}: not a low-voltage pole.")
                if kind == "transformer" and int(source_group) != 12:
                    raise RuntimeError(f"{kind}/{oid}: not a distribution transformer.")
                geometry = dict(feature.geometry or {})
                sr = geometry.get("spatialReference") or shared_sr
                code = int(sr.get("latestWkid") or sr.get("wkid") or 0)
                if code != 4326:
                    raise RuntimeError(f"{kind}/{oid}: WGS84 geometry not verified (wkid {code}).")
                if kind == "line":
                    _lines(geometry)
                else:
                    _xy([geometry["x"], geometry["y"]])
                result[key] = geometry
            for oid in chunk:
                if (kind, oid) not in result:
                    raise RuntimeError(f"{kind}/{oid}: KMZ source geometry missing.")
            time.sleep(0.1)
    return result


def _area_index(cells, duplicates, original_features, progress=None):
    result = {}
    for kind in TITLES:
        subset = {oid: item for (asset, oid), item in original_features.items()
                  if asset == kind}
        if not subset:
            continue
        identified = assign_admin_areas(
            cells, subset, progress=progress, pause_seconds=0.1,
        )
        for oid, area in identified.items():
            result[(kind, oid)] = area
    if len(result) != len(original_features):
        raise RuntimeError("Incomplete administrative location results.")
    return result


def _report_excel(path, duplicates, originals, areas, scanned, invalid,
                  service, boundary, version, scope):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    groups_sheet = wb.create_sheet("Duplicate GlobalIDs")
    details = wb.create_sheet("All Duplicate Records")
    locations = wb.create_sheet("By District Sector Cell")
    notes = wb.create_sheet("Method and Warnings")
    affected = sum(len(members) for members in duplicates.values())
    for row in [
        ("Measure", "Value"), ("Scope", "All Electric Lines + LV Poles + Distribution Transformers"
         if scope == "all" else "LV Lines + LV Poles + Distribution Transformers"),
        ("Source FeatureServer", service), ("Admin source", boundary),
        ("Version", version),
        ("Line records scanned", scanned["line"]), ("Pole records scanned", scanned["pole"]),
        ("Transformer records scanned", scanned["transformer"]),
        ("Total records scanned", sum(scanned.values())),
        ("Duplicate GlobalID groups", len(duplicates)),
        ("Affected records", affected), ("Extra copies", affected-len(duplicates)),
        ("Invalid or missing GUIDs (lines)", invalid["line"]),
        ("Invalid or missing GUIDs (poles)", invalid["pole"]),
        ("Invalid or missing GUIDs (transformers)", invalid["transformer"]),
        ("Read-only", "Yes"),
    ]:
        summary.append(row)
    groups_sheet.append([
        "GlobalID", "Number of records", "Asset types", "Members (Asset/OBJECTID)",
        "Cross-layer duplicate", "Same exact geometry (same-type only)",
        "District(s)", "Sector(s)", "Cell(s)", "Admin Status",
    ])
    details.append([
        "GlobalID", "Asset", "Layer", "OBJECTID", "Other duplicate members",
        "Group size", "Cross-layer", "District", "Sector", "Cell", "Admin Status",
        "Asset Group", "Asset Type", "Created User", "Created Date",
        "Last Edited User", "Last Edited Date", "Same Exact Geometry",
    ])
    locations.append([
        "District", "Sector", "Cell", "Duplicate groups", "Duplicate records",
    ])
    aggregation = defaultdict(lambda: [set(), set()])
    layer_ids = {source: layer for source, layer, _, _ in SOURCES}
    for gid, members in sorted(duplicates.items()):
        members = sorted(members)
        types = {asset for asset, _ in members}
        signature = [_geometry_signature(asset, originals[(asset, oid)]["geometry"])
                     for asset, oid in members]
        identical = (len(types) == 1 and all(value is not None for value in signature)
                     and len(set(signature)) == 1)
        cross = len(types) > 1
        group_areas = [areas[key] for key in members]
        values = lambda field: "; ".join(sorted(
            {value for a in group_areas for value in a[field] if value},
            key=str.casefold,
        ))
        groups_sheet.append([
            gid, len(members), "; ".join(sorted(types)),
            "; ".join(f"{asset.upper()}/{oid}" for asset, oid in members),
            "YES" if cross else "NO",
            "YES" if identical else "NO / NOT ESTABLISHED",
            values("district"), values("sector"), values("cell"),
            "; ".join(sorted({a["status"] for a in group_areas})),
        ])
        for asset, oid in members:
            rec = originals[(asset, oid)]["attributes"]
            area = areas[(asset, oid)]
            details.append([
                gid, TITLES[asset], layer_ids[asset], oid,
                "; ".join(f"{other.upper()}/{other_id}" for other, other_id in members
                          if (other, other_id) != (asset, oid)),
                len(members), "YES" if cross else "NO",
                "; ".join(area["district"]), "; ".join(area["sector"]),
                "; ".join(area["cell"]), area["status"],
                rec.get("ASSETGROUP"), rec.get("ASSETTYPE"),
                rec.get("created_user"), rec.get("created_date"),
                rec.get("last_edited_user"), rec.get("last_edited_date"),
                "YES" if identical else "NO / NOT ESTABLISHED",
            ])
            for area_path in area["area_paths"]:
                cells = area_path.split(" / ", 2)
                if len(cells) == 3:
                    aggregation[tuple(cells)][0].add(gid)
                    aggregation[tuple(cells)][1].add((asset, oid))
    for place, (gids, members) in sorted(aggregation.items()):
        locations.append([*place, len(gids), len(members)])
    for row in [
        ("Definition", "Two or more distinct source features sharing a normalized valid GlobalID."),
        ("Cross-layer", "Same GUID used between different layers is also reported."),
        ("Not geometry duplicate search", "Different GUIDs with identical/nearby geometry are not detected."),
        ("Completeness", "Every included layer uses an initial and final count with full OBJECTID paging."),
        ("Live editing", "Counts could remain stable while source records change; treat results as a review snapshot."),
        ("Administrative lookup", "Spatial intersection; a feature can intersect multiple cells."),
        ("Projection", "KMZ reads duplicate geometries separately with out_sr=4326, verifies spatial reference."),
        ("Safety", "No source edits, repairs or duplicate deletion performed."),
        ("Performance", "Manual only; small identity pages and duplicate-only geometry."),
    ]:
        notes.append(row)
    notes.insert_rows(1)
    notes["A1"] = "Rule"
    notes["B1"] = "Explanation"
    for sheet in wb:
        sheet.freeze_panes="A2"
        sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.fill=PatternFill("solid",fgColor="17365D")
            cell.font=Font(bold=True,color="FFFFFF")
        for column in sheet.columns:
            vals=list(column)[:600]
            sheet.column_dimensions[vals[0].column_letter].width=min(
                75,max(14,max(len(str(c.value or "")) for c in vals)+2)
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith(("=","+","-","@")):
                    cell.data_type="s"  # No spreadsheet formula injection.
    wb.save(path)
    wb.close()


def _render_duplicate_kml(duplicates, geometries):
    root = ET.Element(_n("kml"))
    document=_add(root,"Document")
    _add(document,"name","REG Duplicate GlobalIDs — Lines, Poles, Transformers")
    _add(document,"description",
         "Read-only duplicate GUID audit. Each placemark matches the Excel register.")
    for asset in TITLES:
        style=_add(document,"Style",None,id="dup_"+asset)
        l=_add(style,"LineStyle")
        _add(l,"color",COLORS[asset])
        _add(l,"width","4")
        icon=_add(style,"IconStyle")
        _add(icon,"color",COLORS[asset])
        _add(icon,"scale","1.2")
    for kind in TITLES:
        folder=_add(document,"Folder")
        _add(folder,"name",TITLES[kind]+" — Duplicate GlobalIDs")
        for gid, members in sorted(duplicates.items()):
            for asset, oid in sorted(members):
                if asset != kind:
                    continue
                placemark=_add(folder,"Placemark")
                _add(placemark,"name",f"{kind.upper()} {oid} — duplicate {gid}")
                _add(placemark,"styleUrl","#dup_"+kind)
                _add(placemark,"description",
                     f"GlobalID: {gid}\nAsset: {TITLES[kind]}\n"
                     f"OBJECTID: {oid}\n"
                     "Other members: "+", ".join(
                         f"{a.upper()}/{i}" for a,i in members if (a,i)!=(kind,oid)
                     ))
                data=_add(placemark,"ExtendedData")
                for name,value in (
                    ("GlobalID",gid),("Asset",kind),("OBJECTID",oid),
                    ("Other OBJECTIDs","; ".join(
                        f"{a}/{i}" for a,i in members if (a,i)!=(kind,oid)
                    )),
                ):
                    item=_add(data,"Data",None,name=name)
                    _add(item,"value",value)
                _write_geometry(placemark,kind,geometries[(kind,oid)])
    xml=ET.tostring(root,encoding="utf-8",xml_declaration=True)
    ET.fromstring(xml)
    return xml


def scan_network_duplicate_globalids(
    output_xlsx, classes="all", service_url=None, boundary_service_url=None,
    version="sde.DEFAULT", gis=None, page_size=500, pause_seconds=0.15,
    layer_factory=None, progress=None,
):
    """Publish matching Excel + KMZ only after all three sources verify.

    The filesystem staging area is local. No report is published when any
    source scan, administrative lookup or KMZ projection fails.
    """
    if version.casefold() != "sde.default":
        raise ValueError("Full network duplicate scan requires sde.DEFAULT.")
    output=Path(output_xlsx).expanduser().resolve()
    if output.suffix.lower()!=".xlsx":
        raise ValueError("Output must be .xlsx.")
    output=output.with_name(output.stem+"_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".xlsx")
    kmz=output.with_suffix(".kmz")
    if output.exists() or kmz.exists():
        raise FileExistsError("Refusing to overwrite an existing Excel/KMZ pair.")
    base=(service_url or os.getenv("UN_REPORT_SERVICE_URL") or
          DEFAULT_LINE_SERVICE_URL).rstrip("/")
    boundary=(boundary_service_url or os.getenv("UN_SECTOR_BOUNDARY_URL") or
              DEFAULT_BOUNDARY_SERVICE_URL).rstrip("/")
    report=progress or (lambda message: None)
    if layer_factory is None:
        from arcgis.features import FeatureLayer
        active=_get_gis(gis)
        def layer_factory(url):
            return FeatureLayer(url,active)
    layers={asset:layer_factory(_layer_url(base,layerid))
            for asset,layerid,_,_ in SOURCES}
    cells=layer_factory(_layer_url(boundary,CELL_LAYER_ID))
    groups,counts,invalid=_all_ids(
        layers,version,classes,page_size,pause_seconds,progress=report,
    )
    originals=_fetch_selected(layers,groups,version,progress=report)
    areas=_area_index(cells,groups,originals,progress=report)
    wgs84=_fetch_wgs84(layers,groups,originals,version)
    output.parent.mkdir(parents=True,exist_ok=True)
    staged_excel=output.with_name(output.name+".pending")
    staged_kmz=kmz.with_name(kmz.name+".pending")
    try:
        _report_excel(
            staged_excel,groups,originals,areas,counts,invalid,
            base,boundary,version,classes,
        )
        xml=_render_duplicate_kml(groups,wgs84)
        with ZipFile(staged_kmz,"w",compression=ZIP_DEFLATED) as archive:
            archive.writestr("doc.kml",xml)
        # Do not expose half a pair when KMZ writing/validation fails.
        os.replace(staged_kmz,kmz)
        try:
            os.replace(staged_excel,output)
        except Exception:
            kmz.unlink(missing_ok=True)
            raise
    finally:
        staged_excel.unlink(missing_ok=True)
        staged_kmz.unlink(missing_ok=True)
    return {
        "output_xlsx":str(output), "output_kmz":str(kmz),
        "read_only":True,"scope":classes,"version":version,
        "scanned_by_asset":counts,
        "invalid_or_missing_globalids_by_asset":invalid,
        "duplicate_globalid_groups":len(groups),
        "duplicate_feature_records":sum(len(group) for group in groups.values()),
        "cross_layer_duplicate_groups":sum(
            len({a for a,_ in group})>1 for group in groups.values()
        ),
        "records_without_admin_location":sum(
            area["status"] in ("NOT_FOUND","GEOMETRY_OR_CRS_MISSING")
            for area in areas.values()
        ),
        "kmz_features":len(wgs84),
    }
